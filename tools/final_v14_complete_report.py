#!/usr/bin/env python3
"""Render the complete v14 scientific report; the separate bundle auditor seals it.

Only sealed scientific results and the preserved, audited source-report snapshot
are read. A pending grid stage stays explicitly pending and cannot be promoted
to a completed analysis by this reporting program.
"""
from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
import shutil
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

REPO = Path(__file__).resolve().parents[1]
RUN = REPO / "reports/reruns/final_v14_additions_20260903"
SNAPSHOT = REPO / "reports/snapshots/final_v14_source_working_pre_completion_20260905"
OUT = REPO / "reports/final_v14"
CORRESPONDENCE = RUN / "e4v_correspondence_replay/runs/add086ae82ef5820b15847827e7eb8d9d25d7c3740cc69c38b449ab8e7ed1274/results.json"
GRID = RUN / "grid_offset/locked_extraction/results.json"
TASKS = ("codon", "g12d_broad", "allele1", "allele2", "g12c")
TASK_LABELS = {"codon": "Codon", "g12d_broad": "Broad G12D", "allele1": "Allele 1", "allele2": "Allele 2", "g12c": "G12C"}
COHORTS = ("cptac_primary", "rih_primary", "orion_cpht", "rih_metastatic", "sr1482_metastatic")
COHORT_LABELS = {"cptac_primary": "CPTAC primary", "rih_primary": "RIH primary", "orion_cpht": "Orion", "rih_metastatic": "RIH metastatic", "sr1482_metastatic": "SR1482 metastatic"}
CEILINGS = "single-reader, non-naive names; concept fidelity and association, not mediation or mechanism; source-frozen concept-score association, not proven tissue biology; concept-space consistency, not information absence; no deployment threshold, calibration, clinical utility, or feedback to an Aim 1–3 verdict."


def identity(path: Path) -> dict:
    return {"path": str(path.resolve()), "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "size_bytes": path.stat().st_size}


def verify(item: dict) -> Path:
    path = Path(item["path"])
    observed = identity(path)
    if any(observed[k] != item[k] for k in ("sha256", "size_bytes")):
        raise ValueError(f"Artifact drift: {path}")
    return path


def sealed(path: Path) -> dict:
    pin = json.loads(Path(str(path) + ".seal.json").read_text())
    if identity(path)["sha256"] != pin.get("sha256", pin.get("artifact", {}).get("sha256")):
        raise ValueError(f"Seal mismatch: {path}")
    return json.loads(path.read_text())


def fmt(value, digits: int = 3) -> str:
    if value is None or (isinstance(value, (float, np.floating)) and not np.isfinite(value)):
        return "—"
    if isinstance(value, (float, np.floating)):
        return f"{value:.{digits}f}"
    return str(value)


def ci(point, bounds) -> str:
    if point is None:
        return "Not estimable"
    if bounds is None or any(v is None for v in bounds):
        return f"{fmt(point)} [CI not estimable]"
    return f"{fmt(point)} [{fmt(bounds[0])}, {fmt(bounds[1])}]"


def metric(item: dict) -> str:
    return ci(item.get("estimate"), item.get("ci95"))


def md_table(rows: list[dict] | pd.DataFrame) -> str:
    frame = pd.DataFrame(rows)
    if frame.empty:
        return "No estimable rows are available."
    def cell(value):
        return fmt(value).replace("|", "\\|").replace("\n", " ")
    return "\n".join(["| " + " | ".join(map(str, frame.columns)) + " |", "| " + " | ".join("---" for _ in frame.columns) + " |",
                      *["| " + " | ".join(cell(v) for v in row) + " |" for row in frame.itertuples(index=False, name=None)]])


def add_metric(row: dict, prefix: str, item: dict) -> None:
    row[prefix] = item.get("estimate")
    bounds = item.get("ci95") or [None, None]
    row[prefix + "_ci95_low"], row[prefix + "_ci95_high"] = bounds
    for key in ("fwer_lower99", "fwer_upper99", "status", "n_bootstrap"):
        if key in item:
            row[prefix + "_" + key] = item[key]


def load_inputs() -> tuple[dict, list[dict]]:
    snapshot_pin = SNAPSHOT.with_name(SNAPSHOT.name + ".identity.json")
    snapshot = sealed(snapshot_pin)
    for row in snapshot["report_path_mapping"]:
        verify(row["snapshot"])
    paths = {"m1": RUN / "e4m1_results/results.json", "context": RUN / "e4m1_context/results.json",
             "names": RUN / "e4v_post_reader_xlsx/naming/naming_freeze.json", "m2": RUN / "e4m2_source_frozen/target_results.json",
             "m3": RUN / "module_iii/results.json", "correspondence": CORRESPONDENCE, "grid": GRID,
             "variant_mapping": RUN / "e4v_variant_mapping/results.json"}
    values, pins = {}, [identity(snapshot_pin)]
    for name, path in paths.items():
        if not path.exists() and name == "grid":
            values[name] = None
            continue
        values[name] = sealed(path)
        pins.append(identity(path))
    if values["names"]["name_gate_status"] != "NAME_GATE_FAIL":
        raise ValueError("This report implements the accepted final NAME_GATE_FAIL campaign")
    for key, expected in (("m2", "MODULE_II_ANALYSIS_COMPLETE"), ("m3", "COMPLETE"),
                          ("correspondence", "CORRESPONDENCE_GEOMETRY_COMPLETE"), ("variant_mapping", "CANONICAL_NINE_VARIANT_MAPPING_COMPLETE")):
        if values[key].get("status") != expected:
            raise ValueError(f"Required completed scientific input is unavailable: {key}")
    for name in ("context", "m2", "grid", "variant_mapping"):
        if values.get(name):
            artifacts = values[name].get("artifacts", [])
            for item in artifacts.values() if isinstance(artifacts, dict) else artifacts:
                verify(item)
    if values["m3"]:
        for row in values["m3"]["representations"]["ALL32"].values():
            verify(row["bootstrap"])
    if values["grid"] and values["grid"].get("status") not in ("GRID_OFFSET_SENSITIVITY_COMPLETE", "GRID_OFFSET_NOT_EVALUABLE"):
        raise ValueError("Grid JSON exists but is not a final scientific result")
    return values, pins


def module1_tables(values: dict) -> dict[str, pd.DataFrame]:
    rows = []
    m1 = values["m1"]["ALL32"]
    for population, block in [("pooled", m1["pooled"]), *m1["subcohort_fidelity"].items()]:
        for key, item in block.items():
            if isinstance(item, dict) and "estimate" in item:
                row = {"population": population, "estimand": key}
                add_metric(row, "value", item)
                for count in ("n", "n_patients", "finite_draws", "undefined_draws"):
                    if count in item:
                        row[count] = item[count]
                rows.append(row)
    return {"M1 all estimands": pd.DataFrame(rows)}


def module2_tables(data: dict | None) -> dict[str, pd.DataFrame]:
    if data is None:
        return {}
    rows = []
    for rep, result in data["representations"].items():
        groups = [(c, result["cohorts"][c]) for c in COHORTS]
        groups += [("equal_cohort_macro", result["equal_cohort_macro"]), ("concatenated_247_continuity_only", result["concatenated_247_continuity_only"])]
        for cohort, item in groups:
            row = {"representation": rep, "population": cohort}
            for key in ("patients", "mutant", "wild_type", "concept_auroc", "full_auroc", "concept_minus_full", "concept_one_sided_97_5_lower"):
                row[key] = item.get(key)
            for key in ("concept_ci95", "full_ci95", "concept_minus_full_ci95"):
                row[key + "_low"], row[key + "_high"] = item[key]
            for key, value in item.get("retention_ratio", {}).items():
                if key == "ci95":
                    row["ratio_ci95_low"], row["ratio_ci95_high"] = value or [None, None]
                else:
                    row["ratio_" + key] = value
            rows.append(row)
    root = RUN / "e4m2_source_frozen"
    tables = {"M2 performance": pd.DataFrame(rows)}
    for title, name in (("M2 RIH interactions", "metastatic_interactions_RIH.csv"), ("M2 SR1482 interactions", "metastatic_interactions_SR1482.csv"),
                        ("M2 OOD by KRAS", "target_ood_by_kras.csv"), ("M2 prototype distributions", "target_prototype_distributions_by_kras.csv"),
                        ("M2 OOD before outcomes", "target_ood_before_outcome_join.csv")):
        path = root / name
        if name == "target_ood_before_outcome_join.csv":
            score_seal = sealed(root / "target_score_seal.json")
            verify(next(a for a in score_seal["artifacts"] if Path(a["path"]).name == name))
        tables[title] = pd.read_csv(path, float_precision="round_trip")
    return tables


def module3_tables(data: dict | None) -> dict[str, pd.DataFrame]:
    if data is None:
        return {}
    contract = sealed(RUN / "module_iii/contract.json")
    tasks, draws, inherited = [], [], []
    for task in TASKS:
        item = data["representations"]["ALL32"][task]
        count = contract["task_inputs"]["fine__" + task]
        row = {"task": task, "patients": count["patients"], "positive": count["positive"], "negative": count["patients"] - count["positive"],
               "inherited_v13_consensus": item["inherited_v13_consensus"], **item["statuses"]}
        for name in ("fine", "mil_fine", "concept_minus_mil"):
            add_metric(row, name, item[name])
        row["mil_interval_scope"] = item["mil_fine_interval_scope"]
        tasks.append(row)
        for draw, values in [("canonical", item["canonical"]), *item["repeated_draws"].items()]:
            r = {"task": task, "draw": draw, "patients_per_vector": count["patients"], "positive": count["positive"], "negative": count["patients"] - count["positive"],
                 "role": "descriptive only" if draw == "canonical" else "governed repeated control", "verdict": values.get("verdict", "NOT_USED_FOR_STATUS")}
            for name in ("fine", "control", "control_minus_fine"):
                add_metric(r, name, values[name])
            draws.append(r)
        inherited.append({"task": task, "parent_consensus": item["inherited_v13_consensus"],
                          "canonical_parent_json": json.dumps(item["inherited_v13_canonical_unchanged"], sort_keys=True),
                          "repeated_parent_json": json.dumps(item["inherited_v13_repeated_unchanged"], sort_keys=True)})
    return {"M3 task summary": pd.DataFrame(tasks), "M3 all control draws": pd.DataFrame(draws), "M3 unchanged parent rows": pd.DataFrame(inherited)}


def save_figure(fig, output: Path, name: str) -> None:
    for extension in ("pdf", "svg", "png"):
        fig.savefig(output / "figures" / f"{name}.{extension}", dpi=220, bbox_inches="tight")
    plt.close(fig)


def interval_mark(ax, point: float, bounds, y: float, color: str, label=None) -> None:
    ax.plot(bounds, [y, y], color=color, lw=1.6)
    ax.plot(point, y, "o", color=color, ms=5, label=label)


def figures(data: dict, output: Path) -> None:
    plt.rcParams.update({"font.size": 10, "axes.spines.top": False, "axes.spines.right": False, "pdf.fonttype": 42, "svg.fonttype": "none"})
    if data["m2"]:
        rep = data["m2"]["representations"]["ALL32"]
        entries = [(COHORT_LABELS[c], rep["cohorts"][c]) for c in COHORTS] + [("Primary macro (½ each)", rep["equal_cohort_macro"])]
        fig, axes = plt.subplots(1, 2, figsize=(11.5, 5), sharey=True)
        for y, (_, row) in enumerate(entries):
            for shift, key, color, label in ((-.12, "full", "#244d42", "Frozen full WSI"), (.12, "concept", "#bc7750", "ALL32 concept")):
                interval_mark(axes[0], row[key + "_auroc"], row[key + "_ci95"], y + shift, color, label if y == 0 else None)
            interval_mark(axes[1], row["concept_minus_full"], row["concept_minus_full_ci95"], y, "#725384")
        axes[0].set_yticks(range(len(entries)), [x[0] for x in entries])
        axes[0].invert_yaxis()
        axes[0].axvline(.5, color="#a5b4ad", ls="--", lw=1)
        axes[1].axvline(0, color="#a5b4ad", ls="--", lw=1)
        axes[0].set_xlabel("AUROC · descriptive 95% CI")
        axes[1].set_xlabel("Paired concept − full AUROC · 95% CI")
        fig.legend(*axes[0].get_legend_handles_labels(), loc="upper center", bbox_to_anchor=(.5, .90), ncol=2, frameon=False, fontsize=9)
        axes[0].set_title("A  Source-frozen discrimination", loc="left", weight="bold")
        axes[1].set_title("B  Paired performance gap", loc="left", weight="bold")
        fig.suptitle("Module II · archived cohorts, fixed source model", weight="bold")
        fig.text(.5, .01, "10,000 paired patient draws · only CPTAC primary and RIH primary control H2.1 · no fresh validation claim", ha="center", fontsize=9)
        fig.tight_layout(rect=(0, .035, 1, .83))
        save_figure(fig, output, "aim4_module2")
    if data["m3"]:
        fig, axes = plt.subplots(1, 2, figsize=(10.6, 4.6), sharey=True)
        for y, task in enumerate(TASKS):
            row = data["m3"]["representations"]["ALL32"][task]
            for shift, key, color, label in ((-.12, "mil_fine", "#244d42", "Frozen MIL"), (.12, "fine", "#bc7750", "ALL32 concept")):
                interval_mark(axes[0], row[key]["estimate"], row[key]["ci95"], y + shift, color, label if y == 0 else None)
            interval_mark(axes[1], row["concept_minus_mil"]["estimate"], row["concept_minus_mil"]["ci95"], y, "#725384")
            axes[1].plot(row["concept_minus_mil"]["fwer_lower99"], y, "|", color="black", ms=11)
        axes[0].set_yticks(range(5), [TASK_LABELS[t] for t in TASKS])
        axes[0].invert_yaxis()
        axes[0].axvline(.5, color="#a5b4ad", ls="--", lw=1)
        axes[0].axvline(.6, color="#c6b7a7", ls=":", lw=1)
        axes[1].axvline(0, color="#a5b4ad", ls="--", lw=1)
        fig.legend(*axes[0].get_legend_handles_labels(), loc="upper center", bbox_to_anchor=(.5, .90), ncol=2, frameon=False, fontsize=9)
        axes[0].set_xlabel("Fine-task AUROC · descriptive 95% CI")
        axes[1].set_xlabel("Paired concept − MIL AUROC · 95% CI")
        axes[0].set_title("A  Fine-task ranking", loc="left", weight="bold")
        axes[1].set_title("B  Paired comparison", loc="left", weight="bold")
        fig.suptitle("Module III · related representation and readout", weight="bold")
        fig.text(.5, .01, "20,000 partially paired draws · black ticks = adjusted one-sided 99% lower bounds · statuses also require matched controls", ha="center", fontsize=8.5)
        fig.tight_layout(rect=(0, .035, 1, .83))
        save_figure(fig, output, "aim4_module3")
    c = data["correspondence"]
    image = np.array([[r[axis]["cosine"] for r in c["variants"]] for axis in ("p17", "p28")])
    fig, ax = plt.subplots(figsize=(10.5, 3.7))
    im = ax.imshow(image, vmin=.8, vmax=1, cmap="YlGnBu", aspect="auto")
    ax.set_xticks(range(9), [f"k={r['k']}\ns{r['seed'] % 100:02d}" for r in c["variants"]], fontsize=10)
    ax.set_yticks([0, 1], ["Legacy p17", "Legacy p28"])
    for row in range(2):
        for col in range(9):
            ax.text(col, row, f"{image[row,col]:.3f}", ha="center", va="center", color="white" if image[row,col] > .95 else "black", fontsize=8)
    fig.colorbar(im, ax=ax, label="Backprojected centroid cosine", fraction=.025, pad=.02)
    ax.set_title("Legacy correspondence · direct distinct-pair assignment in every variant", loc="left", weight="bold")
    fig.text(.5, .095, "Seed key: s19 = 20260819; s20 = 20260820; s21 = 20260821.", ha="center", fontsize=9)
    fig.text(.5, .02, "Geometry is descriptive; failed naming gate independently precludes named correspondence or spotlights.", ha="center", fontsize=9)
    fig.tight_layout(rect=(0, .18, 1, 1))
    save_figure(fig, output, "aim4_correspondence")
    if data["grid"] and data["grid"]["status"] == "GRID_OFFSET_SENSITIVITY_COMPLETE":
        g = data["grid"]
        fig, axes = plt.subplots(1, 2, figsize=(10, 4))
        axes[0].hist([r["cosine_similarity"] for r in g["patient_cosine"]], bins=20, color="#437f70", edgecolor="white")
        axes[0].set(xlabel="Paired original/offset profile cosine", ylabel="Patients")
        for r in g["prototype_icc"]:
            if r["icc"] is not None:
                axes[1].plot(r["prototype_id"], r["icc"], "o", color="#437f70", ms=4)
        axes[1].set(xlabel="Canonical prototype", ylabel="Absolute-agreement ICC(2,1)")
        fig.suptitle("Source grid sensitivity · fixed 100-patient diagonal half-stride offset", weight="bold")
        fig.tight_layout()
        save_figure(fig, output, "aim4_grid_offset")
    fig, ax = plt.subplots(figsize=(11, 4))
    ax.set(xlim=(0, 11), ylim=(0, 4))
    ax.axis("off")
    ax.add_patch(FancyBboxPatch((.3, .45), 2.2, 3.1, boxstyle="round,pad=.12", fc="#e4eee8", ec="#244d42"))
    ax.text(1.4, 2, "④ Aim 4\nInterpretation", ha="center", va="center", fontsize=14, weight="bold")
    for y, aim, label, complete in ((3.15, "① Aim 1", "Concept fidelity /\nconcept-only discrimination", True),
                                   (2, "② Aim 2", "Source-frozen cross-cohort\nconcept discrimination", data["m2"] is not None),
                                   (.85, "③ Aim 3", "Concept-space consistency", data["m3"] is not None)):
        ax.add_patch(FancyBboxPatch((8.45, y-.35), 2, .7, boxstyle="round,pad=.10", fc="#f4f1eb", ec="#8b7861"))
        ax.text(9.45, y, aim, ha="center", va="center", fontsize=13)
        ax.add_patch(FancyArrowPatch((2.65, y), (8.25, y), arrowstyle="-|>", mutation_scale=18, lw=1.8, color="#437f70", linestyle="-" if complete else "--"))
        ax.text(5.45, y+.12, label, ha="center", va="bottom", fontsize=10)
    ax.set_title("Figure 1C · one-way interpretive links", loc="left", weight="bold")
    fig.text(.5, .005, "Solid = corresponding module completed with sealed artifacts; dashed = pending. Line style records completion, never feedback or independent validation.", ha="center", fontsize=8.5)
    save_figure(fig, output, "figure1c_interpretive_links")


def module2_text(data: dict | None, tables: dict) -> str:
    if data is None:
        return "## Interpretation of Aim 2: source-frozen cross-cohort concept discrimination\n\nPending sealed target analyses. No target estimate or gate is inferred.\n"
    rep = data["representations"]["ALL32"]
    rows, ratios = [], []
    for c in COHORTS:
        r = rep["cohorts"][c]
        rows.append({"Cohort": COHORT_LABELS[c], "Patients (mutant/WT)": f"{r['patients']} ({r['mutant']}/{r['wild_type']})", "Concept AUROC [95% CI]": ci(r["concept_auroc"], r["concept_ci95"]),
                     "Full AUROC [95% CI]": ci(r["full_auroc"], r["full_ci95"]), "Concept − full [95% CI]": ci(r["concept_minus_full"], r["concept_minus_full_ci95"])})
        ratio = r["retention_ratio"]
        ratios.append({"Cohort": COHORT_LABELS[c], "Above-chance ratio [95% CI]": ci(ratio["point"], ratio["ci95"]), "Finite / 10,000": ratio["finite_draws"], "Undefined draws": ratio["undefined_draws"], "Status": ratio["status"]})
    macro, pool = rep["equal_cohort_macro"], rep["concatenated_247_continuity_only"]
    interactions = []
    for family, item in data["interactions"].items():
        frame = tables[f"M2 {family} interactions"]
        interactions.append({"Family": family, "Patients": item["patients"], "Source-family exposed": item["source_family_exposed"],
                             "Role interaction BH q<0.05 / 32": int((frame.bh_q < .05).sum()),
                             "Change in mutant−WT mean-logit separation [95% CI]": ci(item["direct_logit_separation_change"], item["logit_separation_change_ci95"])})
    return f"""## Interpretation of Aim 2: source-frozen cross-cohort concept discrimination

The fixed ALL32 source model did not establish above-chance discrimination in both controlling primary cohorts: **{data['gates']['ALL32']}**. RIH primary's adjusted lower AUROC bound was {rep['cohorts']['rih_primary']['concept_one_sided_97_5_lower']:.3f}; CPTAC primary's was {rep['cohorts']['cptac_primary']['concept_one_sided_97_5_lower']:.3f}. Both must exceed 0.5. **{data['gates']['NAMED_REF']}** remains fixed by the failed naming gate.

{md_table(rows)}

The equal-cohort primary macro uses fixed ½ weights: concept AUROC **{ci(macro['concept_auroc'], macro['concept_ci95'])}**, full AUROC **{ci(macro['full_auroc'], macro['full_ci95'])}**, and paired concept-minus-full difference **{ci(macro['concept_minus_full'], macro['concept_minus_full_ci95'])}**. The positive macro lower bound does not replace the two-cohort gate. The 247-patient concatenated continuity estimate is concept AUROC {ci(pool['concept_auroc'], pool['concept_ci95'])}, full AUROC {ci(pool['full_auroc'], pool['full_ci95'])}, and paired gap {ci(pool['concept_minus_full'], pool['concept_minus_full_ci95'])}; its cross-cohort rankings are not the controlling estimand.

![Source-frozen cohort performance and gaps](figures/aim4_module2.png)

Above-chance retention ratios are untruncated descriptive contrasts, not shares of signal. Nonpositive bootstrap full-score denominators are undefined; at least 9,500 finite draws are required for a ratio interval. Paired AUROC differences remain controlling.

{md_table(ratios)}

These are locked secondary analyses of previously opened cohorts. Target profiles, distances and logits were sealed without outcome columns before the outcome join. No target centering, feature selection, refitting, calibration, thresholding or favorable-support filtering occurred. Confidence intervals pair concept/full scores within each cohort×KRAS patient draw; macro draws independently resample the two primary cohorts.

### Class-conditional tissue-role characterization and source support

{md_table(interactions)}

All eight RIH dual-role patients were removed from both tissue-role arms before interaction analysis. RIH remains family-naive retrospective evidence; SR1482 is source-family exposed. Every panel reports all four role×KRAS abundance means, both class-specific role shifts, their interaction, 10,000 bootstrap intervals and 10,000 within-KRAS role permutations with BH across all 32 coordinates. The coefficient-weighted sum exactly reconstructs a linear mean-logit contrast; it is not an AUROC decomposition or a causal explanation. No named morphology spotlight is authorized.

Complete panels: [RIH interactions](tables/M2_RIH_interactions.csv), [SR1482 interactions](tables/M2_SR1482_interactions.csv), [source-distance diagnostics before outcomes](tables/M2_OOD_before_outcomes.csv), [diagnostics by KRAS](tables/M2_OOD_by_KRAS.csv), and [all cohort/prototype distributions](tables/M2_prototype_distributions.csv). Distance summaries average nonempty slides equally; no assigned tile is undefined with zero support. Orion is a small retrospective cross-protocol sensitivity after cyclic-IF processing, never a third controlling conventional-primary cohort.
"""


def module3_text(data: dict | None, tables: dict) -> str:
    if data is None:
        return "## Interpretation of Aim 3: concept-space resolution consistency\n\nPending sealed concept-resolution analyses. No task verdict is inferred.\n"
    performance, statuses, draw_rows = [], [], []
    for task in TASKS:
        r = data["representations"]["ALL32"][task]
        count = tables["M3 task summary"].set_index("task").loc[task]
        performance.append({"Task": TASK_LABELS[task], "Patients (+/−)": f"{count['patients']} ({count['positive']}/{count['negative']})",
                            "Concept fine AUROC [95% CI]": metric(r["fine"]), "MIL fine AUROC [new paired 95% CI]": metric(r["mil_fine"]),
                            "Concept − MIL [95% CI]": metric(r["concept_minus_mil"]), "Gap adjusted L99": r["concept_minus_mil"]["fwer_lower99"]})
        statuses.append({"Task": TASK_LABELS[task], "Inherited v13": r["inherited_v13_consensus"], **{k: r["statuses"][k] for k in ("control", "boundary", "fine_signal", "superiority")}})
        for draw, d in [("canonical", r["canonical"]), *r["repeated_draws"].items()]:
            draw_rows.append({"Task": TASK_LABELS[task], "Control draw": draw, "Fine AUROC [95% CI]": metric(d["fine"]),
                              "Control AUROC [95% CI]": metric(d["control"]), "Control − fine [95% CI]": metric(d["control_minus_fine"]),
                              "Fine U99": d["fine"].get("fwer_upper99"), "Control L99": d["control"].get("fwer_lower99"),
                              "Gap L99": d["control_minus_fine"].get("fwer_lower99"), "Draw verdict": d.get("verdict", "DESCRIPTIVE_ONLY")})
    return f"""## Interpretation of Aim 3: concept-space resolution consistency

The task definitions are: codon, G12 versus other KRAS-mutant tumors; broad G12D, G12D versus non-G12D among KRAS mutants; allele 1, G12D versus other G12; allele 2, G12V versus other G12; and G12C, G12C versus other G12. Internal task codes and their frozen positive/negative rosters are unchanged.

The concept representation supplies no supported fine-signal or concept-versus-MIL superiority call in any of the five tasks. Codon, allele 1 and G12C have inadequate repeated matched-WT controls, so their boundary calls are not evaluable. Broad G12D and allele 2 have adequate controls but do not establish the prespecified nonresolution pattern. These task-specific results neither prove absence of information in H&E nor revise any FINAL-v13 verdict.

{md_table(performance)}

MIL point estimates are inherited unchanged; the displayed paired-comparison MIL intervals are newly computed v14 bootstrap intervals. The original parent canonical/repeated rows remain unchanged in [the parent-row table](tables/M3_unchanged_parent_rows.csv) and the sealed result JSON. No v13 ceiling is upgraded or refuted where its inherited consensus was never established.

![Concept fine-task ranking and paired comparison](figures/aim4_module3.png)

The four status families remain separate; a missing learnability control cannot be hidden in a campaign average.

{md_table(statuses)}

### Canonical and all three repeated matched controls

{md_table(draw_rows)}

All repeated rows use 20,000 partially paired patient draws within source-subcohort×frozen-fold strata. Each draw reuses shared positives and fine negatives across all three controls, while independently drawing each control-negative pool. Adjusted one-sided 99% bounds determine statuses; descriptive 95% intervals do not. Canonical controls use a separate 10,000-draw bootstrap and never determine cross-draw status. Full numerical columns, including both adjusted bounds and the practical AUROC≥0.60 flag, are in [the task table](tables/M3_task_summary.csv) and [all control draws](tables/M3_all_control_draws.csv). No named representation was fitted under NAME_GATE_FAIL.
"""


def correspondence_text(data: dict) -> str:
    rows = [{"k": r["k"], "Seed": r["seed"], "Canonical": r["canonical"], "Legacy p17 → prototype": r["p17"]["prototype_id"], "p17 cosine": r["p17"]["cosine"],
             "Legacy p28 → prototype": r["p28"]["prototype_id"], "p28 cosine": r["p28"]["cosine"]} for r in data["variants"]]
    return f"""## Legacy correspondence

All four pinned legacy inputs were recovered byte-identically. Direct, joint one-to-one assignment passes cosine≥0.80 for both anchors in all nine variants. Canonical legacy p17 maps to v14 prototype 17 (cosine {data['axes']['p17']['canonical_geometry']['cosine']:.6f}); legacy p28 maps to v14 prototype 22 (cosine {data['axes']['p28']['canonical_geometry']['cosine']:.6f}). Both named correspondence calls remain **CORRESPONDENCE_NOT_EVALUABLE** because NAME_GATE_FAIL independently applies. Geometry cannot authorize a named spotlight or change a feature.

{md_table(rows)}

![All nine direct legacy geometry comparisons](figures/aim4_correspondence.png)

The legacy vocabulary was target-inclusive/transductive and serves only this historical descriptive comparison. Earlier missing-input reports remain preserved; the restored-archive result supersedes their operational availability status.
"""


def variant_mapping_text(data: dict) -> str:
    rows = [{"k": r["k"], "Seed": r["seed"], "Method": r["method"], "Matched / 32 canonical": r["matched_pairs"],
             "Cosine ≥0.80 / 32": r["pairs_cosine_ge_0_80"], "Unmatched canonical": r["unmatched_canonical"],
             "Matched cosine minimum": r["matched_cosine_min"]} for r in data["variants"]]
    return f"""## Canonical dictionary stability across nine variants

{md_table(rows)}

These descriptive maps compare the full canonical dictionary with each grid member: Hungarian maximum-cosine matching when k=32 and mutual-nearest matching when k differs. The two alternative k=32 seeds retain 29/32 pairs above cosine 0.80, so stability is not uniform across all coordinates. Each k=24 map leaves eight canonical concepts unmatched; k=40 maps retain 30–32 canonical pairs. Shared PCA/sample identities were verified without refitting, target data or changes to names, models or gates. The [complete 288-row mapping table](tables/Canonical_variant_mappings.csv) retains unmatched coordinates. These maps are distinct from the two-anchor legacy correspondence above.
"""


def grid_text(data: dict | None) -> str:
    if data is None:
        return "## Raw-slide grid sensitivity\n\n**PENDING_NOT_TESTED.** The raw-slide computation on the sealed 100-patient, all-listed-slide roster remains in progress. No cosine, ICC, scientific insufficiency verdict or reduced-roster replacement is inferred. The [first-slide operational benchmark](../reruns/final_v14_additions_20260903/grid_offset/locked_extraction/benchmarks/first_slide.json) records measured throughput, sampled resource use and a coarse projection; it supplies no scientific agreement estimate. This report remains a working draft until the final grid result is sealed.\n"
    if data["status"] == "GRID_OFFSET_NOT_EVALUABLE":
        return "## Raw-slide grid sensitivity\n\n**GRID_OFFSET_NOT_EVALUABLE** after the complete authenticated eligibility audit: " + data["reason"] + ".\n\n" + md_table([{"Subcohort": k, "Eligible patients": v} for k, v in data["eligible_counts"].items()]) + "\n\nNo cross-subcohort substitution, patient roster draw or extraction occurred. Canonical-grid dependence remains a limitation.\n"
    undefined = sum(r["icc"] is None for r in data["prototype_icc"])
    rows = [{"Prototype": r["prototype_id"], "ICC(2,1)": r["icc"], "Original support": r["n_nonzero_original"], "Offset support": r["n_nonzero_offset"], "Status": r["status"]} for r in data["prototype_icc"]]
    return f"""## Raw-slide grid sensitivity

All {data['n_patients']} selected patients were analyzed with every listed slide. Original/offset profile cosine had mean **{data['cosine_mean']:.4f}**, median **{data['cosine_median']:.4f}**, and minimum **{data['cosine_min']:.4f}**. {undefined}/32 prototype ICCs were undefined, with support retained below. No selected patient or slide was substituted.

![Paired patient-profile and prototype grid agreement](figures/aim4_grid_offset.png)

{md_table(rows)}

The raw-slide computation used one diagonal half-stride shift, unchanged exact-MPP footprint and canonical HEST masks, boundary omission, frozen UNI-v1, and the frozen reference vocabulary. Float16 storage precision was matched before assignment; patient profiles weight slides equally. This is a descriptive grid-phase sensitivity, not a cached-feature rerun. Low agreement remains a tiling-dependence limitation and cannot select a vocabulary or alter any primary result. [Paired patients](tables/Grid_patient_cosine.csv) and [prototype support/ICC](tables/Grid_prototype_ICC.csv) retain every selected row. The [final operational benchmark](../reruns/final_v14_additions_20260903/grid_offset/locked_extraction/benchmarks/final.json) records actual throughput and sampled resources separately from scientific results.
"""


def deviation_text() -> str:
    return """# FINAL-v14 · Deviations, clarifications and execution history

The frozen preregistration remains the controlling method. The following records are preserved alongside the completed results; none changes an Aim 1–3 parent estimate or verdict.

| Record | Resolution and scientific consequence | Controlling artifact |
|---|---|---|
| Final mentor return accepted with missing fields | Explicit user authorization accepted the return without further form completion. Original categories/comments and missing confidence, artifact, reviewer/date and attestation values were retained. Same-day mentor return is user-reported provenance. | [Pre-unblinding amendment](../reruns/final_v14_additions_20260903/e4v_post_reader_xlsx/governance_amendment.json) |
| Conservative repeatability and name eligibility | Missing structured evidence was not inferred from free text or images. NAME_GATE_FAIL excludes all named models and formal named spotlights; all 32 unsupervised predictors remain unchanged. | [Naming freeze](../reruns/final_v14_additions_20260903/e4v_post_reader_xlsx/naming/naming_freeze.json) |
| Source dictionary and inherited clinical interpretation | Source-only implementation contract binds the shallow grids and literal clinical/MSI normalization. Stage-known comparison uses the inherited fixed OOF-vector restriction; no subgroup refit. | [Source contract](../reruns/final_v14_additions_20260903/e4m1_source/source_fit_contract.json) and [evaluation contract](../reruns/final_v14_additions_20260903/e4m1_results/evaluation_contract.json) |
| Archive outage and independent Module-I evaluation | The unavailable archive was an operational dependency, not a scientific failure. An ordering amendment permitted independently frozen Module-I evaluation while Module-III inputs were unavailable. Later methods could not respond to Module-I results. | [Evaluation ordering](../reruns/final_v14_additions_20260903/e4m1_results/evaluation_contract.json) |
| Attention FP32 boundary | One original attention mass exceeded one by 6.046732892×10⁻⁹. A sealed v2 contract clipped that numerical boundary within a 2×10⁻⁶ tolerance before transformation; original floats remain unchanged. Abundance remained strict. | [Precision receipt](../reruns/final_v14_additions_20260903/e4m1_context/attention_boundary_precision.json) |
| Singular context test | Prototype 13's source-subcohort HC3 joint test was not estimable and retained p=1 within its fixed family. There are 275 attempted and 274 estimable tests. | [Context results](../reruns/final_v14_additions_20260903/e4m1_context/results.json) |
| Decimal constant-column robustness | The separately frozen v2 helper detects exact constant coordinates by range and assigns zero coefficients. It was used for Modules II/III, whose methods were still pending when the helper was frozen. The original Module-I computations and audit remain unchanged. | [Module-II source bundle](../reruns/final_v14_additions_20260903/e4m2_source_frozen/source_scoring_bundle.json); [Module-I no-impact audit](../reruns/final_v14_additions_20260903/e4m1_source_audits/exact_constant_detection_20260905/audit.json) |
| Legacy geometry availability | Initial missing-input receipts and a corrected path diagnosis are preserved. After archive restoration all four original pins matched, and a new versioned geometry result replaced no earlier artifact. NAME_GATE_FAIL still governs named criteria. | [Versioned geometry directory](../reruns/final_v14_additions_20260903/e4v_correspondence_replay/runs/add086ae82ef5820b15847827e7eb8d9d25d7c3740cc69c38b449ab8e7ed1274/results.json) |
| Historical HEST provenance schema | Canonical receipt implementation hashes were recognized using source-only schema evidence and pinned UNI/HEST checkpoint hashes. Historical implementation source was not independently recovered. Older candidate masks were never substituted for canonical inputs. | [Historical schema inventory](../reruns/final_v14_additions_20260903/grid_offset/historical_receipt_schema_20260905/receipt_inventory.json) |
| Canonical mapping after main results | The canonical-to-nine-variant descriptive geometry was computed after the main module results under a pre-computation quarantine contract. It could not select a model, reinterpret names, replace the canonical dictionary or change any prior result or gate. | [Geometry quarantine contract](../reruns/final_v14_additions_20260903/e4v_variant_mapping/contract.json) |

Operational job logs and throughput/resource receipts remain under the governed rerun. The [initial grid resource-monitor launch failed](../reruns/final_v14_additions_20260903/grid_offset/execution_handoff_20260905/grid_offset.monitor_start_failure.log) because the worktree runtime lacked `psutil`, before any scientific child process launched. The unchanged script was then launched with the verified base runtime. This was an operational monitoring interruption, with no scientific method or model change. In Module III the source fit completed before the first RAM poll; evaluation reports observed peak-so-far memory, not an unmeasured final peak. Missing resource telemetry is disclosed rather than reconstructed.

The original audited source-only report is preserved byte-for-byte in [its snapshot](../snapshots/final_v14_source_working_pre_completion_20260905.identity.json). The complete report extends it; original Module-I plots and the montage atlas are copied unchanged. A separate final source manifest and bundle audit bind the final rendered artifacts. No operational pending receipt is silently recoded as a scientific negative result.
"""


def setup_text(data: dict) -> str:
    return f"""# FINAL-v14 · Experimental setup and interpretation scope

This report implements [the preregistered Aim-4 extension](../final_v14_PREREGISTRATION.md) of immutable FINAL-v13. The source population has 1,239 TCGA/SurGen primary patients, 501 mutant and 738 wild type, and 1,389 slides. UNI-v1 features, five frozen folds and five corresponding teacher seeds are inherited. No new MIL model is fitted.

Six PCA bases and fourteen k-means fits comprise the canonical source reference, five outer-training vocabularies and nine full-source grid members in total. Every OOF patient is assigned only through its outer-training vocabulary. Feature fitting is source-only; reference correspondence supports semantic summaries without changing ALL32 coordinates or formal gates. Source patient abundance averages slides equally.

The accepted single mentor return yields NAME_GATE_FAIL. Raw workbook/montage bytes and literal responses are immutable; no missing field is imputed. Thus ALL32 is the modeled representation, and all semantic associations use prototype codes. The original 60-case/420-image whole-section packet remains **GENERATED-UNREAD**.

| Module | Model and population | Uncertainty and formal decisions |
|---|---|---|
| I | 25 teacher-to-concept ridge mappings, five direct classifiers, five joint clinical/concept models and five clinical comparators; source patient-OOF predictions | 10,000 fixed-prediction paired patient draws; source×KRAS strata, with source×site×KRAS for standardized restriction. H1.1/H1.3 use adjusted one-sided 97.5% bounds. |
| Context | All 25 reference prototypes matched in every fold; arcsine-square-root abundance/attention; HC3 OLS with source fixed effects except the source panel | BH separately within each covariate/representation family of 25; 2,000 patient bootstrap OLS refits; 275 attempted tests, 274 estimable. |
| II | One full-source logistic ALL32 refit, C=0.01 (ordered median of five selected source Cs); five archived target cohorts; equal-slide profiles | 10,000 within-cohort KRAS paired patient draws; fixed ½ macro. Both primary cohorts must pass the adjusted AUROC lower-bound rule. All target scores/OOD sealed before outcome join. |
| III | 125 source OOF concept classifiers: five fine tasks, five canonical controls, fifteen repeated controls, each across five folds | 20,000 partially paired draws and task-wise adjusted 99% bounds; canonical controls use a separate descriptive 10,000-draw bootstrap. Four distinct task-status families. |
| Offset grid | Complete label-free source eligibility before seed-20260819 selection of 25 patients per source subcohort; all listed slides | One diagonal half-stride shift; exact-MPP/HEST and UNI preserved; paired patient cosine and ICC(2,1), undefined with support when either grid has zero between-patient variance. |

Module-II target populations are CPTAC primary 94 patients/98 slides (33 mutant), RIH primary 153/155 (70), Orion 40/41 (15), RIH metastatic 85/85 (37), and SR1482 metastatic 74/100 (30). RIH tissue-role interaction excludes eight dual-role patients from both arms; SR1482 primary/metastatic are disjoint but source-family exposed. Orion is descriptive only. These cohorts were opened in FINAL-v13: this is locked secondary analysis, never fresh external validation.

Module-III definitions are codon: G12 versus other KRAS mutants; broad G12D: G12D versus non-G12D among mutants; allele 1: G12D versus other G12; allele 2: G12V versus other G12; and G12C: G12C versus other G12. Each fine task retains its frozen positive/negative patient roster and four matched-WT control rosters.

All score intervals condition on frozen folds, features and fitted algorithms; they do not include model refitting uncertainty. Slides, seeds and folds are not treated as independent observations. Source-distance diagnostics cannot remove a patient or adapt the target model. Tissue-role coefficient sums decompose a linear mean-logit contrast only. Correspondence geometry is descriptive and cannot rescue failed naming.

Claim ceilings, verbatim:

> {CEILINGS}

Multireader naming, causal interventions, counterfactual images, target reweighting/refits, clinical thresholds/calibration/utility, alternate encoders, scales, cellular/spatial models and changes to Aim 1–3 verdicts remain out of scope. [Deviations and execution history](Deviations.md) distinguish scientific methods from operational delays. The final bundle is governed by the separate audit and source manifest; report generation alone does not seal a completed campaign.
"""


def workbook(tables: dict[str, pd.DataFrame], output: Path, complete: bool) -> None:
    book = load_workbook(SNAPSHOT / "Aim4_Results.xlsx")
    for title, frame in tables.items():
        if title in book:
            del book[title]
        sheet = book.create_sheet(title[:31])
        sheet.append(frame.columns.tolist())
        for row in frame.itertuples(index=False, name=None):
            values = []
            for value in row:
                if isinstance(value, (dict, list, tuple)):
                    value = json.dumps(value, sort_keys=True)
                elif isinstance(value, np.generic):
                    value = value.item()
                if isinstance(value, float) and not np.isfinite(value):
                    value = None
                values.append(value)
            sheet.append(values)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="244D42")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for column in sheet.columns:
            width = min(62, max(14, max(len(str(c.value or "")) for c in list(column)[:120]) + 2))
            sheet.column_dimensions[column[0].column_letter].width = width
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if cell.data_type == "n":
                    cell.number_format = "0.000000"
                else:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
    book.save(output / "Aim4_Results.xlsx")


def operating_envelope(data: dict) -> list[dict]:
    return [
        {"Module": "I", "Completion": "Complete", "Supported observation": "Above-chance ALL32 OOF discrimination; partial score fidelity", "Limit": "Restriction persistence not established; observed gaps not captured by this k=32 abundance representation"},
        {"Module": "II", "Completion": "Complete" if data["m2"] else "Pending", "Supported observation": "Archived cohort estimates and paired gaps" if data["m2"] else "Pending", "Limit": "Two-primary-cohort discrimination not established; no fresh external validation" if data["m2"] else "No verdict"},
        {"Module": "III", "Completion": "Complete" if data["m3"] else "Pending", "Supported observation": "Separate task-wise control, boundary and ranking results" if data["m3"] else "Pending", "Limit": "No supported fine-signal/superiority; some controls inadequate; no information-absence conclusion" if data["m3"] else "No verdict"},
    ]


def assemble_tables(data: dict) -> dict[str, pd.DataFrame]:
    """Deterministic scientific-table assembly for the independent bundle verifier."""
    old = load_workbook(SNAPSHOT / "Aim4_Results.xlsx", read_only=True, data_only=False)
    tables = {}
    for sheet in old:
        if sheet.title != "Status and scope":
            rows = list(sheet.values)
            tables[sheet.title] = pd.DataFrame(rows[1:], columns=rows[0])
    old.close()
    tables.update({**module1_tables(data), **module2_tables(data["m2"]), **module3_tables(data["m3"])})
    tables["Correspondence all variants"] = pd.DataFrame([{**{k: r[k] for k in ("k", "seed", "canonical", "joint_cosine_objective")},
                                                           **{axis + "_" + key: v for axis in ("p17", "p28") for key, v in r[axis].items()}} for r in data["correspondence"]["variants"]])
    if data["grid"] and data["grid"]["status"] == "GRID_OFFSET_SENSITIVITY_COMPLETE":
        tables["Grid patient cosine"] = pd.DataFrame(data["grid"]["patient_cosine"])
        tables["Grid prototype ICC"] = pd.DataFrame(data["grid"]["prototype_icc"])
    elif data["grid"]:
        tables["Grid eligible counts"] = pd.DataFrame([{"subcohort": k, "eligible_patients": v} for k, v in data["grid"]["eligible_counts"].items()])
    tables["Canonical variant summary"] = pd.DataFrame(data["variant_mapping"]["variants"])
    tables["Canonical variant mappings"] = pd.read_csv(RUN / "e4v_variant_mapping/canonical_to_variant.csv", float_precision="round_trip")
    tables["Operating envelope"] = pd.DataFrame(operating_envelope(data))
    complete = all(data.get(k) is not None for k in ("m2", "m3", "grid"))
    tables["Status and scope"] = pd.DataFrame([{"item": "analysis_complete", "value": complete}, {"item": "bundle", "value": "Requires separate final audit and seal"},
                                             {"item": "name_gate", "value": "NAME_GATE_FAIL"}, {"item": "whole-section packet", "value": "GENERATED-UNREAD"}, {"item": "claim ceilings", "value": CEILINGS}])
    return tables


def audit_tables() -> dict[str, list]:
    """Replay every CSV/worksheet from sealed scientific inputs and the snapshot."""
    data, _ = load_inputs()
    return {name: [frame.columns.tolist()] + frame.astype(object).where(frame.notna(), None).values.tolist()
            for name, frame in assemble_tables(data).items()}


def markdown_bindings(data: dict, tables: dict) -> dict[str, str]:
    blocks = {"Operating envelope": md_table(operating_envelope(data)),
              "M2 performance": module2_text(data["m2"], tables),
              "M3 task summary": module3_text(data["m3"], tables),
              "Correspondence all variants": correspondence_text(data["correspondence"]),
              "Canonical variant summary": variant_mapping_text(data["variant_mapping"])}
    if data["grid"]:
        title = "Grid prototype ICC" if data["grid"]["status"] == "GRID_OFFSET_SENSITIVITY_COMPLETE" else "Grid eligible counts"
        blocks[title] = grid_text(data["grid"])
    return blocks


def audit_markdown_bindings() -> dict[str, str]:
    """Reconstruct bound report blocks independently of saved STATUS/Markdown."""
    data, _ = load_inputs()
    return markdown_bindings(data, assemble_tables(data))


def render(output: Path) -> dict:
    for name in ("source_manifest.json", "final_bundle_receipt.json", "final_bundle_receipt.json.seal.json"):
        if (output / name).exists():
            raise ValueError(f"Refusing to mutate a completed or partially sealed report: {output / name}")
    data, input_pins = load_inputs()
    if output.resolve() == SNAPSHOT.resolve() or output.is_relative_to(RUN / "e4v_pre_reader"):
        raise ValueError("Cannot render into frozen input directories")
    output.mkdir(parents=True, exist_ok=True)
    (output / "figures").mkdir(exist_ok=True)
    (output / "tables").mkdir(exist_ok=True)
    for path in [SNAPSHOT / "mentor_atlas.html", *sorted((SNAPSHOT / "figures").glob("*"))]:
        target = output / path.relative_to(SNAPSHOT)
        shutil.copyfile(path, target)
    tables = assemble_tables(data)
    complete = all(data.get(k) is not None for k in ("m2", "m3", "grid"))
    envelope = operating_envelope(data)
    blocks = markdown_bindings(data, tables)
    for title, frame in tables.items():
        frame.to_csv(output / "tables" / (title.replace(" ", "_") + ".csv"), index=False, float_format="%.17g")
    workbook(tables, output, complete)
    figures(data, output)
    original = (SNAPSHOT / "Results.md").read_text()
    body = original[original.index("## Interpretation of Aim 1:"):original.index("## Interpretation of Aim 2:")]
    body = body.replace("## Interpretation of Aim 1: concept fidelity and discrimination", "## Interpretation of Aim 1: concept fidelity and concept-only discrimination")
    body = body.replace("These v14 identifiers do not denote the legacy\np17/p28 hypotheses.", "These are v14 coordinate IDs; geometry alone does not establish a named-morphology association.")
    state = "All in-scope analyses have final scientific outputs. This rendered report awaits the separate final audit and bundle seal." if complete else "Working report: Modules I–III and correspondence are complete; the raw-slide grid sensitivity remains pending. No completed FINAL-v14 bundle is declared."
    report = f"""# FINAL-v14 · Aim 4 analysis

**{state}**

The ALL32 concept representation discriminates KRAS above chance in source OOF patients, with limited reconstruction of the full-model score. It does not establish the required two-primary-cohort discrimination gate or a supported fine-task ranking advantage. These results bound the specified representation and readout; all [FINAL-v13](../final_v13/Results.md) Aim 1–3 estimates and verdicts remain unchanged.

[Results workbook](Aim4_Results.xlsx) · [Experimental setup](Experimental_Setup.md) · [Deviations and provenance](Deviations.md) · [Preregistration](../final_v14_PREREGISTRATION.md)

{body}
All pooled and per-source-subcohort fidelity estimates, including Spearman correlation, calibration intercept/slope, untruncated ratios and reconstruction gaps, are retained in [the full Module-I table](tables/M1_all_estimands.csv) and workbook. Stage-known and standardized-restriction rows retain their original audited sheets and source results.

{module2_text(data['m2'], tables)}
{module3_text(data['m3'], tables)}
{correspondence_text(data['correspondence'])}
{variant_mapping_text(data['variant_mapping'])}
{grid_text(data['grid'])}
## Operating envelope and paper integration

{md_table(envelope)}

![One-way interpretive links from Aim 4 to Aims 1–3](figures/figure1c_interpretive_links.png)

Figure 1C line style records completion of the corresponding sealed module, never feedback or independent validation. A completed module can yield a negative, inadequate-control or not-evaluable scientific call. The arrows do not modify an Aim 1–3 verdict.

Claim ceilings, verbatim:

> {CEILINGS}

The FINAL-v13 60-case/420-image whole-section packet remains **GENERATED-UNREAD**. No whole-section interpretation, multireader claim, clinical utility, deployment threshold or target-adaptive analysis is added. Publication plots include [Module-II PDF](figures/aim4_module2.pdf), [Module-III SVG](figures/aim4_module3.svg), and [Figure-1C PDF](figures/figure1c_interpretive_links.pdf). All numeric rows trace to sealed artifacts under `final_v14_additions_20260903`; the separate final source manifest and audit receipt control bundle completion.

## Complete tabular supplements

{md_table([{'Table': title, 'CSV': '[Full table](tables/' + title.replace(' ', '_') + '.csv)'} for title in tables])}
"""
    (output / "Results.md").write_text(report)
    (output / "Experimental_Setup.md").write_text(setup_text(data))
    (output / "Deviations.md").write_text(deviation_text())
    # Bind only this generator's outputs; the final manifest/audit/receipt are separate.
    report_files = [output / n for n in ("Results.md", "Experimental_Setup.md", "Deviations.md", "Aim4_Results.xlsx", "mentor_atlas.html")]
    report_files += sorted((output / "figures").glob("*")) + sorted((output / "tables").glob("*"))
    status = {"status": "REPORT_READY_FOR_FINAL_AUDIT" if complete else "WORKING_FINAL_V14_AIM4_INCOMPLETE",
              "created_utc": dt.datetime.now(dt.timezone.utc).isoformat(), "analysis_complete": complete, "completed_bundle": False,
              "completed": ["Module I", "context and attention", "reader intake and naming", "legacy correspondence", "canonical dictionary variant stability"] + (["Module II"] if data["m2"] else []) + (["Module III"] if data["m3"] else []) + (["grid final scientific output"] if data["grid"] else []),
              "pending": [k for k in ("m2", "m3", "grid") if data[k] is None] + ["final audit and bundle seal"],
              "code": identity(Path(__file__)), "inputs": input_pins,
              "scientific_inputs": input_pins[1:],
              "table_bindings": [{"csv": "tables/" + title.replace(" ", "_") + ".csv", "sheet": title[:31], "markdown": "Results.md", "markdown_role": "summary_or_linked_complete_table",
                                  **({"markdown_table": blocks[title]} if title in blocks else {})} for title in tables],
              "report_artifacts": [{**identity(p), "path": str(p.relative_to(output))} for p in report_files],
              "source_working_snapshot": identity(SNAPSHOT.with_name(SNAPSHOT.name + ".identity.json")),
              "scope": "Report generation does not issue the final bundle seal; any pending in-scope analysis prevents completion."}
    (output / "STATUS.json").write_text(json.dumps(status, indent=2, sort_keys=True) + "\n")
    return status


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=OUT)
    args = parser.parse_args()
    status = render(args.output.resolve())
    print(json.dumps({"status": status["status"], "analysis_complete": status["analysis_complete"], "report_artifacts": len(status["report_artifacts"]), "output": str(args.output.resolve())}, indent=2))


if __name__ == "__main__":
    main()
