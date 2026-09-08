#!/usr/bin/env python3
"""Verify FINAL-v14 identity closure; publish its final receipt only on request.

``audit`` is read-only. ``seal --authorize-final`` publishes the manifest, audit,
and finally the completion receipt, accepting only byte-identical replays. Large
packed inputs reuse only the exact independently deep-hashed pre-reader pins;
their size/mtime guards are checked and they are never described as rehashed.
"""
from __future__ import annotations

import argparse
import csv
import html
import importlib
import json
import math
import re
import sys
from collections.abc import Mapping
from pathlib import Path
from typing import Any
from urllib.parse import unquote, urlsplit

from openpyxl import load_workbook

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO))
from tools.final_v14_module3 import identity, json_bytes, read_json, write_once  # noqa: E402

ROOT = REPO / "reports/reruns/final_v14_additions_20260903"
REPORT = REPO / "reports/final_v14"
PARENT_RECEIPT = REPO / "reports/snapshots/final_v13_pre_v14_20260903.identity.json"
WORKING_RECEIPT = REPO / "reports/snapshots/final_v14_source_working_pre_completion_20260905.identity.json"
FINAL_NAMES = {"source_manifest.json", "Audit.md", "final_bundle_receipt.json",
               "final_bundle_receipt.json.seal.json"}
SHA = re.compile(r"^[0-9a-f]{64}$")
# Fixed by the independent pre-reader auditor, whose source identity is itself
# bound in its completed audit receipt. This closes reuse to the audited bytes.
PRE_READER_PREFLIGHT_SHA256 = "5db760b8ff6bb6359dd86eb7062044b6948eb98a7640326dd3dd698ffe348263"
SOURCE_SUBCOHORTS = {"SR1482", "SR386", "TCGA-COAD", "TCGA-READ"}


class AuditError(RuntimeError):
    """Missing, inconsistent, or incomplete required campaign evidence."""


def require(condition: bool, message: str) -> None:
    if not condition:
        raise AuditError(message)


def records(value: Any):
    """Walk identity objects; ordinary strings and numeric results are not paths."""
    if isinstance(value, dict):
        if isinstance(value.get("path"), str) and SHA.fullmatch(str(value.get("sha256", ""))):
            yield value
        for child in value.values():
            yield from records(child)
    elif isinstance(value, list):
        for child in value:
            yield from records(child)


def verify_identity(path: Path, pin: Mapping[str, Any]) -> dict:
    actual = identity(path)
    require(actual["sha256"] == pin["sha256"], f"Hash mismatch: {path}")
    require("size_bytes" not in pin or actual["size_bytes"] == pin["size_bytes"], f"Size mismatch: {path}")
    return actual


def verify_seal(path: Path) -> dict:
    """Accept the three historical seal wrappers, never an unbound timestamp."""
    seal_path = Path(str(path) + ".seal.json")
    seal = read_json(seal_path)
    pin = seal.get("artifact", seal)
    require(SHA.fullmatch(str(pin.get("sha256", ""))) is not None, f"Unbound seal: {seal_path}")
    if "sha256" in seal:
        require(seal["sha256"] == pin["sha256"], f"Conflicting seal hashes: {seal_path}")
    verify_identity(path, pin)
    return identity(seal_path)


def parent_snapshot(receipt_path: Path = PARENT_RECEIPT, repo: Path = REPO) -> list[dict]:
    receipt = read_json(receipt_path)
    require(receipt.get("diff_clean") is True, "Original v13 snapshot was not diff-clean")
    expected = receipt["files"]
    require(len(expected) == 6, "Expected the six immutable FINAL-v13 files")
    checked = [identity(receipt_path)]
    for key in ("source", "snapshot"):
        base = repo / receipt[key]
        actual = {str(p.relative_to(base)) for p in base.rglob("*") if p.is_file()}
        require(actual == set(expected), f"FINAL-v13 {key} inventory drift")
        checked.extend(verify_identity(base / name, pin) for name, pin in sorted(expected.items()))
    return checked


def working_snapshot(receipt_path: Path = WORKING_RECEIPT) -> tuple[list[dict], dict]:
    verify_seal(receipt_path)
    receipt = read_json(receipt_path)
    require(receipt.get("status") == "AUDITED_WORKING_REPORT_SNAPSHOT_PRESERVED", "Working report snapshot is not preserved")
    checked = [identity(receipt_path), verify_seal(receipt_path)]
    checked.append(verify_identity(Path(receipt["original_audit"]["path"]), receipt["original_audit"]))
    mapping = {}
    for item in receipt["report_path_mapping"]:
        original, snapshot = item["original"], item["snapshot"]
        require(original["sha256"] == snapshot["sha256"] and original["size_bytes"] == snapshot["size_bytes"], "Snapshot mapping changed bytes")
        checked.append(verify_identity(Path(snapshot["path"]), snapshot))
        key = (str(Path(original["path"]).resolve()), original["sha256"])
        require(key not in mapping, "Duplicate working-report snapshot mapping")
        mapping[key] = Path(snapshot["path"])
    require(len(mapping) >= 10, "Working report snapshot inventory incomplete")
    return checked, mapping


def historical_code_mapping(root: Path) -> dict:
    """Resolve the explicit pre-association repair's superseded source bytes."""
    path = root / "e4m1_context/analysis_contract_v2.json"
    verify_seal(path)
    amendment = read_json(path)["pre_association_numerical_amendment"]
    archive = amendment["archived_initial_runner"]
    verify_identity(Path(archive["path"]), archive)
    prior = amendment["prior_contract"]
    verify_identity(Path(prior["path"]), prior)
    prior_code = read_json(Path(prior["path"]))["code_identities"]
    matches = [pin for pin in prior_code if pin["sha256"] == archive["sha256"]]
    require(len(matches) == 1, "Numerical amendment does not bind its superseded code")
    return {(str(Path(matches[0]["path"]).resolve()), archive["sha256"]): Path(archive["path"])}


class IdentityClosure:
    """Check reachable current pins; preserve obsolete receipts as hashed leaves."""
    def __init__(self, campaign: Path, *, path_map: dict | None = None,
                 inherited: dict | None = None):
        self.campaign = campaign.resolve()
        self.path_map = path_map or {}
        self.inherited = inherited or {}
        self.checked: dict[tuple[str, str], dict] = {}
        self.traversed: set[tuple[str, str]] = set()
        self.unresolved_relative: list[dict] = []

    def add(self, pin: dict, *, origin: Path | None = None, recurse: bool = True) -> dict:
        path = Path(pin["path"])
        if not path.is_absolute():
            choices = [REPO / path]
            if origin:
                choices.insert(0, origin.parent / path)
                if origin.name == "source_fit_summary.json":
                    choices.append(origin.parent / "fits" / path)
                # Historical public-package inventories explicitly use paths
                # relative to the package, although their receipt lives outside it.
                if origin.is_relative_to(self.campaign / "e4v_pre_reader"):
                    choices += [self.campaign / "e4v_pre_reader/reader_package/public" / path,
                                REPO / "reviews/v14/FOR_PATHOLOGIST" / path]
            matches = [p for p in choices if p.is_file()]
            require(bool(matches), f"Unresolved relative identity {path} in {origin}")
            path = matches[0]
        path = path.resolve()
        logical = str(path)
        key = (logical, pin["sha256"])
        path = self.path_map.get(key, path)
        if key not in self.checked:
            require(path.is_file(), f"Missing bound artifact: {path}")
            inherited = self.inherited.get(key)
            if inherited:
                stat = path.stat()
                require(stat.st_size == inherited["size_bytes"] and stat.st_mtime_ns == inherited["mtime_ns"], f"Inherited packed-input stat guard changed: {path}")
                actual = {"path": str(path), "sha256": pin["sha256"], "size_bytes": stat.st_size}
                method = "INHERITED_PRE_READER_DEEP_HASH_WITH_CURRENT_SIZE_MTIME_GUARD"
            else:
                actual = verify_identity(path, pin)
                method = "SHA256_REPLAYED_NOW"
            item = {**actual, "verification": method}
            if str(path) != logical:
                item["historical_path"] = logical
            if inherited:
                item["evidence"] = inherited["evidence"]
                item["mtime_ns"] = inherited["mtime_ns"]
            self.checked[key] = item
        if recurse and key not in self.traversed and path.suffix == ".json" and path.is_relative_to(self.campaign):
            self.traversed.add(key)
            value = read_json(path)
            historical_operational = any(part in {"dependency_preflight", "operational_preflight_20260905", "execution_handoff_20260905", "core_preflight_20260905"} for part in path.parts)
            if historical_operational:
                self.checked[key]["closure_role"] = "PRESERVED_HISTORICAL_OPERATIONAL_RECEIPT_LEAF"
            else:
                for child in records(value):
                    self.add(child, origin=path)
            seal_path = Path(str(path) + ".seal.json")
            if seal_path.exists():
                self.add(verify_seal(path), recurse=False)
        return self.checked[key]

    def file(self, path: Path, *, sealed: bool = False, recurse: bool = True) -> None:
        self.add(identity(path), recurse=recurse)
        if sealed:
            self.add(verify_seal(path), recurse=False)

    def inventory(self) -> list[dict]:
        return sorted(self.checked.values(), key=lambda r: (r["path"], r["sha256"]))


def inherited_pack_pins(root: Path) -> dict:
    audit_path = root / "e4v_pre_reader_independent_audit/audit_receipt.json"
    audit = read_json(audit_path)
    require(audit["status"] == "PASS" and audit["gates"]["core_identities"]["status"] == "PASS", "Pre-reader deep identity audit did not pass")
    verify_identity(Path(audit["auditor"]["path"]), audit["auditor"])
    preflight_path = root / "e4v_pre_reader/receipts/preflight.json"
    require(identity(preflight_path)["sha256"] == PRE_READER_PREFLIGHT_SHA256, "Preflight bytes differ from the independent auditor's fixed pin")
    preflight = read_json(preflight_path)
    require(preflight["deep_hash"] is True, "Pre-reader pack was not deep-hashed")
    evidence = [identity(audit_path), identity(preflight_path), identity(Path(audit["auditor"]["path"]))]
    result = {}
    for item in preflight["pack"]["artifacts"].values():
        require(item.get("hash_verified") is True and item["sha256"] == item["expected_sha256"], "Unverified pre-reader packed pin")
        if item["size_bytes"] > 64 * 1024**2:
            result[(str(Path(item["path"]).resolve()), item["sha256"])] = {**item, "evidence": evidence}
    return result


def check_states(data: dict[str, dict]) -> None:
    require(data["report"].get("status") == "REPORT_READY_FOR_FINAL_AUDIT" and data["report"].get("analysis_complete") is True,
            "Report does not declare completed analyses ready for final audit")
    require(data["report"].get("completed_bundle") is False, "Report cannot pre-empt the final receipt")
    require(data["m1_source"]["status"] == "SOURCE_MODELS_SEALED", "Module I source incomplete")
    require(bool(data["m1"].get("gates")) and "ALL32" in data["m1"], "Module I results incomplete")
    m1a = data["m1_audit"]
    require(m1a["status"] == "PASS_SOURCE_COMPLETED_CAMPAIGN_INCOMPLETE", "Wrong Module I audit scope")
    for key in ("source_models", "source_inference", "context"):
        require(m1a[key]["status"] == "PASS", f"Module I {key} independent audit failed")
    require(data["context"]["status"] == "CONTEXT_ASSOCIATIONS_COMPLETE", "Module I context incomplete")
    require(data["m2"]["status"] == "MODULE_II_ANALYSIS_COMPLETE", "Module II incomplete")
    require(data["m2_audit"]["status"] == "INDEPENDENT_MODULE_II_NUMERICAL_AUDIT_PASS", "Module II numerical audit missing")
    require(data["m3"]["status"] == "COMPLETE" and data["m3_source"]["status"] == "SOURCE_FITS_COMPLETE", "Module III incomplete")
    require(data["m3_audit"]["status"] == "PASS", "Module III independent audit failed")
    require(data["correspondence"]["status"] == "CORRESPONDENCE_GEOMETRY_COMPLETE", "Correspondence geometry incomplete")
    require(data["correspondence_audit"]["status"] == "PASS_INDEPENDENT_CORRESPONDENCE_REPLAY", "Correspondence independent audit missing")
    require(data["variant_mapping"]["status"] == "CANONICAL_NINE_VARIANT_MAPPING_COMPLETE", "Canonical-to-variant sensitivity incomplete")
    require(data["variant_audit"]["status"] == "INDEPENDENT_NINE_VARIANT_MAPPING_AUDIT_PASS", "Canonical-to-variant independent audit missing")
    grid, audit = data["grid"], data["grid_audit"]
    require(audit.get("status") == "GRID_OFFSET_INDEPENDENT_AUDIT_PASS", "Grid independent audit missing")
    if grid["status"] == "GRID_OFFSET_NOT_EVALUABLE":
        require(audit.get("prespecified_insufficient_eligible_verified") is True and audit.get("operational_missing_dependencies") is False,
                "Operational grid block cannot become scientific non-evaluability")
        require(grid.get("patient_roster_drawn") is False and grid.get("slide_extraction_run") is False, "Invalid insufficient-eligibility fallback")
        counts = grid["eligible_counts"]
        require(set(counts).issubset(SOURCE_SUBCOHORTS), "Grid eligibility uses an unknown source stratum")
        require(any(counts.get(group, 0) < 25 for group in SOURCE_SUBCOHORTS), "Grid insufficiency has no underfilled source stratum")
    else:
        require(grid["status"] == "GRID_OFFSET_SENSITIVITY_COMPLETE", "Actual grid sensitivity is incomplete")


def table_equal(left: Any, right: Any) -> bool:
    if left is None or left == "":
        return right is None or right == ""
    if isinstance(left, bool) or isinstance(right, bool):
        return str(left).lower() == str(right).lower()
    try:
        a, b = float(left), float(right)
        return math.isfinite(a) and math.isfinite(b) and math.isclose(a, b, rel_tol=1e-12, abs_tol=1e-12)
    except (ValueError, TypeError):
        return str(left) == str(right)


def compare_rows(expected: list[list], actual: list[list], label: str) -> None:
    require(len(expected) == len(actual), f"Table row count differs: {label}")
    for i, (a, b) in enumerate(zip(expected, actual, strict=True)):
        require(len(a) == len(b), f"Table column count differs: {label} row {i}")
        for j, (x, y) in enumerate(zip(a, b, strict=True)):
            require(table_equal(x, y), f"Table mismatch: {label} row {i} column {j}: {x!r} != {y!r}")


def report_tables(report: Path, status: dict, *, expected_tables: dict | None = None,
                  expected_markdown: dict | None = None) -> dict:
    bindings = status.get("table_bindings", [])
    require(bool(bindings), "Report lacks table bindings")
    if expected_tables is None:
        module = importlib.import_module("tools.final_v14_complete_report")
        require(hasattr(module, "audit_tables"), "Report generator needs a pure audit_tables() interface")
        expected_tables = module.audit_tables()
        require(hasattr(module, "audit_markdown_bindings"), "Report generator needs regenerated Markdown bindings")
        expected_markdown = module.audit_markdown_bindings()
    expected_markdown = expected_markdown or {}
    workbook = load_workbook(report / "Aim4_Results.xlsx", read_only=True, data_only=False)
    sheets = set()
    counts = {}
    try:
        for binding in bindings:
            sheet = binding["sheet"]
            require(sheet in workbook.sheetnames and sheet not in sheets, f"Missing/duplicate bound worksheet: {sheet}")
            sheets.add(sheet)
            with (report / binding["csv"]).open(newline="") as handle:
                csv_rows = list(csv.reader(handle))
            require(sheet in expected_tables, f"No scientific replay table: {sheet}")
            expected = expected_tables[sheet]
            if hasattr(expected, "columns"):
                expected = [list(expected.columns)] + expected.where(expected.notna(), None).values.tolist()
            compare_rows(expected, csv_rows, f"scientific inputs → {binding['csv']}")
            actual = [list(row) for row in workbook[sheet].iter_rows(values_only=True)]
            compare_rows(csv_rows, actual, f"CSV → workbook/{sheet}")
            if "markdown_table" in binding:
                require(expected_markdown.get(sheet) == binding["markdown_table"], f"Scientific inputs → Markdown binding drift: {sheet}")
                require(binding["markdown_table"] in (report / binding.get("markdown", "Results.md")).read_text(), f"Markdown table drift: {sheet}")
            counts[sheet] = len(csv_rows) - 1
        require(sheets == set(workbook.sheetnames) == set(expected_tables), "Workbook contains unbound/omitted scientific worksheets")
        require({binding["sheet"] for binding in bindings if "markdown_table" in binding} == set(expected_markdown), "Report omitted a required regenerated Markdown block")
    finally:
        workbook.close()
    return {"status": "PASS", "table_rows": counts, "comparison": "Scientific inputs to CSV to workbook; bound Markdown tables checked verbatim"}


def linked_files(path: Path) -> list[Path]:
    text = path.read_text()
    targets = re.findall(r"\]\(([^)]+)\)", text) if path.suffix == ".md" else []
    targets += re.findall(r'(?:src|href)=[\"\']([^\"\']+)[\"\']', text)
    result = []
    for target in targets:
        target = html.unescape(target).strip().strip("<>")
        url = urlsplit(target)
        if url.scheme or not url.path:
            continue
        resolved = (path.parent / unquote(url.path)).resolve()
        require(resolved.is_file(), f"Broken report artifact link: {path} → {target}")
        result.append(resolved)
    return result


def stage_paths(root: Path, report: Path, grid_audit: Path) -> dict[str, Path]:
    correspondence = list((root / "e4v_correspondence_replay/runs").glob("*/results.json"))
    require(len(correspondence) == 1, "Need one governed correspondence result")
    return {"report": report / "STATUS.json", "m1": root / "e4m1_results/results.json",
            "m1_source": root / "e4m1_source/source_fit_summary.json",
            "m1_audit": root / "post_reader_audit/source_audit_receipt.json",
            "context": root / "e4m1_context/results.json",
            "m2": root / "e4m2_source_frozen/target_results.json",
            "m2_audit": root / "e4m2_independent_audit/numerical_audit_receipt.json",
            "m3": root / "module_iii/results.json", "m3_source": root / "module_iii/source_completion.json",
            "m3_audit": root / "module_iii/audit/audit_completion.json",
            "correspondence": correspondence[0], "correspondence_audit": correspondence[0].with_name("independent_audit.json"),
            "variant_mapping": root / "e4v_variant_mapping/results.json",
            "variant_audit": root / "e4v_variant_mapping/independent_audit.json",
            "grid": root / "grid_offset/locked_extraction/results.json", "grid_audit": grid_audit}


def require_result_binding(value: dict, result_path: Path, label: str) -> None:
    expected = identity(result_path)
    require(any(pin["sha256"] == expected["sha256"] and Path(pin["path"]).resolve() == result_path.resolve()
                for pin in records(value)), f"{label} does not bind its completed numerical result")


def require_source_audit_bindings(m1_audit: dict, paths: dict, mapping: dict) -> None:
    # The original Module-I audit binds its results indirectly through the
    # audited working STATUS. Resolve only its explicitly preserved snapshot.
    report_pin = m1_audit["report"]["report_status"]
    key = (str(Path(report_pin["path"]).resolve()), report_pin["sha256"])
    require(key in mapping, "Module I audit lacks preserved report-status mapping")
    verify_identity(mapping[key], report_pin)
    old_status = read_json(mapping[key])
    require_result_binding(old_status["inputs"], paths["m1"], "Module I audited report status")
    require_result_binding(m1_audit["context"], paths["context"], "Module I context audit")


def audit(root: Path, report: Path, grid_audit: Path) -> dict:
    paths = stage_paths(root, report, grid_audit)
    missing = [str(path) for path in paths.values() if not path.is_file()]
    require(not missing, f"Required campaign evidence pending: {missing}")
    data = {key: read_json(path) for key, path in paths.items()}
    check_states(data)
    parents = parent_snapshot()
    preserved, mapping = working_snapshot()
    report_mapping_count = len(mapping)
    mapping.update(historical_code_mapping(root))
    require_source_audit_bindings(data["m1_audit"], paths, mapping)
    closure = IdentityClosure(root, path_map=mapping, inherited=inherited_pack_pins(root))
    for record in parents + preserved:
        closure.add(record, recurse=False)
    for key, path in paths.items():
        closure.file(path, sealed=key not in {"report", "m1_audit"})
    # These procedure/benchmark roots are not all linked by the result wrappers.
    # Binding them explicitly closes inference code, review provenance and the
    # operational disclosures without treating a benchmark as a scientific gate.
    additional = [
        ("e4v_pre_reader_independent_audit/audit_receipt.json", False),
        ("e4v_pre_reader/receipts/preflight.json", False),
        ("e4m1_results/evaluation_contract.json", True),
        ("e4m1_context/verification.json", True),
        ("e4m1_source_audits/exact_constant_detection_20260905/audit.json", True),
        ("e4v_post_reader_xlsx/naming/naming_freeze.json", True),
        ("e4m2_source_audit/source_audit_receipt.json", True),
        ("e4m2_source_frozen/execution_benchmark_receipt.json", True),
        ("module_iii/benchmarks/source_fit.json", True),
        ("module_iii/benchmarks/evaluation.json", True),
    ]
    if data["grid"]["status"] == "GRID_OFFSET_SENSITIVITY_COMPLETE":
        benchmark = root / "grid_offset/locked_extraction/benchmarks/final.json"
        require(benchmark.is_file(), "Final grid execution benchmark is pending")
        require(read_json(benchmark)["status"] == "GRID_OFFSET_FINAL_BENCHMARK_SEALED", "Final grid execution benchmark is incomplete")
        additional.append((str(benchmark.relative_to(root)), True))
    for relative, sealed in additional:
        closure.file(root / relative, sealed=sealed)
    # Numerical audits must name the exact result, not merely report a PASS word.
    for akey, rkey in (("m2_audit", "m2"), ("m3_audit", "m3"), ("correspondence_audit", "correspondence"), ("variant_audit", "variant_mapping"), ("grid_audit", "grid")):
        require_result_binding(data[akey], paths[rkey], akey)
    for pin in data["report"]["report_artifacts"]:
        path = Path(pin["path"])
        path = path if path.is_absolute() else report / path
        require(path.name not in FINAL_NAMES, "Report inventory creates a final-receipt cycle")
        closure.add({**pin, "path": str(path)}, recurse=False)
    declared_report_files = {str((Path(pin["path"]) if Path(pin["path"]).is_absolute() else report / pin["path"]).resolve())
                             for pin in data["report"]["report_artifacts"]}
    actual_report_files = {str(path.resolve()) for path in report.rglob("*") if path.is_file() and path.name not in FINAL_NAMES and path.name != "STATUS.json"}
    require(declared_report_files == actual_report_files, "Final report inventory contains omitted or unexpected artifacts")
    inputs = data["report"].get("scientific_inputs", data["report"].get("inputs", []))
    for key in ("m1", "context", "m2", "m3", "correspondence", "variant_mapping", "grid"):
        require(any(pin["sha256"] == identity(paths[key])["sha256"] for pin in records(inputs)), f"Report omitted current scientific input {key}")
    for pin in records(data["report"]):
        if Path(pin["path"]).is_absolute():
            closure.add(pin, recurse=False)
    tables = report_tables(report, data["report"])
    links = set()
    for path in report.glob("*"):
        if path.name not in FINAL_NAMES and path.suffix in {".md", ".html"}:
            links.update(linked_files(path))
    for path in sorted(links):
        closure.file(path, recurse=False)
    closure.file(Path(__file__), recurse=False)
    # Preserve historical operational receipts without treating their obsolete
    # incomplete statuses or formerly current report/code pins as active roots.
    historical = []
    for path in sorted(root.rglob("*.seal.json")):
        pin = identity(path)
        closure.add(pin, recurse=False)
        # Authenticate the old receipt itself as well as its preserved seal,
        # while leaving obsolete dependencies out of the active closure.
        original = Path(str(path).removesuffix(".seal.json"))
        require(original.is_file(), f"Orphaned historical seal: {path}")
        verify_seal(original)
        closure.file(original, recurse=False)
        historical.append(pin)
    return {"schema_version": 1, "status": "PASS_FINAL_CAMPAIGN_AUDIT",
            "scientific_stages": {key: identity(path) for key, path in paths.items()},
            "report_consistency": tables, "local_report_links_verified": len(links),
            "parent_snapshot_files_checked": 12, "working_report_path_mapping": report_mapping_count,
            "superseded_context_code_mapping": 1,
            "historical_seals_preserved": len(historical), "artifacts": closure.inventory(),
            "limitations": ["Existing independent numerical audits are authenticated; this packaging audit does not refit models or rerun every bootstrap.",
                            "Explicitly labeled packed inputs inherit the prior independent deep hash and receive current size/mtime guards, not a new content hash.",
                            "Historical no-target-access ordering is supported by frozen receipts/code, not an independent operating-system access log."]}


def publish(report: Path, result: dict) -> dict:
    require(result["status"] == "PASS_FINAL_CAMPAIGN_AUDIT", "Cannot publish incomplete campaign")
    manifest = report / "source_manifest.json"
    audit_path = report / "Audit.md"
    receipt_path = report / "final_bundle_receipt.json"
    write_once(manifest, json_bytes(result))
    lines = ["# FINAL-v14 campaign audit", "", "Status: **PASS**. All required scientific stages are complete or have an independently verified prespecified non-evaluability decision.", "",
             f"Manifest SHA-256: `{identity(manifest)['sha256']}`.", "",
             f"Verified {len(result['artifacts'])} artifact identities, {result['local_report_links_verified']} local report links, the six original and six snapshot FINAL-v13 files, and {result['working_report_path_mapping']} preserved working-report mappings.", "",
             "The manifest records each identity's verification method and the exact independent numerical audit receipts. CSV and workbook tables were replayed from scientific inputs; bound Markdown tables were checked. The final receipt is written after this audit and the manifest, avoiding a self-hash cycle.", "", "Limitations:", ""]
    lines += [f"- {item}" for item in result["limitations"]]
    write_once(audit_path, ("\n".join(lines) + "\n").encode())
    receipt = {"schema_version": 1, "status": "COMPLETED", "completed_bundle": True,
               "manifest": identity(manifest), "audit": identity(audit_path),
               "report_status": identity(report / "STATUS.json"), "verifier": identity(Path(__file__)),
               "completion_rule": "All mandatory scientific outputs and independent audits passed before this receipt was written; STATUS describes the preceding report-ready state."}
    write_once(receipt_path, json_bytes(receipt))
    write_once(Path(str(receipt_path) + ".seal.json"), json_bytes({"status": "SEALED", "artifact": identity(receipt_path)}))
    return receipt


def verify_bundle(report: Path) -> dict:
    receipt_path = report / "final_bundle_receipt.json"
    verify_seal(receipt_path)
    receipt = read_json(receipt_path)
    require(receipt.get("status") == "COMPLETED" and receipt.get("completed_bundle") is True, "Final receipt is not completed")
    for pin in records(receipt):
        verify_identity(Path(pin["path"]), pin)
    manifest = read_json(Path(receipt["manifest"]["path"]))
    require(manifest["status"] == "PASS_FINAL_CAMPAIGN_AUDIT", "Final manifest is not a passed audit")
    for pin in manifest["artifacts"]:
        path = Path(pin["path"])
        if pin["verification"] == "SHA256_REPLAYED_NOW":
            verify_identity(path, pin)
        else:
            require(pin["verification"] == "INHERITED_PRE_READER_DEEP_HASH_WITH_CURRENT_SIZE_MTIME_GUARD", "Unknown verification mode")
            for evidence in pin["evidence"]:
                verify_identity(Path(evidence["path"]), evidence)
            require(path.stat().st_size == pin["size_bytes"] and path.stat().st_mtime_ns == pin["mtime_ns"], "Inherited input stat drift")
    return {"status": "PASS", "receipt": identity(receipt_path), "artifacts": len(manifest["artifacts"])}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("stage", choices=["audit", "seal", "verify"])
    parser.add_argument("--root", type=Path, default=ROOT)
    parser.add_argument("--report", type=Path, default=REPORT)
    parser.add_argument("--grid-audit", type=Path, default=ROOT / "grid_offset/locked_extraction/independent_audit/results.json")
    parser.add_argument("--authorize-final", action="store_true")
    args = parser.parse_args()
    if args.stage == "verify":
        result = verify_bundle(args.report)
    else:
        if args.stage == "seal":
            require(args.authorize_final, "Final sealing requires explicit --authorize-final after coordinator approval")
        result = audit(args.root, args.report, args.grid_audit)
        if args.stage == "seal":
            result = publish(args.report, result)
    print(json.dumps({k: v for k, v in result.items() if k != "artifacts"}, indent=2))


if __name__ == "__main__":
    main()
