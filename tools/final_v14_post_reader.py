#!/usr/bin/env python3
"""Accept the authorized final v14 reader return, preserving missing responses.

This append-only stage requires the sealed coordinator amendment and raw return
before literal ingestion. It never changes the preregistration, reader response,
or old validators. Only sealed intake permits key access. No model is fitted.
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import hashlib
import io
import json
import os
import sys
from collections import Counter
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from scipy.stats import beta

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO))
sys.path.insert(0, str(REPO / "src"))
from tools import final_v14_excel_handoff as handoff  # noqa: E402
from tools import final_v14_pre_reader as pre  # noqa: E402

PRE = REPO / "reports/reruns/final_v14_additions_20260903/e4v_pre_reader"
POST = REPO / "reports/reruns/final_v14_additions_20260903/e4v_post_reader_xlsx"
COMPONENT = "final_v14_authorized_post_reader"


class PostReaderError(RuntimeError):
    """A required identity, order, or response-preservation contract failed."""


def now() -> str:
    return dt.datetime.now(dt.timezone.utc).isoformat()


def identity(path: Path) -> dict[str, Any]:
    if path.is_symlink() or not path.is_file():
        raise PostReaderError(f"Expected regular nonsymlink file: {path}")
    return {"path": str(path.resolve()), "sha256": pre.sha256_file(path),
            "size_bytes": path.stat().st_size}


def check_identity(record: dict[str, Any], path: Path | None = None) -> None:
    observed = identity(path or Path(record["path"]))
    for field in ("sha256", "size_bytes"):
        if observed[field] != record[field]:
            raise PostReaderError(f"Identity drift ({field}): {observed['path']}")


def read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def write_once(path: Path, payload: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        if path.read_bytes() != payload:
            raise PostReaderError(f"Refusing to replace a frozen artifact: {path}")
        return
    with path.open("xb") as stream:
        stream.write(payload)
        stream.flush()
        os.fsync(stream.fileno())
    path.chmod(0o400)


def json_bytes(value: Any) -> bytes:
    return (json.dumps(value, indent=2, sort_keys=True, ensure_ascii=False,
                       allow_nan=False) + "\n").encode("utf-8")


def publish(path: Path, value: Any) -> dict[str, Any]:
    write_once(path, json_bytes(value))
    record = identity(path)
    seal_path = path.with_suffix(path.suffix + ".seal.json")
    if seal_path.exists():
        seal = read_json(seal_path)
        check_identity(seal["artifact"], path)
        if seal.get("sha256") != record["sha256"]:
            raise PostReaderError(f"Seal digest disagreement: {path}")
    else:
        write_once(seal_path, json_bytes({"schema_version": 1, "status": "SEALED",
                   "created_utc": now(), "sha256": record["sha256"], "artifact": record}))
    return record


def verify_seal(path: Path) -> dict[str, Any]:
    seal = read_json(path.with_suffix(path.suffix + ".seal.json"))
    check_identity(seal["artifact"], path)
    if seal.get("status") != "SEALED" or seal.get("sha256") != identity(path)["sha256"]:
        raise PostReaderError(f"Invalid seal: {path}")
    return read_json(path)


def verify_authorization(post: Path = POST) -> dict[str, Any]:
    amendment = verify_seal(post / "governance_amendment.json")
    if amendment.get("status") != "USER_AUTHORIZED_FINAL_REVIEW_ACCEPTANCE_BEFORE_UNBLINDING":
        raise PostReaderError("Missing explicit final-review acceptance amendment")
    if amendment.get("key_access_before_this_amendment") is not False:
        raise PostReaderError("Amendment was not made before key access")
    for field in ("raw_return_receipt", "raw_return_workbook", "preregistration"):
        check_identity(amendment[field])
    receipt = read_json(post / "raw_return/receipt.json")
    if receipt["status"] != "RAW_RETURN_SEALED_BEFORE_VALIDATION_OR_KEY_ACCESS":
        raise PostReaderError("Raw workbook was not sealed before validation/key access")
    check_identity(receipt["artifact"], post / "raw_return" / receipt["artifact"]["path"])
    return amendment


def verify_frozen_delivery(post: Path = POST) -> dict[str, Any]:
    """Replay original identities; allow only the separately sealed return extra."""
    source = handoff.validate_source_package()
    receipt = read_json(handoff.HANDOFF_RECEIPT)
    handoff._validate_handoff_receipt(receipt["verification"])
    mirror = handoff.AUDIT_ROOT / "public"
    expected = handoff.tree_inventory(mirror)
    actual = handoff.tree_inventory(handoff.DELIVERY_PACKAGE)
    extra_name = handoff.COMPLETED_WORKBOOK_NAME
    extras = [item for item in actual if item["path"] == extra_name]
    actual_originals = [item for item in actual if item["path"] != extra_name]
    if actual_originals != expected:
        raise PostReaderError("Original delivered package differs from its sealed mirror")
    if extras:
        check_identity(identity(post / "raw_return" / extra_name),
                       handoff.DELIVERY_PACKAGE / extra_name)
    manifest = read_json(mirror / "HANDOFF_MANIFEST.json")
    for record in manifest["files"]:
        check_identity(record, mirror / record["path"])
    expected_names = handoff._package_expected_files(source["rows"])
    if {item["path"] for item in expected} != expected_names:
        raise PostReaderError("Sealed mirror contains an unexpected file inventory")
    if {p.name for p in handoff.DELIVERY_ROOT.iterdir()} != {
        "FOR_PATHOLOGIST", "README_COORDINATOR.md", "PRE_READER_STATUS.json"
    }:
        raise PostReaderError("Delivery root inventory changed")
    return {"status": "PASS", "original_files_verified": len(expected),
            "original_files_changed": 0, "accepted_sealed_return_extra": bool(extras),
            "handoff_receipt": identity(handoff.HANDOFF_RECEIPT),
            "source": source}


def literal_workbook_rows(payload: bytes, expected: list[dict[str, Any]],
                          rubric: str) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """Read responses literally; format guards are not evidence of completeness."""
    handoff._safe_zip_members(payload, "final authorized return")
    workbook = load_workbook(io.BytesIO(payload), data_only=False, keep_links=True)
    if workbook.sheetnames != handoff.EXPECTED_SHEETS:
        raise PostReaderError("Workbook sheet inventory/order changed")
    if getattr(workbook, "_external_links", []):
        raise PostReaderError("External workbook links are forbidden")
    handoff._reject_formulas(workbook, "final authorized return")
    handoff._reject_undeclared_workbook_content(workbook, "final authorized return")
    if workbook[handoff.INSTRUCTIONS_SHEET]["A2"].value != handoff.excel_instructions():
        raise PostReaderError("Reader instructions changed")
    if workbook[handoff.RUBRIC_SHEET]["A2"].value != rubric:
        raise PostReaderError("Reader rubric changed")
    blank = load_workbook(handoff.AUDIT_ROOT / "public" / handoff.WORKBOOK_NAME,
                          data_only=False)
    opts = workbook[handoff.OPTIONS_SHEET]
    base_opts = blank[handoff.OPTIONS_SHEET]
    if list(opts.values) != list(base_opts.values):
        raise PostReaderError("Literal ontology/options changed")
    # Excel may quote sheet names in defined references and rewrite conditional
    # formatting. These interface serializations never enter any response value.
    for name, definition in blank.defined_names.items():
        observed = workbook.defined_names.get(name)
        if observed is None or list(observed.destinations) != list(definition.destinations):
            raise PostReaderError(f"Defined option range changed: {name}")
    review = workbook[handoff.REVIEW_SHEET]
    if [review.cell(1, i).value for i in range(1, 14)] != handoff.FORM_COLUMNS:
        raise PostReaderError("Response schema changed")
    if len(expected) != 40:
        raise PostReaderError("Expected exactly 40 presentations")
    result = []
    for row_index, frozen in enumerate(expected, start=2):
        row = {field: review.cell(row_index, col).value
               for col, field in enumerate(handoff.FORM_COLUMNS, start=1)}
        fixed = [row[field] for field in handoff.FIXED_COLUMNS]
        if fixed != [int(frozen["presentation_order"]), frozen["blinded_code"],
                     int(frozen["n_tiles"])]:
            raise PostReaderError(f"Fixed presentation identity changed: row {row_index}")
        hyperlink = review.cell(row_index, 2).hyperlink
        if hyperlink is None or hyperlink.target != f"montages/{row['blinded_code']}.jpg":
            raise PostReaderError(f"Montage hyperlink changed: row {row_index}")
        for field in ("primary_category", "secondary_category_1", "secondary_category_2"):
            if row[field] is not None and row[field] not in handoff.ONTOLOGY:
                raise PostReaderError(f"Nonmissing category outside frozen ontology: {row_index}/{field}")
        if row["confidence_1_to_5"] is not None and row["confidence_1_to_5"] not in range(1, 6):
            raise PostReaderError(f"Invalid nonmissing confidence: {row_index}")
        if row["artifact_uninterpretable"] not in (None, "yes", "no"):
            raise PostReaderError(f"Invalid nonmissing artifact response: {row_index}")
        if isinstance(row["review_date"], (dt.date, dt.datetime)):
            row["review_date"] = row["review_date"].isoformat()
        result.append(row)
    missing = {field: sum(row[field] in (None, "") for row in result)
               for field in handoff.RESPONSE_COLUMNS}
    return result, {"rows": 40, "missing_response_counts": missing,
                    "literal_response_values_preserved": True,
                    "interface_normalization": "Parsed equivalent named-range references; ignored format-only serialization differences.",
                    "recorded_blinding_attestation_present": any(
                        row["blinding_attestation"] == "confirmed_no_key_access" for row in result)}


def intake(post: Path = POST) -> dict[str, Any]:
    verify_authorization(post)
    destination = post / "literal_intake/receipt.json"
    if destination.exists():
        result = verify_seal(destination)
        for record in result["artifacts"].values():
            check_identity(record)
        verify_frozen_delivery(post)
        return result
    delivery = verify_frozen_delivery(post)
    source = delivery.pop("source")
    raw = post / "raw_return" / handoff.COMPLETED_WORKBOOK_NAME
    rows, validation = literal_workbook_rows(raw.read_bytes(), source["rows"], source["rubric"])
    directory = destination.parent
    literal = directory / "literal_response_values.json"
    literal_record = publish(literal, {"rows": rows, "policy": "Missing cells remain null; text is unchanged."})
    buffer = io.StringIO(newline="")
    writer = csv.DictWriter(buffer, fieldnames=handoff.FORM_COLUMNS, lineterminator="\n")
    writer.writeheader()
    writer.writerows(rows)
    csv_path = directory / handoff.COMPLETED_CSV_NAME
    write_once(csv_path, buffer.getvalue().encode("utf-8"))
    feedback = directory / "reader_qualitative_feedback.json"
    feedback_record = publish(feedback, {"interpretation_scope": "Unchanged descriptive reader comments; cannot fill missing ontology, confidence, or artifact answers.",
        "comments": [{"presentation_order": row["presentation_order"], "blinded_code": row["blinded_code"],
                      "free_text_description": row["free_text_description"]}
                     for row in rows if row["free_text_description"] not in (None, "")]})
    result = {"schema_version": 1, "component": COMPONENT,
              "status": "AUTHORIZED_FINAL_RETURN_LITERAL_INTAKE_SEALED",
              "created_utc": now(), "validation": validation, "delivery_verification": delivery,
              "artifacts": {"raw_return": identity(raw), "raw_receipt": identity(post / "raw_return/receipt.json"),
                            "governance_amendment": identity(post / "governance_amendment.json"),
                            "literal_response_values": literal_record,
                            "canonical_completed_csv": identity(csv_path), "qualitative_feedback": feedback_record},
              "key_accessed_by_intake": False, "preregistered_complete_form_validation_passed": False}
    publish(destination, result)
    return result


def cp_interval(successes: int, total: int = 8) -> list[float]:
    if not 0 <= successes <= total or total <= 0:
        raise ValueError("Invalid binomial counts")
    return [0.0 if successes == 0 else float(beta.ppf(0.025, successes, total - successes + 1)),
            1.0 if successes == total else float(beta.ppf(0.975, successes + 1, total - successes))]


def category_agreement(a: dict[str, Any], b: dict[str, Any]) -> str:
    first, second = a["primary_category"], b["primary_category"]
    if not first or not second:
        return "missing_primary"
    if first == second:
        return "exact"
    if (first in (b["secondary_category_1"], b["secondary_category_2"])
            or second in (a["secondary_category_1"], a["secondary_category_2"])):
        return "partial"
    return "different"


def derive_naming(rows: list[dict[str, Any]], key: pd.DataFrame,
                  mapping: pd.DataFrame, duplicate_preflight: str) -> dict[str, Any]:
    by_code = {row["blinded_code"]: row for row in rows}
    if len(by_code) != 40 or len(key) != 40 or set(key["code"]) != set(by_code):
        raise PostReaderError("Reader/key code bijection failed")
    control = key.loc[key["is_controlling_read"].eq(True)]
    if sorted(control["prototype_id"].astype(int)) != list(range(32)):
        raise PostReaderError("Expected one predesignated controlling read per reference prototype")
    if not control["occurrence"].eq(0).all():
        raise PostReaderError("Controlling occurrences changed")
    if sorted(key["presentation_order"].astype(int)) != list(range(40)):
        raise PostReaderError("Embargoed key must retain its zero-based presentation indices")
    for item in key.itertuples(index=False):
        # The original key uses 0..39; both released forms display 1..40.
        if by_code[item.code]["presentation_order"] != int(item.presentation_order) + 1:
            raise PostReaderError("Key/presentation order mismatch")
        if hasattr(item, "n_displayed_tiles") and by_code[item.code]["n_tiles"] != int(item.n_displayed_tiles):
            raise PostReaderError("Key/form displayed-tile support mismatch")
    named, controlling, pairs = [], [], []
    for prototype in range(32):
        block = key.loc[key["prototype_id"].eq(prototype)]
        controlling_code = str(control.loc[control["prototype_id"].eq(prototype), "code"].iloc[0])
        row = by_code[controlling_code]
        reasons = []
        if row["n_tiles"] != 12:
            reasons.append("MONTAGE_SUPPORT_INSUFFICIENT")
        if not row["primary_category"]:
            reasons.append("MISSING_PRIMARY_CATEGORY")
        elif row["primary_category"] == "mixed/other interpretable":
            reasons.append("MIXED_OTHER_PRIMARY")
        if row["confidence_1_to_5"] is None:
            reasons.append("MISSING_CONFIDENCE")
        elif row["confidence_1_to_5"] < 3:
            reasons.append("CONFIDENCE_BELOW_3")
        if row["artifact_uninterpretable"] is None:
            reasons.append("MISSING_ARTIFACT_RESPONSE")
        elif row["artifact_uninterpretable"] == "yes":
            reasons.append("ARTIFACT_UNINTERPRETABLE")
        if not reasons:
            named.append(prototype)
        controlling.append({"prototype_id": prototype, "controlling_code": controlling_code,
                            "eligible_for_named_ref": not reasons, "exclusion_reasons": reasons,
                            "response": row})
        if len(block) == 2:
            duplicate = block.loc[~block["is_controlling_read"].eq(True)]
            if len(duplicate) != 1 or int(duplicate.iloc[0]["occurrence"]) != 1:
                raise PostReaderError("Duplicate occurrence key is invalid")
            other = by_code[str(duplicate.iloc[0]["code"])]
            raw_agreement = category_agreement(row, other)
            flags_recorded = all(value["artifact_uninterpretable"] == "no" for value in (row, other))
            agreement = raw_agreement if flags_recorded else "unconfirmed_artifact_status_or_flagged"
            pairs.append({"prototype_id": prototype, "controlling_code": controlling_code,
                          "duplicate_code": other["blinded_code"], "agreement": agreement,
                          "categorical_agreement_descriptive_only": raw_agreement,
                          "agreement_established": agreement in ("exact", "partial")})
        elif len(block) != 1:
            raise PostReaderError("Invalid presentation multiplicity")
    if len(pairs) != 8:
        raise PostReaderError("Duplicate denominator is not eight")
    agreements = sum(pair["agreement_established"] for pair in pairs)
    conditions = {"duplicate_support_preflight": duplicate_preflight == "DUPLICATE_PREFLIGHT_PASS",
                  "at_least_16_named_ref": len(named) >= 16, "at_least_6_of_8_agreements": agreements >= 6}
    named_oof, counts, attributable_sets = {}, {}, []
    for fold in range(5):
        block = mapping.loc[mapping["outer_fold"].eq(fold)]
        if sorted(block["source_prototype_id"].astype(int)) != list(range(32)) or block["reference_prototype_id"].nunique() != 32:
            raise PostReaderError(f"Fold {fold} mapping is not a bijection")
        matched = block.loc[block["cosine_similarity"].ge(0.8)]
        eligible = matched.loc[matched["reference_prototype_id"].isin(named)]
        named_oof[str(fold)] = sorted(eligible["source_prototype_id"].astype(int).tolist())
        counts[str(fold)] = {"feature_count": len(eligible), "reference_coverage": sorted(eligible["reference_prototype_id"].astype(int).tolist()),
                             "unmatched_source_coordinates": sorted(block.loc[block["cosine_similarity"].lt(0.8), "source_prototype_id"].astype(int).tolist())}
        attributable_sets.append(set(matched["reference_prototype_id"].astype(int)))
    category_counts = Counter(item["response"]["primary_category"] or "MISSING" for item in controlling)
    return {"name_gate_status": "NAME_GATE_PASS" if all(conditions.values()) else "NAME_GATE_FAIL",
            "name_gate_conditions": conditions, "ALL32": list(range(32)), "NAMED_REF": named,
            "NAMED_OOF": named_oof, "fold_coverage": counts,
            "ATTRIBUTABLE_REF": sorted(set.intersection(*attributable_sets)),
            "named_ref_count": len(named), "controlling_reads": controlling,
            "controlling_primary_category_counts": dict(sorted(category_counts.items())),
            "repeatability": {"numerator": agreements, "denominator": 8,
                              "ci95_clopper_pearson": cp_interval(agreements),
                              "exact": sum(pair["agreement"] == "exact" for pair in pairs),
                              "partial": sum(pair["agreement"] == "partial" for pair in pairs),
                              "pairs": pairs, "missing_artifact_policy": "Conservative no established agreement, as sealed before unblinding.",
                              "categorical_concordance_ignoring_missing_artifact_fields_descriptive_only": dict(Counter(pair["categorical_agreement_descriptive_only"] for pair in pairs))}}


def naming(post: Path = POST, pre_root: Path = PRE) -> dict[str, Any]:
    intake_receipt = intake(post)
    path = post / "naming/naming_freeze.json"
    if path.exists():
        result = verify_seal(path)
        for record in result["source_identities"].values():
            check_identity(record)
        return result
    # All accesses to the embargoed key occur only after the calls above verify
    # the raw seal, pre-unblinding amendment, frozen delivery, and literal intake.
    reader_receipt_path = pre_root / "receipts/reader_package.json"
    reader = read_json(reader_receipt_path)
    key_record = reader["artifacts"]["embargoed_occurrence_key"]
    check_identity(key_record)
    key = pd.read_csv(key_record["path"])
    mapping_receipt = read_json(pre_root / "receipts/mappings.json")
    mapping_record = mapping_receipt["artifacts"]["fold_to_reference"]
    check_identity(mapping_record)
    mapping = pd.read_csv(mapping_record["path"])
    literal_record = intake_receipt["artifacts"]["literal_response_values"]
    rows = read_json(Path(literal_record["path"]))["rows"]
    result = derive_naming(rows, key, mapping, reader["duplicate_preflight_status"])
    result.update({"schema_version": 1, "component": COMPONENT, "created_utc": now(),
                   "status": "NAMING_FROZEN_AFTER_AUTHORIZED_LITERAL_INTAKE",
                   "source_identities": {"amendment": identity(post / "governance_amendment.json"),
                                         "literal_intake_receipt": identity(post / "literal_intake/receipt.json"),
                                         "raw_workbook": intake_receipt["artifacts"]["raw_return"],
                                         "reader_package_receipt": identity(reader_receipt_path),
                                         "embargoed_key": identity(Path(key_record["path"])),
                                         "mapping": identity(Path(mapping_record["path"])),
                                         "runner": identity(Path(__file__))},
                   "recorded_blinding_attestation_present": False,
                   "session_timing": "Same-day completion/return per user; workbook dates remain missing.",
                   "reader_authority": "Single mentor/study pathologist per user; prior legacy review exposure; missing workbook reviewer ID preserved.",
                   "named_model_policy": "No named models or name-dependent gates under NAME_GATE_FAIL.",
                   "qualitative_feedback": intake_receipt["artifacts"]["qualitative_feedback"],
                   "interpretation_scope": "ALL32 unsupervised concept space. Literal reader observations are descriptive and cannot repair missing structured evidence."})
    publish(path, result)
    return result


def joint_anchor_assignment(cosines: np.ndarray) -> tuple[int, int]:
    values = np.asarray(cosines, dtype=float)
    if values.ndim != 2 or values.shape[0] != 2 or values.shape[1] < 2 or not np.isfinite(values).all():
        raise ValueError("Expected finite 2 by k cosine matrix")
    objectives = [(float(values[0, a] + values[1, b]), a, b)
                  for a in range(values.shape[1]) for b in range(values.shape[1]) if a != b]
    maximum = max(item[0] for item in objectives)
    return min((a, b) for score, a, b in objectives if maximum - score <= 1e-12)


def correspondence(post: Path = POST, pre_root: Path = PRE) -> dict[str, Any]:
    name_result = naming(post, pre_root)
    destination = post / "correspondence/results.json"
    if destination.exists():
        result = verify_seal(destination)
        for record in result.get("source_identities", {}).values():
            check_identity(record)
        return result
    records = {"naming_freeze": identity(post / "naming/naming_freeze.json")}
    failures = []
    for name, (path, digest) in pre.LEGACY_INPUTS.items():
        try:
            record = identity(path)
            if record["sha256"] != digest:
                raise PostReaderError("Pinned SHA-256 mismatch")
            records[name] = record
        except (OSError, PostReaderError) as error:
            failures.append({"input": name, "reason": str(error)})
    variants = []
    if not failures:
        try:
            with np.load(pre.LEGACY_INPUTS["legacy_vocab_npz"][0], allow_pickle=False) as archive:
                anchors = np.asarray(archive["pca_mean"], dtype=np.float64) + np.asarray(archive["centroids"][[17, 28]], dtype=np.float64) @ np.asarray(archive["pca_components"], dtype=np.float64)
            anchors /= np.linalg.norm(anchors, axis=1, keepdims=True)
            for k in pre.VARIANT_K:
                for seed in pre.VARIANT_SEEDS:
                    path = pre._vocabulary_path(pre_root, fold=None, k=k, seed=seed)
                    receipt = read_json(path.with_suffix(path.suffix + ".receipt.json"))
                    # The frozen receipt recursively binds the array payload.
                    pre._validate_artifact_tree(receipt, f"vocabulary k{k}/seed{seed}")
                    vocabulary = pre._load_vocabulary(path)
                    vectors = np.asarray(vocabulary.pca_mean, dtype=np.float64) + np.asarray(vocabulary.centroids, dtype=np.float64) @ np.asarray(vocabulary.pca_components, dtype=np.float64)
                    vectors /= np.linalg.norm(vectors, axis=1, keepdims=True)
                    cosines = anchors @ vectors.T
                    pair = joint_anchor_assignment(cosines)
                    records[f"k{k}_seed{seed}"] = identity(path)
                    variants.append({"k": k, "seed": seed, "canonical": k == 32 and seed == pre.VOCAB_SEED,
                                     "p17": {"reference_prototype_id": pair[0], "cosine": float(cosines[0, pair[0]]), "geometry_pass": bool(cosines[0, pair[0]] >= 0.8)},
                                     "p28": {"reference_prototype_id": pair[1], "cosine": float(cosines[1, pair[1]]), "geometry_pass": bool(cosines[1, pair[1]] >= 0.8)}})
        except (KeyError, ValueError, OSError) as error:
            failures.append({"input": "required_geometry", "reason": str(error)})
    axes = {}
    categories = {"p17": "extracellular mucin/mucinous pattern", "p28": "malignant gland-forming epithelium/gland–lumen"}
    for axis, category in categories.items():
        axis_status = "CORRESPONDENCE_NOT_EVALUABLE"
        passes = sum(variant[axis]["geometry_pass"] for variant in variants)
        canonical = next((variant[axis] for variant in variants if variant["canonical"]), None)
        compatibility = None
        if not failures and canonical is not None and name_result["name_gate_status"] == "NAME_GATE_PASS":
            read = next(row["response"] for row in name_result["controlling_reads"] if row["prototype_id"] == canonical["reference_prototype_id"])
            compatibility = "exact" if read["primary_category"] == category else "partial" if category in (read["secondary_category_1"], read["secondary_category_2"]) else "different"
            axis_status = "CORRESPONDENCE_CRITERIA_MET" if canonical["geometry_pass"] and passes >= 7 and compatibility in ("exact", "partial") else "CORRESPONDENCE_CRITERIA_NOT_MET"
        axes[axis] = {"status": axis_status, "legacy_category": category,
                      "geometry_passes": passes, "variant_denominator": 9,
                      "canonical_geometry": canonical, "controlling_category_compatibility": compatibility,
                      "name_gate_status": name_result["name_gate_status"]}
    result = {"schema_version": 1, "component": COMPONENT, "created_utc": now(),
              "status": "CORRESPONDENCE_REPORTED", "source_identities": records,
              "input_failures": failures, "variants": variants, "axes": axes,
              "geometry_method": "Direct joint one-to-one legacy p17/p28 to each variant in backprojected normalized UNI-v1 space; objective tie tolerance 1e-12, lexical pair.",
              "claim_limit": "Legacy vocabulary was target-inclusive/transductive. Geometry cannot authorize names when NAME_GATE_FAIL."}
    publish(destination, result)
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("stage", choices=("intake", "naming", "correspondence", "verify", "run"))
    args = parser.parse_args()
    if args.stage == "intake":
        result = intake()
    elif args.stage == "naming":
        result = naming()
    elif args.stage == "verify":
        verify_authorization()
        verify_frozen_delivery()
        for relative in ("literal_intake/receipt.json", "naming/naming_freeze.json", "correspondence/results.json"):
            result = verify_seal(POST / relative)
            for record in result.get("artifacts", {}).values():
                check_identity(record)
            for record in result.get("source_identities", {}).values():
                check_identity(record)
        result = {"status": "PASS"}
    else:
        result = correspondence()
    print(json.dumps({"stage": args.stage, "status": result["status"],
                      "name_gate_status": result.get("name_gate_status"),
                      "output_root": str(POST)}, indent=2))


if __name__ == "__main__":
    main()
