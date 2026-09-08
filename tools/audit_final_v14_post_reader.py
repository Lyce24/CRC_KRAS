#!/usr/bin/env python3
"""Audit completed v14 source work without refitting or opening target outcomes."""
from __future__ import annotations

import argparse
import html
import json
import re
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from sklearn.metrics import average_precision_score, roc_auc_score
from threadpoolctl import threadpool_limits

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO))
from tools import audit_final_v14_context as context_audit  # noqa: E402
from tools import final_v14_context as context  # noqa: E402
from tools import final_v14_module1 as source  # noqa: E402
from tools import final_v14_module1_evaluate as evaluation  # noqa: E402
from tools import final_v14_post_reader as reader  # noqa: E402

ROOT = source.ROOT
OUT = ROOT / "post_reader_audit"


def check(condition: bool, message: str) -> None:
    if not condition:
        raise ValueError(message)


def parent_snapshot() -> dict:
    original = REPO / "reports/final_v13"
    snapshot = REPO / "reports/snapshots/final_v13_pre_v14_20260903"
    files = sorted(p.relative_to(original) for p in original.rglob("*") if p.is_file())
    check(files == sorted(p.relative_to(snapshot) for p in snapshot.rglob("*") if p.is_file()), "Parent/snapshot file inventory differs")
    for relative in files:
        check(reader.identity(original / relative)["sha256"] == reader.identity(snapshot / relative)["sha256"], f"Parent changed: {relative}")
    return {"status": "PASS", "byte_identical_files": len(files),
            "original_snapshot_receipt": reader.identity(REPO / "reports/snapshots/final_v13_pre_v14_20260903.identity.json")}


def audit_source() -> tuple[dict, pd.DataFrame]:
    frame = source.patient_frame(source.PRE)
    summary = source.read_sealed(source.OUT / "source_fit_summary.json")
    contract = source.read_sealed(source.OUT / "source_fit_contract.json")
    for pin in contract["input_pins"] + contract["code"]:
        source.verify_record(pin)
    predictions = pd.read_parquet(source.verify_record(summary["predictions"]))
    check(predictions.patient_id.tolist() == frame.patient_id.tolist(), "Prediction patient order differs")
    check(set(summary["representations"]) == {"ALL32"}, "Failed name gate must suppress named models")
    check(len(frame) == 1239 and int(frame.label.sum()) == 501, "Source population census differs")
    model_counts = {"direct": 0, "ridge": 0, "joint": 0, "clinical": 0}
    candidate_sets = 0
    for fold in range(5):
        item = source.read_sealed(source.OUT / f"fits/ALL32/all/fold_{fold}.json")
        train = frame.loc[frame.k_fold.ne(fold)].reset_index(drop=True)
        held = frame.loc[frame.k_fold.eq(fold)].reset_index(drop=True)
        check(item["train_patients"] == train.patient_id.tolist(), "Outer train roster mismatch")
        check(item["test_patients"] == held.patient_id.tolist(), "Outer test roster mismatch")
        check(not set(item["train_patients"]) & set(item["test_patients"]), "Source leakage across outer fold")
        check(item["coordinates"] == list(range(32)), "ALL32 coordinate selection changed")
        X = source.aligned_matrix(source.PRE / f"profiles/outer_fold_{fold}/patient_profiles.parquet", frame)
        Xtr, Xte = X[frame.k_fold.ne(fold)], X[frame.k_fold.eq(fold)]
        Ctr, Cte, encoder = source.clinical_design(train, held)
        check(len(item["inner_splits"]) == 4, "Module-I requires four inner folds")
        observed_validation = []
        for split in item["inner_splits"]:
            a, b = set(split["train"]), set(split["validation"])
            check(not a & b and a | b == set(range(len(train))), "Inner split does not partition source-training patients")
            observed_validation += split["validation"]
        check(sorted(observed_validation) == list(range(len(train))), "Inner validation coverage differs")
        for name, model in item["models"].items():
            check(model["status"] == "ESTIMABLE", f"Required source model unavailable: {fold}/{name}")
            kind = "ridge" if name.startswith("ridge_seed") else "direct" if name == "concept" else name
            check(kind in model_counts, "Unknown source model")
            model_counts[kind] += 1
            if name == "clinical":
                score = Cte @ np.asarray(model["coef"]) + model["intercept"]
                check(model["clinical_encoder"] == encoder, "Inherited clinical preprocessing differs")
            else:
                a, b = (np.column_stack([Xtr, Ctr]), np.column_stack([Xte, Cte])) if name == "joint" else (Xtr, Xte)
                np.testing.assert_allclose(model["scaler_mean"], a.mean(0), atol=1e-12)
                expected_scale = a.std(0)
                expected_scale[expected_scale == 0] = 1
                np.testing.assert_allclose(model["scaler_scale"], expected_scale, atol=1e-12)
                score = ((b - np.asarray(model["scaler_mean"])) / np.asarray(model["scaler_scale"])) @ np.asarray(model["coef"]) + model["intercept"]
                candidates = model["candidates"]
                check([x["penalty"] for x in candidates] == list(source.modeling.PENALTIES), "Penalty grid changed")
                for candidate in candidates:
                    valid = len(candidate["inner_losses"]) == 4 and all(x is not None and np.isfinite(x) for x in candidate["inner_losses"]) and not candidate["failures"]
                    check(candidate["eligible"] == valid, "Candidate convergence/finite-loss eligibility differs")
                    if valid:
                        np.testing.assert_allclose(candidate["mean_loss"], np.mean(candidate["inner_losses"]), atol=1e-12)
                eligible = [x for x in candidates if x["eligible"]]
                best = min(x["mean_loss"] for x in eligible)
                ties = [x["penalty"] for x in eligible if abs(x["mean_loss"] - best) <= 1e-12]
                selected = max(ties) if kind == "ridge" else min(ties)
                check(model["selected_penalty"] == selected, "Stronger-regularization tie break differs")
                candidate_sets += 1
            np.testing.assert_allclose(score, model["predictions"], atol=1e-12, rtol=0)
            saved = predictions.loc[frame.k_fold.eq(fold), f"ALL32__all__{name}"].to_numpy(float)
            np.testing.assert_allclose(score, saved, atol=1e-12, rtol=0)
    check(model_counts == {"direct": 5, "ridge": 25, "joint": 5, "clinical": 5}, "Required source model count differs")
    ridge = predictions[[f"ALL32__all__ridge_seed{s}" for s in range(42, 47)]].to_numpy(float)
    check(np.isfinite(ridge).all(), "Missing ridge seed predictions")
    np.testing.assert_allclose(ridge.mean(1), predictions.ALL32__all__ridge_mean, atol=1e-12)
    return {"status": "PASS", "patients": 1239, "model_counts": model_counts,
            "candidate_grids_replayed": candidate_sets, "heldout_prediction_tolerance": 1e-12,
            "source_summary": reader.identity(source.OUT / "source_fit_summary.json")}, predictions


def audit_inference(predictions: pd.DataFrame) -> dict:
    result = source.read_sealed(evaluation.OUT / "results.json")
    contract = source.read_sealed(evaluation.OUT / "evaluation_contract.json")
    for key in ("source_summary", "module3_dependency", "code", "derived_stage_source"):
        source.verify_record(contract[key])
    stages = pd.read_csv(contract["derived_stage_source"]["path"], usecols=contract["derived_stage_read_columns"])
    stages = stages.loc[stages.patient_uid.isin(predictions.patient_id) & stages.specimen_role.str.casefold().eq("primary")]
    stages = stages.drop_duplicates("patient_uid").set_index("patient_uid")
    predictions = predictions.copy()
    predictions["stage_known_derived"] = stages.loc[predictions.patient_id, "stage_group_major_filled"].isin(["I", "II", "III", "IV"]).to_numpy()
    check(int(predictions.stage_known_derived.sum()) == 1060, "Derived-stage evaluation census differs")
    points = evaluation.metrics(predictions, "ALL32", np.arange(len(predictions)))
    intervals = 0
    with np.load(source.verify_record(result["bootstrap"]["arrays"]), allow_pickle=False) as archive:
        check(len(archive.files) == 75 and all(archive[key].shape == (10000,) for key in archive.files), "Bootstrap array census/budget differs")
        for key, block in result["ALL32"]["pooled"].items():
            np.testing.assert_allclose(block["estimate"], points[key], atol=1e-12)
            check(block == evaluation.interval(points[key], archive[f"ALL32/{key}"], ratio=key.endswith("retention_ratio")), f"Pooled CI mismatch: {key}")
            intervals += 1
        for cohort, estimates in result["ALL32"]["subcohort_fidelity"].items():
            frame = predictions.loc[predictions.subcohort.eq(cohort)]
            local = evaluation.fidelity(frame.label.to_numpy(int), frame.full_logit.to_numpy(float), frame.ALL32__all__ridge_mean.to_numpy(float))
            for key, block in estimates.items():
                check(block == evaluation.interval(local[key], archive[f"ALL32/subcohort/{cohort}/{key}"], ratio=key == "retention_ratio"), "Subcohort fidelity CI differs")
                intervals += 1
        for reference, estimates in result["ALL32"]["standardized_restriction"].items():
            for display, key in (("standardized_A_auroc", "A"), ("standardized_D_auroc", "D"), ("standardized_D_minus_A", "D_minus_A")):
                block = estimates[display]
                check(block == evaluation.interval(block["estimate"], archive[f"ALL32/standardized/{reference}/{key}"]), "Standardized restriction CI differs")
                intervals += 1
        # Independently replay three shared patient draws, with sklearn's AUROC.
        cells = [np.flatnonzero(predictions.subcohort.eq(c).to_numpy() & predictions.label.eq(k).to_numpy()) for c in sorted(predictions.subcohort.unique()) for k in (0, 1)]
        rng = np.random.default_rng(20260827)
        for draw in range(10000):
            idx = np.concatenate([rng.choice(cell, len(cell), replace=True) for cell in cells])
            if draw not in (0, 4999, 9999):
                continue
            frame = predictions.iloc[idx]
            y, full, ridge, concept = (frame.label.to_numpy(int), frame.full_logit.to_numpy(float), frame.ALL32__all__ridge_mean.to_numpy(float), frame.ALL32__all__concept.to_numpy(float))
            values = {"all/concept/auroc": roc_auc_score(y, concept), "all/concept/auprc": average_precision_score(y, concept),
                      "all/concept_minus_full/auroc": roc_auc_score(y, concept) - roc_auc_score(y, full),
                      "fidelity/r2_oof": 1 - np.sum((full - ridge) ** 2) / np.sum((full - full.mean()) ** 2)}
            for key, value in values.items():
                np.testing.assert_allclose(archive[f"ALL32/{key}"][draw], value, atol=1e-12)
    y = predictions.label.to_numpy(int)
    direct = roc_auc_score(y, predictions.ALL32__all__concept)
    full = roc_auc_score(y, predictions.full_logit)
    ridge = predictions.ALL32__all__ridge_mean.to_numpy(float)
    r2 = 1 - np.sum((predictions.full_logit - ridge) ** 2) / np.sum((predictions.full_logit - predictions.full_logit.mean()) ** 2)
    np.testing.assert_allclose([direct, full, r2], [points["all/concept/auroc"], points["all/full/auroc"], points["fidelity/r2_oof"]], atol=1e-12)
    h11 = result["ALL32"]["pooled"]["all/concept/auroc"]["ci95"][0] > 0.5
    h13 = result["ALL32"]["standardized_restriction"]["A_all_primary"]["standardized_D_minus_A"]["ci95"][0] > 0
    check(result["gates"]["H1.1"].endswith("_SUPPORTED") == h11, "H1.1 gate differs")
    check(result["gates"]["H1.3"].endswith("_SUPPORTED") == h13, "H1.3 gate differs")
    return {"status": "PASS", "bootstrap_intervals_replayed": intervals, "bootstrap_arrays": 75,
            "draws_per_array": 10000, "independent_sklearn_draws": [0, 4999, 9999],
            "concept_auroc": float(direct), "full_auroc": float(full), "ridge_r2": float(r2), "gates": result["gates"]}


def audit_report() -> dict:
    root = REPO / "reports/final_v14"
    status = reader.read_json(root / "STATUS.json")
    check(status["completed_bundle"] is False, "Partial source work was called a completed campaign")
    for record in status["inputs"]:
        source.verify_record(record)
    for record in status["report_artifacts"]:
        reader.check_identity(record, root / record["path"])
    report = (root / "Results.md").read_text()
    results = source.read_sealed(evaluation.OUT / "results.json")
    for key in ("all/concept/auroc", "all/concept_minus_full/auroc", "fidelity/r2_oof", "all/full/auroc"):
        block = results["ALL32"]["pooled"][key]
        check(f"{block['estimate']:.3f} [{block['ci95'][0]:.3f}, {block['ci95'][1]:.3f}]" in report, f"Report headline does not match artifact: {key}")
    check("not yet complete" in report and "NAME_GATE_FAIL" in report, "Report completion/name boundaries missing")
    for target in re.findall(r"\]\(([^)]+)\)", report):
        if not target.startswith(("http:", "https:", "#")):
            check((root / target.split("#")[0]).exists(), f"Broken report artifact link: {target}")
    atlas = (root / "mentor_atlas.html").read_text()
    images = re.findall(r'<img[^>]+src="([^"]+)"', atlas)
    check(len(images) == 32 and atlas.count("<article ") == 32, "Atlas is not 32 controlling concepts")
    names = source.read_sealed(source.NAMING)
    for item, image in zip(names["controlling_reads"], images, strict=True):
        image_path = (root / html.unescape(image)).resolve()
        check(image_path.name == item["controlling_code"] + ".jpg", "Atlas points to a duplicate/wrong montage")
        check(image_path.is_file(), "Atlas image missing")
        comment = item["response"]["free_text_description"]
        if comment:
            check(html.escape(comment) in atlas, "Atlas changed original free text")
    workbook = load_workbook(root / "Aim4_Results.xlsx", read_only=True, data_only=False)
    return {"status": "PASS", "artifacts": len(status["report_artifacts"]), "atlas_controlling_montages": 32,
            "report_status": reader.identity(root / "STATUS.json"), "workbook_sheets": workbook.sheetnames,
            "full_campaign_complete": False}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--skip-report", action="store_true", help="Audit numerical work while the presentation artifacts are still being rebuilt")
    arguments = parser.parse_args()
    with threadpool_limits(limits=1):
        parent = parent_snapshot()
        reader.intake()
        names = reader.naming()
        source_result, predictions = audit_source()
        inference = audit_inference(predictions)
        context_result = context_audit.run()
        report = None if arguments.skip_report else audit_report()
    value = {"schema_version": 1, "component": "final_v14_completed_source_independent_audit", "status": "PASS_SOURCE_COMPLETED_CAMPAIGN_INCOMPLETE",
             "created_utc": reader.now(), "auditor": reader.identity(Path(__file__)), "parent_snapshot": parent,
             "reader_return": {"raw_workbook": reader.identity(reader.POST / "raw_return/completed_review_form.xlsx"), "name_gate": names["name_gate_status"],
                               "named_ref_count": names["named_ref_count"], "original_montages_verified": 40},
             "source_models": source_result, "source_inference": inference, "context": context_result,
             "report": report, "scientific_scope": "Completed Module-I source work only. No full Aim-4 completion claim; Module-II targets, Module-III ladder, geometric correspondence, and raw grid sensitivity remain pending inaccessible archive inputs."}
    if not arguments.skip_report:
        reader.publish(OUT / "source_audit_receipt.json", value)
    print(json.dumps({"status": value["status"], "source": source_result, "inference": inference, "report": report}, indent=2))


if __name__ == "__main__":
    main()
