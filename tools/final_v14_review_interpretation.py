"""Descriptive companion to the frozen v14 review; no fitting or relabeling."""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import math
import os
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

REPO = Path(__file__).resolve().parents[1]
RUN = REPO / "reports/reruns/final_v14_additions_20260903"
OUT = REPO / "reports/final_v14_review_interpretation"
MENTOR = REPO / "reports/final_v14/tables/Mentor_annotations.csv"
NAMES = RUN / "e4v_post_reader_xlsx/naming/naming_freeze.json"
LITERAL = RUN / "e4v_post_reader_xlsx/literal_intake/literal_response_values.json"
CONTEXT = RUN / "e4m1_context/results.json"
CORRESPONDENCE = RUN / "e4v_correspondence_replay/runs/add086ae82ef5820b15847827e7eb8d9d25d7c3740cc69c38b449ab8e7ed1274/results.json"
PANEL_KEYS = [
    "abundance__oof_score", "abundance__kras", "attention__oof_score", "attention__kras",
    "abundance__age_at_diagnosis", "abundance__braf", "abundance__msi_dmmr",
    "abundance__sex", "abundance__source_subcohort", "abundance__stage_group_major",
    "abundance__tumor_site_group",
]
NOT_ATTRIBUTABLE = "NOT_IN_ATTRIBUTABLE_REF"


def identity(path: Path) -> dict:
    raw = path.read_bytes()
    return {"path": str(path.resolve()), "sha256": hashlib.sha256(raw).hexdigest(), "size_bytes": len(raw)}


def read_json(path: Path) -> dict:
    return json.loads(path.read_text())


def load_inputs() -> tuple[dict, list[dict]]:
    pins = {}

    def read(path: Path, expected: dict | None = None, sealed: bool = False) -> dict:
        pin = identity(path)
        if expected is not None:
            assert pin == expected, f"Changed source artifact: {path}"
        if sealed:
            seal_path = Path(str(path) + ".seal.json")
            seal = read_json(seal_path)
            assert seal["sha256"] == pin["sha256"] == seal["artifact"]["sha256"]
            pins[str(seal_path)] = identity(seal_path)
        pins[str(path)] = pin
        return read_json(path)

    names = read(NAMES, sealed=True)
    literal = read(LITERAL, sealed=True)
    context = read(CONTEXT, sealed=True)
    correspondence = read(CORRESPONDENCE, sealed=True)
    assert names["name_gate_status"] == context["name_gate_status"] == "NAME_GATE_FAIL"
    assert names["ALL32"] == list(range(32))
    assert context["status"] == "CONTEXT_ASSOCIATIONS_COMPLETE"
    assert correspondence["status"] == "CORRESPONDENCE_GEOMETRY_COMPLETE"
    assert len(correspondence["variants"]) == 9
    with MENTOR.open(newline="") as handle:
        mentor = list(csv.DictReader(handle))
    pins[str(MENTOR)] = identity(MENTOR)
    assert [int(row["Concept"]) for row in mentor] == list(range(32))
    reads = {r["prototype_id"]: r for r in names["controlling_reads"]}
    raw = {r["blinded_code"]: r for r in literal["rows"]}
    assert len(reads) == 32 and len(raw) == 40
    fields = {"Controlling code": "blinded_code", "Primary category": "primary_category",
              "Secondary 1": "secondary_category_1", "Secondary 2": "secondary_category_2",
              "Mentor comment": "free_text_description", "Confidence": "confidence_1_to_5",
              "Artifact flag": "artifact_uninterpretable"}
    for row in mentor:
        controlling = reads[int(row["Concept"])]
        response = controlling["response"]
        assert response == raw[controlling["controlling_code"]]
        for column, field in fields.items():
            expected = "" if response[field] is None else str(response[field])
            if column == "Confidence" and row[column] and expected:
                assert float(row[column]) == float(expected)
            else:
                assert row[column] == expected, (row["Concept"], column)
        assert row["Named-set eligible"] == str(controlling["eligible_for_named_ref"])
    panels = {}
    for key in PANEL_KEYS:
        pin = context["artifacts"][key]
        panel = read(Path(pin["path"]), pin)
        assert [r["prototype_id"] for r in panel["rows"]] == names["ATTRIBUTABLE_REF"]
        assert all(r["bh_family_size"] == 25 for r in panel["rows"])
        panels[key] = panel
    coefficient_pin = context["artifacts"]["coefficient_stability"]
    coefficients = read(Path(coefficient_pin["path"]), coefficient_pin, sealed=True)
    assert [r["prototype_id"] for r in coefficients["rows"]] == names["ATTRIBUTABLE_REF"]
    for row in coefficients["rows"]:
        for family, expected in (("logistic", 5), ("ridge", 25)):
            record = row[family]
            assert record["available"] == record["expected"] == expected
            assert sum(record[k] for k in ("positive", "negative", "zero")) == expected
    table_pin = context["artifacts"]["association_tests"]
    assert identity(Path(table_pin["path"])) == table_pin
    pins[table_pin["path"]] = table_pin
    with Path(table_pin["path"]).open(newline="") as handle:
        test_rows = list(csv.DictReader(handle))
    assert len(test_rows) == 275
    for row in test_rows:
        key = row["representation"] + "__" + row["covariate"]
        source = next(r for r in panels[key]["rows"] if r["prototype_id"] == int(row["prototype_id"]))
        for field, value in row.items():
            expected = source[field]
            if isinstance(expected, (int, float)):
                assert math.isclose(float(value), expected, rel_tol=1e-14, abs_tol=1e-15)
            else:
                assert value == ("" if expected is None else str(expected))
    return {"names": names, "raw": raw, "reads": reads, "mentor": mentor,
            "context": context, "panels": panels, "coefficients": coefficients,
            "correspondence": correspondence}, sorted(pins.values(), key=lambda p: p["path"])


def compact_json(value) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False)


def assemble(data: dict) -> dict[str, list[dict]]:
    attributable = set(data["names"]["ATTRIBUTABLE_REF"])
    panels = {key: {r["prototype_id"]: r for r in value["rows"]} for key, value in data["panels"].items()}
    coefficients = {r["prototype_id"]: r for r in data["coefficients"]["rows"]}
    pairs = {r["prototype_id"]: r for r in data["names"]["repeatability"]["pairs"]}
    wide, long_rows, duplicates, geometry = [], [], [], []
    for mentor in data["mentor"]:
        pid = int(mentor["Concept"])
        response = data["reads"][pid]["response"]
        row = {"prototype_id": pid, **{key: value for key, value in response.items()},
               "eligible_for_named_ref": data["reads"][pid]["eligible_for_named_ref"],
               "name_gate_status": data["names"]["name_gate_status"],
               "attribution_status": "ATTRIBUTABLE_REF" if pid in attributable else NOT_ATTRIBUTABLE,
               "attribution_scope": "Mapped at cosine >=0.80 in all five outer folds" if pid in attributable else "Fails the all-five-fold mapping criterion; no reference-coordinate context inference",
               "original_montage": f"../../reviews/v14_xlsx/FOR_PATHOLOGIST/montages/{response['blinded_code']}.jpg"}
        pair = pairs.get(pid)
        duplicate = data["raw"][pair["duplicate_code"]] if pair else None
        row["duplicate_presentation_status"] = "HAS_RESAMPLED_PRESENTATION" if pair else "NO_DUPLICATE_DESIGNATED"
        row["duplicate_category_concordance_descriptive"] = pair["categorical_agreement_descriptive_only"] if pair else None
        for field in response:
            row["duplicate_" + field] = duplicate[field] if duplicate else None
        if pair:
            duplicates.append({"prototype_id": pid, **pair,
                               "controlling_primary_literal": response["primary_category"],
                               "duplicate_primary_literal": duplicate["primary_category"],
                               "controlling_comment_literal": response["free_text_description"],
                               "duplicate_comment_literal": duplicate["free_text_description"]})
        for key in PANEL_KEYS:
            panel_row = panels[key].get(pid)
            row[key + "__status"] = panel_row["status"] if panel_row else NOT_ATTRIBUTABLE
            for field in ("n_complete_cases", "n_missing", "p_value", "q_value_bh", "partial_r2", "not_estimable_reason"):
                row[key + "__" + field] = panel_row[field] if panel_row else None
            if key in PANEL_KEYS[:4]:
                contrast = panel_row["contrasts"][0] if panel_row else None
                for field in ("adjusted_coefficient", "standardized_effect"):
                    row[key + "__" + field] = contrast[field] if contrast else None
                    interval = contrast[field + "_bootstrap"] if contrast else None
                    for i, label in enumerate(("ci95_lower", "ci95_upper")):
                        row[key + "__" + field + "__" + label] = interval["ci95"][i] if interval and interval["ci95"] else None
            representation, covariate = key.split("__")
            long_row = {"prototype_id": pid, "representation": representation, "covariate": covariate,
                        "status": panel_row["status"] if panel_row else NOT_ATTRIBUTABLE}
            for field in ("n_complete_cases", "n_missing", "hc3_statistic", "hc3_distribution", "tested_degrees_of_freedom",
                          "p_value", "q_value_bh", "bh_family_size", "partial_r2", "not_estimable_reason"):
                long_row[field] = panel_row[field] if panel_row else None
            long_row["contrasts_json"] = compact_json(panel_row["contrasts"]) if panel_row else None
            long_row["partial_r2_bootstrap_json"] = compact_json(panel_row["partial_r2_bootstrap"]) if panel_row else None
            long_row["design_json"] = compact_json(data["panels"][key]["design"])
            long_rows.append(long_row)
        for family in ("logistic", "ridge"):
            record = coefficients[pid][family] if pid in coefficients else None
            row[family + "__status"] = "DESCRIPTIVE_STABILITY_AVAILABLE" if record else NOT_ATTRIBUTABLE
            for field in ("median_standardized_coefficient", "positive", "negative", "zero", "available", "expected"):
                row[family + "__" + field] = record[field] if record else None
            row[family + "__coefficients_json"] = compact_json(record["coefficients"]) if record else None
        wide.append(row)
    for variant in data["correspondence"]["variants"]:
        row = {key: variant[key] for key in ("k", "seed", "canonical")}
        for anchor in ("p17", "p28"):
            for key, value in variant[anchor].items():
                row[f"legacy_{anchor}__{key}"] = value
            row[f"legacy_{anchor}__named_status"] = data["correspondence"]["axes"][anchor]["status"]
        geometry.append(row)
    assert len(wide) == 32 and len(long_rows) == 352 and len(duplicates) == 8 and len(geometry) == 9
    assert Counter(r["status"] for r in long_rows) == {"ESTIMABLE": 274, "NOT_ESTIMABLE": 1, NOT_ATTRIBUTABLE: 77}
    return {"All32_evidence": wide, "Context_associations": long_rows,
            "Duplicate_presentations": duplicates, "Legacy_geometry": geometry}


def csv_text(rows: list[dict]) -> str:
    buffer = io.StringIO(newline="")
    writer = csv.DictWriter(buffer, fieldnames=list(rows[0]), lineterminator="\n")
    writer.writeheader()
    writer.writerows(rows)
    return buffer.getvalue()


def cell(value) -> str:
    if value is None or value == "":
        return "Unrecorded"
    return str(value).replace("|", "\\|").replace("\r", "").replace("\n", "<br>")


def table(headers: list[str], rows: list[list]) -> str:
    return "\n".join(["| " + " | ".join(headers) + " |", "| " + " | ".join("---" for _ in headers) + " |",
                      *("| " + " | ".join(cell(v) for v in row) + " |" for row in rows)])


def annotation_text(response: dict) -> str:
    terms = [response[key] for key in ("primary_category", "secondary_category_1", "secondary_category_2")]
    return "; ".join(term for term in terms if term) or "Primary category unrecorded"


def numerical_cell(row: dict, key: str) -> str:
    status = row[key + "__status"]
    if status != "ESTIMABLE":
        return "Outside mapping set" if status == NOT_ATTRIBUTABLE else status
    effect = row[key + "__standardized_effect"]
    low = row[key + "__standardized_effect__ci95_lower"]
    high = row[key + "__standardized_effect__ci95_upper"]
    return f"{effect:+.3f} [{low:+.3f}, {high:+.3f}]; q={row[key + '__q_value_bh']:.3g}"


def render_markdown(data: dict, tables: dict[str, list[dict]]) -> str:
    wide = tables["All32_evidence"]
    counts = data["names"]["repeatability"]["categorical_concordance_ignoring_missing_artifact_fields_descriptive_only"]
    pathology = table(["v14 ID / original montage", "Literal recorded categories", "Literal mentor comment"], [
        [f"[p{r['prototype_id']:02d}]({r['original_montage']})", annotation_text(r), r["free_text_description"]]
        for r in wide])
    numeric = table(["v14 ID", "Abundance vs OOF score", "Abundance vs KRAS", "Attention vs OOF score", "Attention vs KRAS"], [
        [f"p{r['prototype_id']:02d}", *(numerical_cell(r, key) for key in PANEL_KEYS[:4])] for r in wide])
    stability = table(["v14 ID", "Ridge coefficient signs + / − / 0", "Direct classifier signs + / − / 0"], [
        [f"p{r['prototype_id']:02d}", *(f"{r[family + '__positive']} / {r[family + '__negative']} / {r[family + '__zero']} ({r[family + '__available']} fits)" if r[family + '__available'] is not None else "Outside mapping set" for family in ("ridge", "logistic"))]
        for r in wide])
    families = table(["Representation", "Covariate", "Complete-case n", "Estimable / fixed family", "BH q<0.05"], [
        [panel["representation"], panel["covariate"], panel["n_complete_cases"],
         f"{sum(r['status'] == 'ESTIMABLE' for r in panel['rows'])}/25",
         sum(r["status"] == "ESTIMABLE" and r["q_value_bh"] < .05 for r in panel["rows"])]
        for panel in data["panels"].values()])
    repeat = table(["v14 ID", "Original / resampled code", "Literal primary: original", "Literal primary: resampled", "Descriptive category agreement"], [
        [f"p{r['prototype_id']:02d}", r["controlling_code"] + " / " + r["duplicate_code"], r["controlling_primary_literal"], r["duplicate_primary_literal"], r["categorical_agreement_descriptive_only"]]
        for r in tables["Duplicate_presentations"]])
    geometry = table(["k", "Seed", "Legacy p17 → v14 ID", "Cosine", "Legacy p28 → v14 ID", "Cosine"], [
        [r["k"], r["seed"], r["legacy_p17__prototype_id"], f"{r['legacy_p17__cosine']:.6f}", r["legacy_p28__prototype_id"], f"{r['legacy_p28__cosine']:.6f}"]
        for r in tables["Legacy_geometry"]])
    return f"""# FINAL-v14: interpretation of the completed mentor review

The review identifies recognizable tissue structure in the selected examples, while showing that several prototype groups contain mixtures and require precise pathology wording. The useful conclusion is a qualified tissue interpretation of the dictionary. These annotations do not establish that every assigned tile has the same morphology, or that a named tissue pattern explains KRAS prediction.

This companion brings the accepted mentor return together with already completed numerical results. It adds no model, association test, category, imputed response, feature selection or change to a sealed gate. The [main analysis](../final_v14/Results.md) and its completion process remain separate. **NAME_GATE_FAIL** remains in force. The observations below are descriptive; statistical inference remains under prototype IDs.

[All 32 evidence rows](All32_evidence.csv) · [Workbook](Review_Interpretation.xlsx) · [Full context panels](Context_associations.csv) · [Input identities and verification](input_identity_manifest.json)

## What the mentor's observations add

1. **The examples contain recognizable tissue patterns.** Recorded categories include gland-forming epithelium (for example p03, p12 and p22), mucinous pattern (p17), adipose (p16/p29), smooth muscle (p15/p19), blood/vessels (p07), and benign colonic epithelium (p20/p23). These are literal annotations of the selected montages, with the qualifications below; they are not new validated prototype names.
2. **Benign fibrosis and desmoplasia should remain distinct in the interpretation.** The comments describe p01/p02 as largely desmoplastic, p10/p31 as benign fibrosis, and p04 as benign fibrosis with inflammatory cells except one desmoplastic example. The p24 comment describes fibrosis, fibrin deposits and cautery and explicitly asks that desmoplasia and fibrous stroma be separated. The frozen category system is unchanged; this finer distinction stays in the reader's text.
3. **Several groups contain different tissues within the same montage.** For p30 the mentor lists fibrosis, lymphoid tissue, a blood vessel and small tumor components. For p27 the comment identifies benign fibrosis, nerves and lymphoid tissue. For p23, two examples within a primarily benign epithelial group appear malignant or dysplastic. A single category cannot convey all of these observations.
4. **Lumen-poor tumor clusters are not automatically poorly differentiated tumor.** For p25 the mentor describes tumor clusters almost without lumina but explicitly rejects calling them solid and poorly differentiated. Some p05/p09 examples have papillary features; several p21 examples may not be tumor. These qualifications are retained alongside the original category fields.
5. **Edge effects and damaged material qualify necrosis/debris interpretations.** For p06, two tiles are explicitly described as tissue edges rather than cancer necrosis. The p11 comment lists red blood cells, debris, an uncertain example and fibrosis at a tissue edge. The p24 comment also identifies cautery. None of these remarks supplies a missing structured artifact rating.

The montage design deliberately favors examples close to each centroid, with at most one tile per patient and source-balanced targets. Selection starts with the closest-distance decile and expands to the quartile or all eligible candidates when needed. The 12 tiles per controlling montage are selected exemplars, **not a random sample of all assigned tiles**. Consequently neither a cluster-wide purity percentage nor the prevalence of a tissue pattern can be estimated from this review.

## Complete pathology cross-index

All 32 IDs appear in fixed numerical order, including mixed and uncertain groups. Categories and comments below are reproduced literally; “Unrecorded” is a display marker for a missing field. Blank fields stay blank in the CSV/workbook. Both controlling and resampled responses, including their original blanks, are retained in the evidence CSV.

{pathology}

## Existing coordinate-level evidence

Twenty-five reference coordinates meet cosine ≥0.80 mapping in all five outer folds. The other seven (p00, p06, p12, p15, p25, p26 and p27) remain in the all-32 prediction models and in this cross-index, but receive no reference-coordinate context association or coefficient summary. Their missing summaries are **not zero effects** and are unrelated to whether the mentor recognized the tissue.

The following table reports standardized adjusted effects [descriptive bootstrap 95% CI] and HC3/BH q-values. OOF score means the full-model native logit per one standard deviation; KRAS means mutant minus wild type. The response is arcsine-square-root abundance, or separately transformed attention mass, and the displayed effect is divided by that response's standard deviation. Each panel adjusts for source subcohort. Positive and negative values describe the coordinate, not a named tissue mechanism. The adjacent pathology cross-index is an annotation lookup; it does not turn these tests into validated named-morphology associations. No favorable subset of coordinates is singled out.

{numeric}

Nine abundance panels cover score, KRAS, age, BRAF, MSI, sex, source, stage and site; two separate attention panels cover score and KRAS. Covariates are examined separately with source fixed effects, except the source panel itself. This is not a model simultaneously adjusted for all clinical factors. In total 275 tests were attempted and 274 were estimable. The non-estimable p13 joint source-subcohort test remains NOT_ESTIMABLE with p=q=1 in its fixed family; preserved bootstrap values do not repair that failed test. The other 77 cells in the 352-row complete grid explicitly mark the seven coordinates outside the mapping set. All multilevel contrasts, sample counts, effects, intervals, reasons and original status fields are retained in [the full context table](Context_associations.csv); no test or interval was recomputed here.

{families}

Source-subcohort associations in 24/25 coordinates show that the representation varies with data-source context as well as morphology. Associations with age, site, MSI, BRAF and stage provide additional context. They do not by themselves establish causal confounding or explain why a classifier succeeds or fails. BH correction remains separate within each prespecified 25-coordinate representation/covariate family; there is no new cross-family selection rule in this companion.

The coefficient counts below summarize the already fitted 25 ridge models and five direct classifiers. They describe consistency of coefficient sign across folds/teacher seeds, with no independence claim, significance test, predictor selection or causal interpretation. Exact coefficients and medians are in the all-32 evidence CSV/workbook.

{stability}

## Resampled-presentation consistency

The eight pairs have {counts['exact']} exact category agreements, {counts['partial']} partial agreement and {counts['missing_primary']} pairs with a missing primary category. These are independently resampled tile presentations under different blinded codes, sometimes with disjoint patients; they measure joint presentation/name consistency, **not same-image intra-reader reliability**. The table therefore does not isolate reader variability from the variability of the selected examples.

{repeat}

Missing structured artifact responses prevent all eight pairs from establishing the formal agreement criterion under the sealed conservative amendment. That formal 0/8 does not mean the mentor disagreed on every pair or failed to recognize the tissue. The accepted return remains final; this companion does not request replacement answers or infer missing ratings. [Both literal comments for every pair](Duplicate_presentations.csv) preserve information that a category-only comparison misses.

## Historical geometry, without a naming upgrade

The original legacy p17 anchor maps to v14 p17 at cosine 0.971566; legacy p28 maps to v14 p22 at cosine 0.982369 in the canonical dictionary. Both anchors exceed 0.80 in all nine variants under direct joint one-to-one matching. The legacy vocabulary was target-inclusive/transductive, so this comparison is historical and descriptive. It does not establish new external validation. The reader's literal categories for these coordinates remain in the uniform all-32 table; **both formal named correspondence calls remain CORRESPONDENCE_NOT_EVALUABLE under NAME_GATE_FAIL**. Neither the geometry nor this companion changes feature definitions or rescues a naming gate.

{geometry}

## How this companion should be used

Use the comments to explain what selected examples look like and where a short label would be misleading. Use the numerical tables to describe the existing coordinate-level findings and their limits. Keep these two types of evidence distinct when writing biological conclusions. The raw-slide grid sensitivity assesses technical stability of assignments and abundance; it does not determine pathology purity. The 60-case/420-image whole-section packet remains GENERATED-UNREAD. This companion is a descriptive synthesis of accepted evidence, not a new confirmatory analysis or a replacement for the sealed FINAL-v14 report.
"""


def write_workbook(path: Path, tables: dict[str, list[dict]]) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, rows in tables.items():
        sheet = workbook.create_sheet(title)
        fields = list(rows[0])
        sheet.append(fields)
        for row in rows:
            sheet.append([row[key] for key in fields])
        sheet.freeze_panes = "C2"
        sheet.auto_filter.ref = sheet.dimensions
        for header in sheet[1]:
            header.fill = PatternFill("solid", fgColor="193B35")
            header.font = Font(color="FFFFFF", bold=True)
            header.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.row_dimensions[1].height = 48
        for column in sheet.columns:
            name = column[0].value
            sheet.column_dimensions[column[0].column_letter].width = 55 if "comment" in name or "description" in name else 24
            for item in column[1:]:
                item.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(item.value, str):
                    item.data_type = "s"
        for index in range(2, sheet.max_row + 1):
            sheet.row_dimensions[index].height = 45
    workbook.save(path)


def verify_outputs(tables: dict, markdown: str) -> dict:
    checks = 0
    for name, rows in tables.items():
        assert (OUT / f"{name}.csv").read_text() == csv_text(rows)
        checks += len(rows)
    assert (OUT / "Review_Interpretation.md").read_text() == markdown
    workbook = load_workbook(OUT / "Review_Interpretation.xlsx", read_only=True, data_only=False)
    assert workbook.sheetnames == list(tables)
    for name, rows in tables.items():
        values = list(workbook[name].values)
        fields = list(rows[0])
        assert list(values[0]) == fields
        assert len(values) == len(rows) + 1
        for actual, expected in zip(values[1:], rows, strict=True):
            for value, key in zip(actual, fields, strict=True):
                source = expected[key]
                if isinstance(source, float):
                    assert math.isclose(value, source, rel_tol=1e-14, abs_tol=1e-15), (name, key)
                else:
                    assert value == (None if source == "" else source), (name, key, value, source)
                checks += 1
    workbook.close()
    return {"status": "PASS", "csv_rows_reassembled": sum(map(len, tables.values())),
            "workbook_cells_and_rows_checked": checks, "all32_ids": list(range(32)),
            "association_status_counts": dict(Counter(r["status"] for r in tables["Context_associations"])),
            "literal_controlling_rows_verified": 32, "literal_return_rows_retained": 40,
            "new_models_or_tests": 0, "new_pathology_categories": 0}


def build() -> dict:
    manifest_path = OUT / "input_identity_manifest.json"
    if manifest_path.exists():
        return verify()
    data, pins = load_inputs()
    tables = assemble(data)
    markdown = render_markdown(data, tables)
    OUT.mkdir(parents=True, exist_ok=True)
    for name, rows in tables.items():
        (OUT / f"{name}.csv").write_text(csv_text(rows))
    (OUT / "Review_Interpretation.md").write_text(markdown)
    write_workbook(OUT / "Review_Interpretation.xlsx", tables)
    checked = verify_outputs(tables, markdown)
    artifacts = [identity(OUT / f"{name}.csv") for name in tables]
    artifacts += [identity(OUT / "Review_Interpretation.md"), identity(OUT / "Review_Interpretation.xlsx")]
    manifest = {"status": "DESCRIPTIVE_REVIEW_COMPANION_COMPLETE", "created_utc": datetime.now(timezone.utc).isoformat(),
                "code": identity(Path(__file__)), "scientific_inputs": pins, "artifacts": artifacts,
                "verification": checked, "name_gate_status": "NAME_GATE_FAIL",
                "scope": "Descriptive companion only; no modification of sealed analyses, no new tests, labels, imputation or cluster-wide purity estimate."}
    temporary = OUT / ".input_identity_manifest.tmp"
    temporary.write_text(json.dumps(manifest, indent=2, ensure_ascii=False, allow_nan=False) + "\n")
    os.replace(temporary, manifest_path)
    return manifest


def verify() -> dict:
    manifest = read_json(OUT / "input_identity_manifest.json")
    assert manifest["code"] == identity(Path(__file__))
    data, pins = load_inputs()
    assert manifest["scientific_inputs"] == pins
    for pin in manifest["artifacts"]:
        assert identity(Path(pin["path"])) == pin
    tables = assemble(data)
    checked = verify_outputs(tables, render_markdown(data, tables))
    assert manifest["verification"] == checked
    return manifest


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("command", nargs="?", choices=("build", "verify"), default="build")
    args = parser.parse_args()
    result = build() if args.command == "build" else verify()
    print(json.dumps({"status": result["status"], "output": str(OUT), "verification": result["verification"]}, indent=2))


if __name__ == "__main__":
    main()
