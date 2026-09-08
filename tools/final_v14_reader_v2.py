#!/usr/bin/env python3
"""Ingest the finalized v14 reader revision without changing prior analyses.

The finalized workbook changes the rubric after the existing analyses. Its
literal annotations are current descriptive evidence; the original naming gate
and all fitted representations remain historical, immutable results. Helper
HYPERLINK formulas are stored as strings and are never evaluated.
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import io
import json
import sys
from collections import Counter
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO))
from tools import final_v14_post_reader as original  # noqa: E402

handoff = original.handoff
ROOT = original.POST.with_name("e4v_post_reader_xlsx_v2")
WORKBOOK = handoff.DELIVERY_PACKAGE / "completed_review_form_v2.xlsx"
SCOPE = "FINALIZED_REVISED_RUBRIC_POST_ANALYSIS"
HELPERS = ("Column1", "Column2")
ADDED_CATEGORIES = ("desmoplastic", "benign fibrosis")
REMOVED_CATEGORY = "desmoplastic/fibrous stroma"
V2_HEADERS = handoff.FORM_COLUMNS[:2] + list(HELPERS) + handoff.FORM_COLUMNS[2:]


def require(value: bool, message: str) -> None:
    if not value:
        raise original.PostReaderError(message)


def json_value(value: Any) -> Any:
    if isinstance(value, (dt.date, dt.datetime)):
        return value.isoformat()
    return value


def parse_workbook(payload: bytes, baseline_payload: bytes,
                   expected_rows: list[dict]) -> dict:
    """Validate fixed presentation identities and retain the revised form literally."""
    handoff._safe_zip_members(payload, "finalized revision v2")
    workbook = load_workbook(io.BytesIO(payload), data_only=False, keep_links=True)
    baseline = load_workbook(io.BytesIO(baseline_payload), data_only=False, keep_links=True)
    require(workbook.sheetnames == handoff.EXPECTED_SHEETS, "Revision changed sheet inventory/order")
    require(not getattr(workbook, "_external_links", []), "External workbook links are forbidden")
    review = workbook[handoff.REVIEW_SHEET]
    require(review.max_row == 41 and review.max_column == 15, "Expected 40 rows and the 15-column v2 schema")
    headers = [cell.value for cell in review[1]]
    require(headers == V2_HEADERS, "Unexpected v2 headers or helper-column locations")
    require(len(expected_rows) == 40, "Expected exactly 40 frozen presentations")
    require(list(workbook[handoff.INSTRUCTIONS_SHEET].values) ==
            list(baseline[handoff.INSTRUCTIONS_SHEET].values), "Reader instructions changed beyond the authorized rubric revision")
    old_rubric = baseline[handoff.RUBRIC_SHEET]["A2"].value
    rubric = workbook[handoff.RUBRIC_SHEET]["A2"].value
    expected_rubric = old_rubric.replace(f"- {REMOVED_CATEGORY}\n", "- desmoplastic\n- benign fibrosis\n")
    require(rubric == expected_rubric, "Unexpected revision of the literal rubric")
    old_options = [list(row) for row in baseline[handoff.OPTIONS_SHEET].values]
    options = [list(row) for row in workbook[handoff.OPTIONS_SHEET].values]
    categories = [row[1] for row in options if len(row) > 1 and row[1] is not None]
    expected_categories = [category for category in handoff.ONTOLOGY if category != REMOVED_CATEGORY]
    expected_categories[4:4] = list(ADDED_CATEGORIES)
    require(categories == expected_categories, "Unexpected revised ontology or option order")
    require(all(len(row) == 5 for row in options), "Unexpected options-column inventory")
    for column in (0, 2, 3, 4):
        require([row[column] for row in options if row[column] is not None] ==
                [row[column] for row in old_options if row[column] is not None],
                f"Non-ontology options changed in column {column + 1}")
    # Only the two workbook interface columns differ from the response schema.
    # Formula contents are checked as text; no links, programs or formulas run.
    for sheet in workbook:
        for row in sheet:
            for cell in row:
                if cell.data_type == "f":
                    require(sheet.title == handoff.REVIEW_SHEET and cell.column == 4
                            and 2 <= cell.row <= 41, "Formula outside the declared hyperlink helper column")
                    require(cell.value == f'=HYPERLINK(C{cell.row}&B{cell.row}&".jpg", "Open")',
                            f"Unexpected helper formula at {cell.coordinate}")
    rows, helpers = [], []
    for index, frozen in enumerate(expected_rows, start=2):
        literal = {header: json_value(review.cell(index, column).value)
                   for column, header in enumerate(headers, start=1)}
        row = {field: literal[field] for field in handoff.FORM_COLUMNS}
        require(literal["Column2"] == f'=HYPERLINK(C{index}&B{index}&".jpg", "Open")',
                f"Missing or unexpected hyperlink helper at row {index}")
        require([row[field] for field in handoff.FIXED_COLUMNS] ==
                [frozen[field] for field in handoff.FIXED_COLUMNS],
                f"Frozen presentation identity changed at row {index}")
        for field in ("primary_category", "secondary_category_1", "secondary_category_2"):
            require(row[field] is None or row[field] in categories,
                    f"Response outside revised ontology at row {index}/{field}")
        selected = [row[field] for field in ("primary_category", "secondary_category_1", "secondary_category_2")
                    if row[field] is not None]
        require(len(selected) == len(set(selected)), f"Repeated category at row {index}")
        confidence = row["confidence_1_to_5"]
        require(confidence is None or (type(confidence) is int and 1 <= confidence <= 5),
                f"Invalid confidence at row {index}")
        require(row["artifact_uninterpretable"] in (None, "yes", "no"),
                f"Invalid artifact response at row {index}")
        rows.append(row)
        helpers.append({"presentation_order": row["presentation_order"], "blinded_code": row["blinded_code"],
                        **{field: literal[field] for field in HELPERS},
                        "blinded_code_hyperlink": review.cell(index, 2).hyperlink.target
                        if review.cell(index, 2).hyperlink else None})
    missing = {field: sum(row[field] in (None, "") for row in rows)
               for field in handoff.RESPONSE_COLUMNS}
    changes = [{"presentation_order": row["presentation_order"], "blinded_code": row["blinded_code"],
                "field": field, "previous": previous[field], "finalized": row[field]}
               for previous, row in zip(expected_rows, rows, strict=True)
               for field in handoff.FORM_COLUMNS if previous[field] != row[field]]
    result = {"reader_revision": 2, "review_scope": SCOPE, "rows": rows,
              "helper_columns": helpers,
              "workbook_context": {"headers": headers, "sheet_names": workbook.sheetnames,
                                   "instructions": workbook[handoff.INSTRUCTIONS_SHEET]["A2"].value,
                                   "rubric": rubric, "options": options,
                                   "structure": {sheet.title: {"state": sheet.sheet_state,
                                       "protected": sheet.protection.sheet,
                                       "freeze_panes": sheet.freeze_panes,
                                       "tables": {name: sheet.tables[name].ref for name in sheet.tables}}
                                       for sheet in workbook},
                                   "defined_names": {name: value.attr_text for name, value in workbook.defined_names.items()},
                                   "file_metadata": {"modified": json_value(workbook.properties.modified),
                                                     "last_modified_by": workbook.properties.lastModifiedBy},
                                   "metadata_policy": "File metadata is preserved but does not establish the missing reviewer, review date or blinding attestation."},
              "validation": {"presentation_identity_bijection": True,
                             "literal_response_values_preserved": True,
                             "helper_formulas_evaluated": False,
                             "missing_response_counts": missing,
                             "scientific_rating_fields_complete": all(missing[field] == 0 for field in
                                 ("primary_category", "confidence_1_to_5", "artifact_uninterpretable")),
                             "original_complete_form_validation_passed": False,
                             "original_complete_form_limitations": ["Revised ontology and two added interface columns",
                                 "No recorded status, reviewer identifier, date or blinding attestation"],
                             "confidence_counts": dict(sorted(Counter(row["confidence_1_to_5"] for row in rows).items(), key=lambda item: str(item[0]))),
                             "artifact_counts": dict(Counter(row["artifact_uninterpretable"] for row in rows)),
                             "ontology": {"previous": list(handoff.ONTOLOGY), "finalized": categories,
                                          "removed": [REMOVED_CATEGORY], "added": list(ADDED_CATEGORIES),
                                          "changed": True}},
              "response_changes": changes,
              "changed_field_counts": dict(sorted(Counter(row["field"] for row in changes).items()))}
    workbook.close()
    baseline.close()
    return result


def sources() -> tuple[dict, dict, list[dict]]:
    names_path = original.POST / "naming/naming_freeze.json"
    literal_path = original.POST / "literal_intake/literal_response_values.json"
    names = original.verify_seal(names_path)
    literal = original.verify_seal(literal_path)
    for pin in names["source_identities"].values():
        original.check_identity(pin)
    raw_pin = names["source_identities"]["raw_workbook"]
    original.check_identity(raw_pin)
    reader_path = original.PRE / "receipts/reader_package.json"
    mapping_path = original.PRE / "receipts/mappings.json"
    key_pin = original.read_json(reader_path)["artifacts"]["embargoed_occurrence_key"]
    mapping_pin = original.read_json(mapping_path)["artifacts"]["fold_to_reference"]
    for pin in (key_pin, mapping_pin):
        original.check_identity(pin)
    pins = [original.identity(path) for path in (names_path, literal_path, reader_path, mapping_path,
                                                Path(original.__file__), Path(handoff.__file__))]
    pins += [raw_pin, key_pin, mapping_pin]
    return {"names": names, "literal": literal, "raw": raw_pin,
            "key": key_pin, "mapping": mapping_pin,
            "duplicate_preflight": original.read_json(reader_path)["duplicate_preflight_status"]}, names, pins


def revised_summary(parsed: dict, historical: dict, key: pd.DataFrame,
                    mapping: pd.DataFrame, duplicate_preflight: str) -> dict:
    replay = original.derive_naming(parsed["rows"], key, mapping, duplicate_preflight)
    controlling = []
    old_reads = {row["prototype_id"]: row for row in historical["controlling_reads"]}
    for record in replay["controlling_reads"]:
        controlling.append({**record, "response_eligibility_descriptive": record["eligible_for_named_ref"],
                            "response_exclusion_reasons": record["exclusion_reasons"],
                            "eligible_for_named_ref": old_reads[record["prototype_id"]]["eligible_for_named_ref"],
                            "exclusion_reasons": old_reads[record["prototype_id"]]["exclusion_reasons"]})
    return {"reader_revision": 2, "review_scope": SCOPE,
            "status": "FINALIZED_REVISED_READER_EVIDENCE_DESCRIPTIVE",
            "revision_reason": "User designated completed_review_form_v2.xlsx finalized and requested report/paper updates.",
            "name_gate_status": historical["name_gate_status"],
            "historical_name_gate": {"name_gate_status": historical["name_gate_status"],
                                     "named_ref_count": historical["named_ref_count"],
                                     "repeatability": historical["repeatability"],
                                     "name_gate_conditions": historical["name_gate_conditions"]},
            "response_only_gate_status": replay["name_gate_status"],
            "response_only_name_gate_conditions": replay["name_gate_conditions"],
            "descriptive_eligible_ref": replay["NAMED_REF"],
            "descriptive_eligible_ref_count": replay["named_ref_count"],
            "descriptive_eligible_oof": replay["NAMED_OOF"],
            "descriptive_fold_coverage": replay["fold_coverage"],
            "ALL32": historical["ALL32"], "ATTRIBUTABLE_REF": historical["ATTRIBUTABLE_REF"],
            "NAMED_REF": historical["NAMED_REF"], "NAMED_OOF": historical["NAMED_OOF"],
            "named_ref_count": historical["named_ref_count"],
            "controlling_reads": controlling,
            "controlling_primary_category_counts": replay["controlling_primary_category_counts"],
            "repeatability": {**replay["repeatability"],
                              "scope": "Descriptive consistency of independently resampled presentations under the finalized revised rubric",
                              "missing_artifact_policy": "Both literal artifact responses must be no; no response is imputed."},
            "revised_rule_replay": replay,
            "missing_metadata_counts": {field: parsed["validation"]["missing_response_counts"][field]
                                        for field in ("review_status", "reviewer_id", "review_date", "blinding_attestation")},
            "scientific_rating_fields_complete": parsed["validation"]["scientific_rating_fields_complete"],
            "ontology": parsed["validation"]["ontology"],
            "session_timing": "Revision received after existing analyses; completion date/session timing are unrecorded in the workbook.",
            "reader_authority": "Finalized revision of the existing mentor return per user; workbook reviewer identifier and blinding attestation are unrecorded.",
            "interpretation_scope": "Current literal tissue annotations and response-rule replay are descriptive. The changed rubric cannot retroactively change the historical gate, authorize named models, or refit existing ALL32 analyses."}


def mentor_rows(summary: dict) -> list[dict]:
    rows = []
    for item in summary["controlling_reads"]:
        response = item["response"]
        rows.append({"Concept": item["prototype_id"], "Controlling code": item["controlling_code"],
                     "Primary category": response["primary_category"], "Secondary 1": response["secondary_category_1"],
                     "Secondary 2": response["secondary_category_2"], "Mentor comment": response["free_text_description"],
                     "Confidence": response["confidence_1_to_5"], "Artifact flag": response["artifact_uninterpretable"],
                     "Named-set eligible": item["eligible_for_named_ref"],
                     "Historical named-set eligible": item["eligible_for_named_ref"],
                     "V2 response eligibility": item["response_eligibility_descriptive"],
                     "V2 response exclusion reasons": "; ".join(item["response_exclusion_reasons"]),
                     "Reader revision": 2})
    return rows


def csv_bytes(rows: list[dict]) -> bytes:
    buffer = io.StringIO(newline="")
    writer = csv.DictWriter(buffer, fieldnames=list(rows[0]), lineterminator="\n")
    writer.writeheader()
    writer.writerows(rows)
    return buffer.getvalue().encode("utf-8")


def seal_raw_return(source: Path, output: Path) -> dict:
    payload = source.read_bytes()
    raw = output / "raw_return/completed_review_form_v2.xlsx"
    original.write_once(raw, payload)
    pin = original.identity(raw)
    original.publish(output / "raw_return/receipt.json", {
        "status": "RAW_FINALIZED_REVISION_PRESERVED_BEFORE_REVISION_PROCESSING",
        "reader_revision": 2, "artifact": pin, "source": original.identity(source),
        "previous_analysis_key_access": "Already occurred in the historical campaign; no new pre-unblinding claim."})
    return pin


def ingest(source: Path = WORKBOOK, output: Path = ROOT) -> dict:
    if (output / "receipt.json").exists():
        receipt = verify(output)
        original.check_identity(receipt["raw_workbook"], source)
        return receipt
    raw_pin = seal_raw_return(source, output)
    context, historical, input_pins = sources()
    amendment = {"status": "USER_AUTHORIZED_FINAL_REVIEW_REVISION_AFTER_EXISTING_ANALYSES",
                 "reader_revision": 2, "review_scope": SCOPE,
                 "authorizing_instruction": "with the new completed_review_form_v2 (finalized), update final_v14 and the paper as well.",
                 "raw_workbook": raw_pin, "historical_inputs": input_pins,
                 "historical_gate_preserved": historical["name_gate_status"],
                 "prior_key_access_and_completed_analyses": True,
                 "revision_session_timing_established": False,
                 "revision_scope": "Accept current literal annotations and revised rubric as finalized descriptive evidence; preserve earlier receipts, all numerical analyses, and the original naming gate."}
    original.publish(output / "governance_amendment.json", amendment)
    parsed = parse_workbook(Path(raw_pin["path"]).read_bytes(), Path(context["raw"]["path"]).read_bytes(),
                            context["literal"]["rows"])
    summary = revised_summary(parsed, historical, pd.read_csv(context["key"]["path"]),
                              pd.read_csv(context["mapping"]["path"]), context["duplicate_preflight"])
    artifacts = []
    for relative, value in (("literal_intake/literal_response_values.json", parsed),
                            ("naming/naming_freeze.json", summary)):
        artifacts.append(original.publish(output / relative, value))
    for relative, rows in (("literal_intake/completed_review_form_v2.csv", parsed["rows"]),
                           ("tables/Mentor_annotations.csv", mentor_rows(summary))):
        original.write_once(output / relative, csv_bytes(rows))
        artifacts.append(original.identity(output / relative))
    artifacts += [original.identity(output / name) for name in ("governance_amendment.json", "raw_return/receipt.json")]
    receipt = {"status": "FINALIZED_READER_V2_INGESTION_COMPLETE", "reader_revision": 2, "review_scope": SCOPE,
               "raw_workbook": raw_pin, "code": original.identity(Path(__file__)),
               "scientific_inputs": input_pins, "artifacts": artifacts,
               "historical_name_gate_status": historical["name_gate_status"],
               "response_only_gate_status": summary["response_only_gate_status"],
               "descriptive_eligible_ref_count": summary["descriptive_eligible_ref_count"],
               "descriptive_repeatability": {key: summary["repeatability"][key] for key in
                                             ("numerator", "denominator", "exact", "partial", "ci95_clopper_pearson")},
               "existing_numerical_analyses_modified": False}
    original.publish(output / "receipt.json", receipt)
    return verify(output)


def verify(output: Path = ROOT) -> dict:
    receipt = original.verify_seal(output / "receipt.json")
    require(receipt["code"] == original.identity(Path(__file__)), "Reader successor code identity drift")
    original.check_identity(receipt["raw_workbook"])
    for pin in receipt["scientific_inputs"] + receipt["artifacts"]:
        original.check_identity(pin)
        if Path(pin["path"]).suffix == ".json" and Path(pin["path"]).is_relative_to(output):
            original.verify_seal(Path(pin["path"]))
    amendment = original.verify_seal(output / "governance_amendment.json")
    require(amendment["status"] == "USER_AUTHORIZED_FINAL_REVIEW_REVISION_AFTER_EXISTING_ANALYSES",
            "Revision timing/provenance is not preserved")
    context, historical, inputs = sources()
    require(inputs == receipt["scientific_inputs"], "Historical input inventory drift")
    parsed = parse_workbook(Path(receipt["raw_workbook"]["path"]).read_bytes(),
                            Path(context["raw"]["path"]).read_bytes(), context["literal"]["rows"])
    # Normalize only JSON dictionary key serialization (integer confidence counts).
    parsed = json.loads(json.dumps(parsed))
    require(parsed == original.verify_seal(output / "literal_intake/literal_response_values.json"),
            "Literal revision replay differs")
    summary = revised_summary(parsed, historical, pd.read_csv(context["key"]["path"]),
                              pd.read_csv(context["mapping"]["path"]), context["duplicate_preflight"])
    require(summary == original.verify_seal(output / "naming/naming_freeze.json"), "Revised-response replay differs")
    require(csv_bytes(parsed["rows"]) == (output / "literal_intake/completed_review_form_v2.csv").read_bytes(),
            "Literal response CSV differs")
    require(csv_bytes(mentor_rows(summary)) == (output / "tables/Mentor_annotations.csv").read_bytes(),
            "Mentor annotation table differs")
    require(receipt["historical_name_gate_status"] == historical["name_gate_status"] and
            receipt["response_only_gate_status"] == summary["response_only_gate_status"] and
            receipt["descriptive_eligible_ref_count"] == summary["descriptive_eligible_ref_count"],
            "Receipt scoring summary differs")
    require(receipt["descriptive_repeatability"] == {key: summary["repeatability"][key]
            for key in ("numerator", "denominator", "exact", "partial", "ci95_clopper_pearson")},
            "Receipt repeatability summary differs")
    return receipt


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("stage", choices=("ingest", "verify"))
    parser.add_argument("--workbook", type=Path, default=WORKBOOK)
    parser.add_argument("--output", type=Path, default=ROOT)
    args = parser.parse_args()
    result = ingest(args.workbook, args.output) if args.stage == "ingest" else verify(args.output)
    print(json.dumps({key: value for key, value in result.items()
                      if key not in ("artifacts", "scientific_inputs")}, indent=2))


if __name__ == "__main__":
    main()
