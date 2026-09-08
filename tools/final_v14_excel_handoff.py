#!/usr/bin/env python3
"""Build, verify, and ingest the FINAL-v14 Excel reader handoff.

The original ``reviews/v14`` CSV handoff is frozen and hash-sealed.  This
tool creates a separately versioned, public-only Excel interface derived from
that package without changing any frozen FINAL-v14 computation or reader
artifact.  It also provides the blinded, fail-closed ingestion step for a
returned workbook; ingestion seals the raw XLSX and exports the canonical CSV
before any embargoed key is opened.
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import hashlib
import io
import json
import os
import re
import shutil
import tempfile
import zipfile
from collections.abc import Iterable, Mapping, Sequence
from pathlib import Path, PurePosixPath
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

REPO_ROOT = Path(__file__).resolve().parents[1]
SOURCE_REVIEW_ROOT = REPO_ROOT / "reviews/v14"
SOURCE_PACKAGE = SOURCE_REVIEW_ROOT / "FOR_PATHOLOGIST"
DELIVERY_ROOT = REPO_ROOT / "reviews/v14_xlsx"
DELIVERY_PACKAGE = DELIVERY_ROOT / "FOR_PATHOLOGIST"
AUDIT_ROOT = (
    REPO_ROOT
    / "reports/reruns/final_v14_additions_20260903/e4v_pre_reader_xlsx_handoff"
)
HANDOFF_RECEIPT = AUDIT_ROOT / "handoff_receipt.json"
POST_READER_PARENT = (
    REPO_ROOT
    / "reports/reruns/final_v14_additions_20260903/e4v_post_reader_xlsx"
)
POST_READER_OUTPUT = POST_READER_PARENT / "human_reader_1"

PINNED_SOURCE_MANIFEST_SHA256 = (
    "1b6f7af9f953649b81dc70756cd6981b69bd6f45fe66fdea80c40570f4fe3830"
)
PINNED_SOURCE_FORM_SHA256 = (
    "3aed5eb8dfa5336d3c256833089ee1e415fab2f427d8925779573983ec9dca03"
)
PINNED_INDEPENDENT_AUDIT_SHA256 = (
    "744cc8e50f8e3b42995a2fe0fafae5fd0e0a511945b2e63c989de16e2c0781a8"
)
INDEPENDENT_AUDIT_RECEIPT = (
    REPO_ROOT
    / "reports/reruns/final_v14_additions_20260903/"
    "e4v_pre_reader_independent_audit/audit_receipt.json"
)

SCHEMA_VERSION = 1
INTERFACE_REVISION = 1
PACKAGE_NAME = "FINAL-v14 blinded morphology naming session — Excel interface"
WORKBOOK_NAME = "review_form.xlsx"
COMPLETED_WORKBOOK_NAME = "completed_review_form.xlsx"
COMPLETED_CSV_NAME = "completed_review_form.csv"
REVIEW_SHEET = "Review Form"
INSTRUCTIONS_SHEET = "Instructions"
RUBRIC_SHEET = "Rubric"
OPTIONS_SHEET = "_Options"
TABLE_NAME = "FINALv14MorphologyReview"
SHEET_PASSWORD = "v14-interface-guard"

FORM_COLUMNS = [
    "presentation_order",
    "blinded_code",
    "n_tiles",
    "review_status",
    "primary_category",
    "secondary_category_1",
    "secondary_category_2",
    "free_text_description",
    "confidence_1_to_5",
    "artifact_uninterpretable",
    "reviewer_id",
    "review_date",
    "blinding_attestation",
]
FIXED_COLUMNS = FORM_COLUMNS[:3]
RESPONSE_COLUMNS = FORM_COLUMNS[3:]
ONTOLOGY = [
    "malignant gland-forming epithelium/gland–lumen",
    "malignant solid or poorly differentiated epithelium",
    "extracellular mucin/mucinous pattern",
    "normal or benign colonic epithelium",
    "desmoplastic/fibrous stroma",
    "smooth muscle",
    "lymphoid/inflammatory tissue",
    "necrosis/debris",
    "adipose",
    "blood/vessel",
    "mixed/other interpretable",
]
STATUS_OPTIONS = ["complete"]
CONFIDENCE_OPTIONS = [1, 2, 3, 4, 5]
ARTIFACT_OPTIONS = ["yes", "no"]
ATTESTATION_OPTIONS = ["confirmed_no_key_access"]
EXPECTED_SHEETS = [REVIEW_SHEET, INSTRUCTIONS_SHEET, RUBRIC_SHEET, OPTIONS_SHEET]

FORBIDDEN_PUBLIC_TERMS = {
    "patient_id",
    "slide_id",
    "prototype_id",
    "subcohort",
    "target_label",
    "kras",
    "outer_fold",
    "centroid_distance",
    "model_score",
    "association_statistic",
    "sr1482",
    "sr386",
    "tcga-coad",
    "tcga-read",
}

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
EDITABLE_FILL = PatternFill("solid", fgColor="FFF2CC")
FIXED_FILL = PatternFill("solid", fgColor="D9EAF7")
WARNING_FILL = PatternFill("solid", fgColor="F4CCCC")
THIN_GRAY = Side(style="thin", color="B7B7B7")
CELL_BORDER = Border(left=THIN_GRAY, right=THIN_GRAY, top=THIN_GRAY, bottom=THIN_GRAY)


class HandoffError(RuntimeError):
    """A fail-closed handoff or completed-workbook validation error."""


def utc_now() -> str:
    return dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds")


def sha256_bytes(payload: bytes) -> str:
    return hashlib.sha256(payload).hexdigest()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while chunk := handle.read(8 * 1024 * 1024):
            digest.update(chunk)
    return digest.hexdigest()


def json_bytes(value: Any) -> bytes:
    return (
        json.dumps(value, indent=2, sort_keys=True, ensure_ascii=False, allow_nan=False)
        + "\n"
    ).encode("utf-8")


def identity(path: Path, *, relative_to: Path | None = None) -> dict[str, Any]:
    resolved = path.resolve(strict=True)
    return {
        "path": (
            resolved.relative_to(relative_to.resolve(strict=True)).as_posix()
            if relative_to is not None
            else str(resolved)
        ),
        "size_bytes": resolved.stat().st_size,
        "sha256": sha256_file(resolved),
    }


def tree_inventory(root: Path) -> list[dict[str, Any]]:
    if root.is_symlink() or not root.is_dir():
        raise HandoffError(f"artifact tree is missing or is a symlink: {root}")
    if any(path.is_symlink() for path in root.rglob("*")):
        raise HandoffError(f"artifact tree contains a symlink: {root}")
    return [
        identity(path, relative_to=root)
        for path in sorted(root.rglob("*"))
        if path.is_file()
    ]


def write_new(path: Path, payload: bytes, *, mode: int = 0o600) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        descriptor = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_EXCL, mode)
    except FileExistsError as exc:
        raise HandoffError(f"refusing to overwrite existing artifact: {path}") from exc
    with os.fdopen(descriptor, "wb") as handle:
        handle.write(payload)
        handle.flush()
        os.fsync(handle.fileno())


def _safe_zip_members(payload: bytes, label: str) -> list[str]:
    try:
        archive = zipfile.ZipFile(io.BytesIO(payload), "r")
    except (zipfile.BadZipFile, OSError) as exc:
        raise HandoffError(f"{label}: not a valid XLSX ZIP archive") from exc
    with archive:
        infos = archive.infolist()
        names = [item.filename for item in infos]
        if len(names) != len(set(names)):
            raise HandoffError(f"{label}: duplicate ZIP member")
        if len(names) > 250:
            raise HandoffError(f"{label}: excessive ZIP member count")
        total = 0
        for item in infos:
            member = PurePosixPath(item.filename)
            if member.is_absolute() or ".." in member.parts or "\\" in item.filename:
                raise HandoffError(f"{label}: unsafe ZIP path {item.filename!r}")
            if item.file_size > 8 * 1024 * 1024:
                raise HandoffError(f"{label}: oversized ZIP member {item.filename!r}")
            total += item.file_size
        if total > 32 * 1024 * 1024:
            raise HandoffError(f"{label}: excessive uncompressed workbook size")
        dangerous_prefixes = (
            "customui/",
            "customxml/",
            "xl/activex/",
            "xl/connections",
            "xl/ctrlprops/",
            "xl/dialogsheets/",
            "xl/embeddings/",
            "xl/externallinks/",
            "xl/macrosheets/",
            "xl/model/",
            "xl/pivotcache/",
            "xl/querytables/",
            "xl/slicercaches/",
            "xl/timelines/",
            "xl/webextensions/",
        )
        forbidden = [
            name
            for name in names
            if name.casefold().endswith((".bin", ".ole"))
            or name.casefold().startswith(dangerous_prefixes)
        ]
        if forbidden:
            raise HandoffError(f"{label}: forbidden active/external content: {forbidden}")
        if "[Content_Types].xml" not in names:
            raise HandoffError(f"{label}: XLSX content-types declaration is missing")
        content_types = archive.read("[Content_Types].xml").lower()
        dangerous_content_types = (
            b"activex",
            b"connections",
            b"customxml",
            b"externallink",
            b"macroenabled",
            b"oleobject",
            b"querytable",
            b"vbaproject",
            b"webextension",
        )
        if any(term in content_types for term in dangerous_content_types):
            raise HandoffError(f"{label}: forbidden OOXML content type")
    return names


def _validate_relative_hyperlink_relationships(
    payload: bytes, expected_targets: set[str], label: str
) -> None:
    import xml.etree.ElementTree as element_tree

    observed: set[str] = set()
    with zipfile.ZipFile(io.BytesIO(payload), "r") as archive:
        for name in archive.namelist():
            if not name.endswith(".rels"):
                continue
            try:
                root = element_tree.fromstring(archive.read(name))
            except element_tree.ParseError as exc:
                raise HandoffError(f"{label}: malformed relationship XML {name}") from exc
            for relation in root:
                if relation.attrib.get("TargetMode") != "External":
                    continue
                relation_type = relation.attrib.get("Type", "")
                target = relation.attrib.get("Target", "")
                if not relation_type.endswith("/hyperlink") or target not in expected_targets:
                    raise HandoffError(
                        f"{label}: forbidden external relationship {target!r}"
                    )
                observed.add(target)
    if observed != expected_targets:
        raise HandoffError(f"{label}: montage hyperlink relationship census differs")


def _read_source_form(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle)
        if reader.fieldnames != FORM_COLUMNS:
            raise HandoffError("source review form schema drifted")
        rows = list(reader)
    if len(rows) != 40:
        raise HandoffError("source review form is not exactly 40 rows")
    for offset, row in enumerate(rows, start=1):
        if int(row["presentation_order"]) != offset:
            raise HandoffError("source presentation order drifted")
        if not re.fullmatch(r"[A-Z2-7]{6}", row["blinded_code"]):
            raise HandoffError("source blinded code is malformed")
        n_tiles = int(row["n_tiles"])
        if not 1 <= n_tiles <= 12:
            raise HandoffError("source montage tile count is invalid")
        if any(row[column] != "" for column in RESPONSE_COLUMNS):
            raise HandoffError("source review form is not a blank master")
    codes = [row["blinded_code"] for row in rows]
    if len(codes) != len(set(codes)):
        raise HandoffError("source blinded codes are not unique")
    return rows


def _ontology_from_rubric(text: str) -> list[str]:
    categories = [line[2:] for line in text.splitlines() if line.startswith("- ")]
    if categories != ONTOLOGY:
        raise HandoffError("source rubric ontology drifted")
    return categories


def validate_source_package() -> dict[str, Any]:
    if SOURCE_PACKAGE.is_symlink() or not SOURCE_PACKAGE.is_dir():
        raise HandoffError("frozen source handoff is missing or is a symlink")
    manifest_path = SOURCE_PACKAGE / "HANDOFF_MANIFEST.json"
    manifest_sha = sha256_file(manifest_path)
    if manifest_sha != PINNED_SOURCE_MANIFEST_SHA256:
        raise HandoffError("frozen source handoff manifest identity drifted")
    sidecar = (SOURCE_PACKAGE / "HANDOFF_MANIFEST.sha256").read_text(
        encoding="ascii"
    )
    if sidecar != f"{manifest_sha}  HANDOFF_MANIFEST.json\n":
        raise HandoffError("frozen source handoff manifest sidecar failed")
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    records = manifest.get("files")
    if not isinstance(records, list):
        raise HandoffError("frozen source manifest has no file inventory")
    expected_paths: set[str] = set()
    for record in records:
        relative = PurePosixPath(str(record.get("path", "")))
        if relative.is_absolute() or ".." in relative.parts:
            raise HandoffError("frozen source manifest contains an unsafe path")
        relative_text = relative.as_posix()
        if relative_text in expected_paths:
            raise HandoffError("frozen source manifest contains duplicate paths")
        expected_paths.add(relative_text)
        path = SOURCE_PACKAGE / relative_text
        if path.is_symlink() or not path.is_file():
            raise HandoffError(f"frozen source artifact is missing: {relative_text}")
        if path.stat().st_size != int(record.get("size_bytes", -1)):
            raise HandoffError(f"frozen source size mismatch: {relative_text}")
        if sha256_file(path) != str(record.get("sha256", "")):
            raise HandoffError(f"frozen source hash mismatch: {relative_text}")
    observed_paths = {
        path.relative_to(SOURCE_PACKAGE).as_posix()
        for path in SOURCE_PACKAGE.rglob("*")
        if path.is_file()
    }
    if observed_paths != expected_paths | {
        "HANDOFF_MANIFEST.json",
        "HANDOFF_MANIFEST.sha256",
    }:
        raise HandoffError("frozen source package file census drifted")
    if any(path.is_symlink() for path in SOURCE_PACKAGE.rglob("*")):
        raise HandoffError("frozen source package contains a symlink")

    source_form = SOURCE_PACKAGE / "review_form.csv"
    if sha256_file(source_form) != PINNED_SOURCE_FORM_SHA256:
        raise HandoffError("frozen source blank form identity drifted")
    rows = _read_source_form(source_form)
    rubric = (SOURCE_PACKAGE / "RUBRIC.md").read_text(encoding="utf-8")
    _ontology_from_rubric(rubric)
    codes = {row["blinded_code"] for row in rows}
    montage_paths = sorted((SOURCE_PACKAGE / "montages").glob("*.jpg"))
    if {path.stem for path in montage_paths} != codes:
        raise HandoffError("frozen source montage census differs from blank form")

    status = json.loads(
        (SOURCE_REVIEW_ROOT / "PRE_READER_STATUS.json").read_text(encoding="utf-8")
    )
    if status.get("status") != "READY_FOR_PATHOLOGIST":
        raise HandoffError("frozen source handoff is not reader-ready")
    if sha256_file(INDEPENDENT_AUDIT_RECEIPT) != PINNED_INDEPENDENT_AUDIT_SHA256:
        raise HandoffError("independent FINAL-v14 audit receipt identity drifted")
    independent = json.loads(INDEPENDENT_AUDIT_RECEIPT.read_text(encoding="utf-8"))
    if independent.get("status") != "PASS" or not independent.get("handoff_ready"):
        raise HandoffError("independent FINAL-v14 audit is not passing")
    return {
        "manifest": manifest,
        "manifest_sha256": manifest_sha,
        "form_sha256": PINNED_SOURCE_FORM_SHA256,
        "rows": rows,
        "rubric": rubric,
        "salt": (SOURCE_PACKAGE / "PUBLIC_SALT_SHA256.txt").read_text(
            encoding="ascii"
        ),
        "independent_audit_sha256": PINNED_INDEPENDENT_AUDIT_SHA256,
    }


def excel_instructions() -> str:
    return """# Blinded morphology naming session — Excel form

Use this entire extracted folder locally so the montage links continue to work. Open `review_form.xlsx` and immediately save a copy named `completed_review_form.xlsx`; enter responses only in that copy. In the `Review Form` worksheet, click a blinded code to open its montage. Review all 40 montages once, in row order, during one approximately two-hour session. Each montage contains up to twelve 256-pixel tissue fields on a fixed 4-by-3 canvas; unused cells are white. The nominal field width is approximately 128 micrometres.

Yellow cells are editable. Use the dropdown arrows for review status, primary and secondary categories, confidence, artifact/uninterpretable, and blinding attestation. For every blinded code, select `complete`, exactly one primary category, confidence from 1 through 5, and artifact/uninterpretable as `yes` or `no`. You may select up to two distinct secondary categories; neither may repeat the primary category. Type a short morphology-only description, reviewer identifier, and review date. Use the same reviewer identifier and date in every row; Excel fill-down is permitted. Enter dates as `YYYY-MM-DD`. Do not describe what a pattern might predict.

Red highlighting indicates a missing status, an incomplete completed row, a repeated primary/secondary category, or an inconsistent reviewer identifier/date and must be resolved before return. Begin the morphology description and reviewer identifier with a letter or number; do not enter formulas. Do not sort the table, rename worksheets, insert or delete rows or columns, change blinded codes, or edit blue cells. Do not access material outside this folder.

Enter `confirmed_no_key_access` through the attestation dropdown in every row to confirm that no hidden key or outcome information was accessed before the session was completed. Return only `completed_review_form.xlsx` to the study coordinator.
"""


def _write_text_sheet(sheet: Any, text: str, title: str) -> None:
    sheet.sheet_view.showGridLines = False
    sheet["A1"] = title
    sheet["A1"].font = Font(size=16, bold=True, color="1F4E78")
    sheet["A2"] = text
    sheet["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    sheet.column_dimensions["A"].width = 120
    sheet.row_dimensions[2].height = 360
    sheet.protection.set_password(SHEET_PASSWORD)
    sheet.protection.sheet = True


def _add_named_list(
    workbook: Workbook,
    sheet: Any,
    *,
    column: int,
    name: str,
    values: Sequence[Any],
) -> None:
    for row_index, value in enumerate(values, start=1):
        sheet.cell(row=row_index, column=column, value=value)
    letter = sheet.cell(row=1, column=column).column_letter
    reference = f"'{OPTIONS_SHEET}'!${letter}$1:${letter}${len(values)}"
    workbook.defined_names.add(DefinedName(name, attr_text=reference))


def _list_validation(
    formula: str,
    *,
    prompt_title: str,
    prompt: str,
    allow_blank: bool,
) -> DataValidation:
    validation = DataValidation(
        type="list",
        formula1=f"={formula}",
        allow_blank=allow_blank,
        showDropDown=False,
        showErrorMessage=True,
        showInputMessage=True,
        errorStyle="stop",
        errorTitle="Invalid selection",
        error="Select a value from the dropdown list.",
        promptTitle=prompt_title,
        prompt=prompt,
    )
    return validation


def _canonicalize_xlsx(raw: bytes, source_sealed_utc: str) -> bytes:
    import xml.etree.ElementTree as element_tree

    sealed = dt.datetime.fromisoformat(source_sealed_utc)
    if sealed.tzinfo is None:
        raise HandoffError("workbook canonicalization requires a timezone-aware seal")
    fixed_timestamp = sealed.astimezone(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    source = zipfile.ZipFile(io.BytesIO(raw), "r")
    output = io.BytesIO()
    with source, zipfile.ZipFile(
        output, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
    ) as destination:
        for name in sorted(source.namelist()):
            member_payload = source.read(name)
            if name == "docProps/core.xml":
                try:
                    core = element_tree.fromstring(member_payload)
                except element_tree.ParseError as exc:
                    raise HandoffError("workbook core properties XML is malformed") from exc
                terms = "http://purl.org/dc/terms/"
                for field in ("created", "modified"):
                    node = core.find(f"{{{terms}}}{field}")
                    if node is None:
                        raise HandoffError(
                            f"workbook core properties omit dcterms:{field}"
                        )
                    node.text = fixed_timestamp
                member_payload = element_tree.tostring(
                    core, encoding="utf-8", xml_declaration=True
                )
            info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.create_system = 3
            info.external_attr = 0o600 << 16
            destination.writestr(
                info,
                member_payload,
                compress_type=zipfile.ZIP_DEFLATED,
                compresslevel=9,
            )
    return output.getvalue()


def build_workbook(
    rows: Sequence[Mapping[str, str]], rubric: str, source_sealed_utc: str
) -> bytes:
    workbook = Workbook()
    workbook.iso_dates = True
    review = workbook.active
    review.title = REVIEW_SHEET
    instructions = workbook.create_sheet(INSTRUCTIONS_SHEET)
    rubric_sheet = workbook.create_sheet(RUBRIC_SHEET)
    options = workbook.create_sheet(OPTIONS_SHEET)

    sealed = dt.datetime.fromisoformat(source_sealed_utc).replace(tzinfo=None)
    workbook.properties.creator = "OceanPath"
    workbook.properties.lastModifiedBy = "OceanPath"
    workbook.properties.title = "FINAL-v14 blinded morphology review form"
    workbook.properties.subject = "Blinded morphology naming session"
    workbook.properties.description = "Excel interface revision 1"
    workbook.properties.created = sealed
    workbook.properties.modified = sealed
    workbook.security.lockStructure = True
    workbook.security.set_workbook_password(SHEET_PASSWORD)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"

    for column_index, header in enumerate(FORM_COLUMNS, start=1):
        cell = review.cell(row=1, column=column_index, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = CELL_BORDER
        cell.protection = Protection(locked=True)
    review.row_dimensions[1].height = 42

    for row_index, source_row in enumerate(rows, start=2):
        values: list[Any] = [
            int(source_row["presentation_order"]),
            source_row["blinded_code"],
            int(source_row["n_tiles"]),
            *("" for _ in RESPONSE_COLUMNS),
        ]
        for column_index, value in enumerate(values, start=1):
            cell = review.cell(row=row_index, column=column_index, value=value)
            editable = column_index >= 4
            cell.fill = EDITABLE_FILL if editable else FIXED_FILL
            cell.border = CELL_BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.protection = Protection(locked=not editable)
        code = source_row["blinded_code"]
        review.cell(row=row_index, column=2).hyperlink = f"montages/{code}.jpg"
        review.cell(row=row_index, column=2).style = "Hyperlink"
        review.cell(row=row_index, column=2).fill = FIXED_FILL
        review.cell(row=row_index, column=2).protection = Protection(locked=True)
        review.cell(row=row_index, column=12).number_format = "yyyy-mm-dd"
        review.row_dimensions[row_index].height = 48

    widths = {
        "A": 13,
        "B": 14,
        "C": 9,
        "D": 15,
        "E": 47,
        "F": 47,
        "G": 47,
        "H": 54,
        "I": 15,
        "J": 20,
        "K": 20,
        "L": 14,
        "M": 31,
    }
    for column, width in widths.items():
        review.column_dimensions[column].width = width
    review.freeze_panes = "D2"
    review.sheet_view.showGridLines = False
    review.sheet_view.zoomScale = 80
    review.sheet_view.zoomScaleNormal = 80
    review.auto_filter.ref = None
    review.sheet_properties.pageSetUpPr.fitToPage = True
    review.page_setup.fitToWidth = 1
    review.page_setup.fitToHeight = 0
    review.print_title_rows = "1:1"

    table = Table(displayName=TABLE_NAME, ref="A1:M41")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    review.add_table(table)

    _add_named_list(
        workbook, options, column=1, name="_review_status", values=STATUS_OPTIONS
    )
    _add_named_list(
        workbook, options, column=2, name="_morphology_categories", values=ONTOLOGY
    )
    _add_named_list(
        workbook, options, column=3, name="_confidence", values=CONFIDENCE_OPTIONS
    )
    _add_named_list(
        workbook, options, column=4, name="_artifact", values=ARTIFACT_OPTIONS
    )
    _add_named_list(
        workbook, options, column=5, name="_attestation", values=ATTESTATION_OPTIONS
    )
    options.sheet_state = "veryHidden"
    options.protection.set_password(SHEET_PASSWORD)
    options.protection.sheet = True

    validations = [
        (
            _list_validation(
                "_review_status",
                prompt_title="Completion status",
                prompt="Select complete after reviewing this montage.",
                allow_blank=False,
            ),
            "D2:D41",
        ),
        (
            _list_validation(
                "_morphology_categories",
                prompt_title="Primary morphology",
                prompt="Select exactly one primary category.",
                allow_blank=False,
            ),
            "E2:E41",
        ),
        (
            _list_validation(
                "_morphology_categories",
                prompt_title="Optional secondary morphology",
                prompt="Select a distinct secondary category or leave blank.",
                allow_blank=True,
            ),
            "F2:F41",
        ),
        (
            _list_validation(
                "_morphology_categories",
                prompt_title="Optional secondary morphology",
                prompt="Select a distinct secondary category or leave blank.",
                allow_blank=True,
            ),
            "G2:G41",
        ),
        (
            _list_validation(
                "_confidence",
                prompt_title="Confidence",
                prompt="Select an integer from 1 (very low) to 5 (very high).",
                allow_blank=False,
            ),
            "I2:I41",
        ),
        (
            _list_validation(
                "_artifact",
                prompt_title="Artifact/uninterpretable",
                prompt="Select yes if artifact makes interpretation unreliable.",
                allow_blank=False,
            ),
            "J2:J41",
        ),
        (
            DataValidation(
                type="date",
                operator="between",
                formula1="DATE(2000,1,1)",
                formula2="DATE(2100,12,31)",
                allow_blank=False,
                showErrorMessage=True,
                showInputMessage=True,
                errorStyle="stop",
                errorTitle="Invalid date",
                error="Enter a valid date as YYYY-MM-DD.",
                promptTitle="Review date",
                prompt="Enter the date as YYYY-MM-DD.",
            ),
            "L2:L41",
        ),
        (
            _list_validation(
                "_attestation",
                prompt_title="Blinding attestation",
                prompt="Confirm that no hidden key or outcome information was accessed.",
                allow_blank=False,
            ),
            "M2:M41",
        ),
    ]
    for validation, cell_range in validations:
        review.add_data_validation(validation)
        validation.add(cell_range)

    duplicate_rule = FormulaRule(
        formula=[
            'OR(AND($F2<>"",$F2=$E2),AND($G2<>"",$G2=$E2),'
            'AND($F2<>"",$F2=$G2))'
        ],
        fill=WARNING_FILL,
    )
    incomplete_rule = FormulaRule(
        formula=[
            'AND($D2="complete",OR($E2="",$H2="",$I2="",$J2="",'
            '$K2="",$L2="",$M2=""))'
        ],
        fill=WARNING_FILL,
    )
    missing_status_rule = FormulaRule(
        formula=['$D2<>"complete"'],
        fill=WARNING_FILL,
    )
    reviewer_mismatch_rule = FormulaRule(
        formula=['AND($K2<>"",$K$2<>"",$K2<>$K$2)'],
        fill=WARNING_FILL,
    )
    date_mismatch_rule = FormulaRule(
        formula=['AND($L2<>"",$L$2<>"",$L2<>$L$2)'],
        fill=WARNING_FILL,
    )
    review.conditional_formatting.add("E2:G41", duplicate_rule)
    review.conditional_formatting.add("D2:M41", incomplete_rule)
    review.conditional_formatting.add("D2:D41", missing_status_rule)
    review.conditional_formatting.add("K2:K41", reviewer_mismatch_rule)
    review.conditional_formatting.add("L2:L41", date_mismatch_rule)

    review.protection.set_password(SHEET_PASSWORD)
    review.protection.sheet = True
    review.protection.insertRows = True
    review.protection.deleteRows = True
    review.protection.insertColumns = True
    review.protection.deleteColumns = True
    review.protection.sort = True
    review.protection.autoFilter = True
    review.protection.formatRows = True
    review.protection.formatColumns = True

    _write_text_sheet(
        instructions,
        excel_instructions(),
        "FINAL-v14 blinded morphology naming session — instructions",
    )
    _write_text_sheet(rubric_sheet, rubric, "Fixed morphology rubric")
    workbook.active = 0

    raw = io.BytesIO()
    workbook.save(raw)
    payload = _canonicalize_xlsx(raw.getvalue(), source_sealed_utc)
    _safe_zip_members(payload, "generated workbook")
    return payload


def _defined_name_text(workbook: Any, name: str) -> str:
    definition = workbook.defined_names.get(name)
    if definition is None:
        raise HandoffError(f"workbook: missing defined name {name}")
    return str(definition.attr_text)


def _validate_option_definitions(workbook: Any, label: str) -> None:
    expected_names = {
        "_review_status": "'_Options'!$A$1:$A$1",
        "_morphology_categories": "'_Options'!$B$1:$B$11",
        "_confidence": "'_Options'!$C$1:$C$5",
        "_artifact": "'_Options'!$D$1:$D$2",
        "_attestation": "'_Options'!$E$1:$E$1",
    }
    if set(workbook.defined_names) != set(expected_names):
        raise HandoffError(f"{label}: defined-name census drifted")
    for name, reference in expected_names.items():
        if _defined_name_text(workbook, name) != reference:
            raise HandoffError(f"{label}: defined name {name} drifted")
    option_sheet = workbook[OPTIONS_SHEET]
    expected_lists = {
        1: STATUS_OPTIONS,
        2: ONTOLOGY,
        3: CONFIDENCE_OPTIONS,
        4: ARTIFACT_OPTIONS,
        5: ATTESTATION_OPTIONS,
    }
    for column, expected_values in expected_lists.items():
        observed = [
            option_sheet.cell(row, column).value
            for row in range(1, len(expected_values) + 1)
        ]
        if observed != list(expected_values):
            raise HandoffError(f"{label}: option list {column} drifted")
    expected_cells = {
        *(f"A{row}" for row in range(1, 2)),
        *(f"B{row}" for row in range(1, 12)),
        *(f"C{row}" for row in range(1, 6)),
        *(f"D{row}" for row in range(1, 3)),
        *(f"E{row}" for row in range(1, 2)),
    }
    observed_cells = {
        cell.coordinate
        for row in option_sheet.iter_rows()
        for cell in row
        if cell.value is not None or cell.comment is not None or cell.hyperlink is not None
    }
    if observed_cells != expected_cells:
        raise HandoffError(f"{label}: option worksheet cell census drifted")
    if any(
        cell.comment is not None or cell.hyperlink is not None
        for row in option_sheet.iter_rows()
        for cell in row
    ):
        raise HandoffError(f"{label}: option worksheet annotations are forbidden")


def _validate_data_validations(review: Any, label: str) -> None:
    expected_validations = {
        "D2:D41": ("list", "=_review_status", None, False),
        "E2:E41": ("list", "=_morphology_categories", None, False),
        "F2:F41": ("list", "=_morphology_categories", None, True),
        "G2:G41": ("list", "=_morphology_categories", None, True),
        "I2:I41": ("list", "=_confidence", None, False),
        "J2:J41": ("list", "=_artifact", None, False),
        "L2:L41": ("date", "DATE(2000,1,1)", "DATE(2100,12,31)", False),
        "M2:M41": ("list", "=_attestation", None, False),
    }
    observed_validations: dict[str, tuple[Any, Any, Any, bool]] = {}
    for validation in review.data_validations.dataValidation:
        key = str(validation.sqref)
        observed_validations[key] = (
            validation.type,
            validation.formula1,
            validation.formula2,
            bool(validation.allow_blank),
        )
        if not validation.showErrorMessage or validation.errorStyle != "stop":
            raise HandoffError(f"{label}: validation is not fail-closed")
    if observed_validations != expected_validations:
        raise HandoffError(f"{label}: data-validation contract drifted")


def _reject_undeclared_workbook_content(workbook: Any, label: str) -> None:
    expected_text_cells = {
        INSTRUCTIONS_SHEET: {"A1", "A2"},
        RUBRIC_SHEET: {"A1", "A2"},
    }
    for sheet_name, expected_cells in expected_text_cells.items():
        sheet = workbook[sheet_name]
        observed = {
            cell.coordinate
            for row in sheet.iter_rows()
            for cell in row
            if cell.value is not None
            or cell.comment is not None
            or cell.hyperlink is not None
        }
        if observed != expected_cells:
            raise HandoffError(f"{label}: {sheet_name} cell census drifted")
        if any(
            cell.comment is not None or cell.hyperlink is not None
            for row in sheet.iter_rows()
            for cell in row
        ):
            raise HandoffError(f"{label}: {sheet_name} annotations are forbidden")
    review = workbook[REVIEW_SHEET]
    for row in review.iter_rows():
        for cell in row:
            if cell.row > 41 or cell.column > 13:
                if cell.value is not None or cell.comment is not None or cell.hyperlink is not None:
                    raise HandoffError(
                        f"{label}: undeclared review content at {cell.coordinate}"
                    )
            elif cell.comment is not None:
                raise HandoffError(f"{label}: cell comments are forbidden")
    for sheet in workbook.worksheets:
        if getattr(sheet, "_images", []) or getattr(sheet, "_charts", []):
            raise HandoffError(f"{label}: embedded drawings/charts are forbidden")


def _cell_value(value: Any) -> str:
    return "" if value is None else str(value)


def validate_template_workbook(
    payload: bytes,
    rows: Sequence[Mapping[str, str]],
    rubric: str,
    *,
    label: str = "template workbook",
) -> dict[str, Any]:
    _safe_zip_members(payload, label)
    codes = [row["blinded_code"] for row in rows]
    expected_links = {f"montages/{code}.jpg" for code in codes}
    _validate_relative_hyperlink_relationships(payload, expected_links, label)
    try:
        workbook = load_workbook(io.BytesIO(payload), data_only=False, keep_links=True)
    except Exception as exc:  # openpyxl exposes multiple parser exception types
        raise HandoffError(f"{label}: openpyxl could not parse workbook") from exc
    if workbook.sheetnames != EXPECTED_SHEETS:
        raise HandoffError(f"{label}: worksheet census/order drifted")
    if getattr(workbook, "_external_links", []):
        raise HandoffError(f"{label}: external workbook links are forbidden")
    if not workbook.security.lockStructure:
        raise HandoffError(f"{label}: workbook structure protection drifted")
    if workbook[OPTIONS_SHEET].sheet_state != "veryHidden":
        raise HandoffError(f"{label}: option worksheet is not veryHidden")
    review = workbook[REVIEW_SHEET]
    if review.freeze_panes != "D2" or not review.protection.sheet:
        raise HandoffError(f"{label}: review worksheet guardrails drifted")
    if set(review.tables) != {TABLE_NAME} or review.tables[TABLE_NAME].ref != "A1:M41":
        raise HandoffError(f"{label}: Excel table contract drifted")
    headers = [review.cell(1, column).value for column in range(1, 14)]
    if headers != FORM_COLUMNS:
        raise HandoffError(f"{label}: header schema drifted")
    hyperlinks: set[str] = set()
    for row_index, expected in enumerate(rows, start=2):
        observed_fixed = [review.cell(row_index, column).value for column in range(1, 4)]
        wanted_fixed: list[Any] = [
            int(expected["presentation_order"]),
            expected["blinded_code"],
            int(expected["n_tiles"]),
        ]
        if observed_fixed != wanted_fixed:
            raise HandoffError(f"{label}: fixed row {row_index} differs from source")
        for column in range(4, 14):
            cell = review.cell(row_index, column)
            if _cell_value(cell.value) != "" or cell.data_type == "f":
                raise HandoffError(f"{label}: response cell is not blank")
            if cell.protection.locked:
                raise HandoffError(f"{label}: response cell is unexpectedly locked")
        for column in range(1, 4):
            if not review.cell(row_index, column).protection.locked:
                raise HandoffError(f"{label}: fixed cell is unexpectedly editable")
        link = review.cell(row_index, 2).hyperlink
        if link is None or link.target != f"montages/{expected['blinded_code']}.jpg":
            raise HandoffError(f"{label}: montage hyperlink drifted")
        hyperlinks.add(str(link.target))
    if hyperlinks != expected_links:
        raise HandoffError(f"{label}: montage hyperlink census drifted")
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    raise HandoffError(f"{label}: formulas are forbidden")

    _validate_option_definitions(workbook, label)
    _validate_data_validations(review, label)
    _reject_undeclared_workbook_content(workbook, label)
    if workbook[INSTRUCTIONS_SHEET]["A2"].value != excel_instructions():
        raise HandoffError(f"{label}: embedded instructions drifted")
    if workbook[RUBRIC_SHEET]["A2"].value != rubric:
        raise HandoffError(f"{label}: embedded rubric drifted")

    public_text = "\n".join(
        _cell_value(cell.value)
        for worksheet in workbook.worksheets
        for row in worksheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str)
    )
    leaks = sorted(
        term for term in FORBIDDEN_PUBLIC_TERMS if term.casefold() in public_text.casefold()
    )
    if leaks:
        raise HandoffError(f"{label}: forbidden blinded terms found: {leaks}")
    return {
        "status": "PASS",
        "rows": len(rows),
        "dropdown_columns": 7,
        "date_validation_columns": 1,
        "montage_hyperlinks": len(hyperlinks),
        "macros": 0,
        "external_data_links": 0,
    }


def _package_expected_files(rows: Sequence[Mapping[str, str]]) -> set[str]:
    return {
        "HANDOFF_MANIFEST.json",
        "HANDOFF_MANIFEST.sha256",
        "INSTRUCTIONS.md",
        "RUBRIC.md",
        "PUBLIC_SALT_SHA256.txt",
        WORKBOOK_NAME,
        *(f"montages/{row['blinded_code']}.jpg" for row in rows),
    }


def validate_delivery_package(
    source: Mapping[str, Any] | None = None,
) -> dict[str, Any]:
    source = source or validate_source_package()
    if DELIVERY_ROOT.is_symlink() or not DELIVERY_ROOT.is_dir():
        raise HandoffError("Excel handoff root is missing or is a symlink")
    if {path.name for path in DELIVERY_ROOT.iterdir()} != {
        "FOR_PATHOLOGIST",
        "README_COORDINATOR.md",
        "PRE_READER_STATUS.json",
    }:
        raise HandoffError("Excel handoff root census drifted")
    package = DELIVERY_PACKAGE
    if package.is_symlink() or not package.is_dir():
        raise HandoffError("Excel handoff package is missing or is a symlink")
    expected_files = _package_expected_files(source["rows"])
    observed_files = {
        path.relative_to(package).as_posix()
        for path in package.rglob("*")
        if path.is_file()
    }
    if observed_files != expected_files:
        raise HandoffError("Excel handoff package file census drifted")
    if any(path.is_symlink() for path in package.rglob("*")):
        raise HandoffError("Excel handoff package contains a symlink")

    manifest_path = package / "HANDOFF_MANIFEST.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    expected_keys = {
        "schema_version",
        "package",
        "interface_revision",
        "sealed_utc",
        "n_montages",
        "public_salt_sha256",
        "source_handoff_manifest_sha256",
        "source_review_form_sha256",
        "source_independent_audit_receipt_sha256",
        "files",
    }
    try:
        sealed = dt.datetime.fromisoformat(str(manifest.get("sealed_utc", "")))
    except ValueError as exc:
        raise HandoffError("Excel handoff sealed timestamp is malformed") from exc
    salt = (package / "PUBLIC_SALT_SHA256.txt").read_text(encoding="ascii").strip()
    if (
        set(manifest) != expected_keys
        or manifest.get("schema_version") != SCHEMA_VERSION
        or manifest.get("package") != PACKAGE_NAME
        or manifest.get("interface_revision") != INTERFACE_REVISION
        or manifest.get("n_montages") != 40
        or manifest.get("public_salt_sha256") != salt
        or manifest.get("source_handoff_manifest_sha256")
        != PINNED_SOURCE_MANIFEST_SHA256
        or manifest.get("source_review_form_sha256") != PINNED_SOURCE_FORM_SHA256
        or manifest.get("source_independent_audit_receipt_sha256")
        != PINNED_INDEPENDENT_AUDIT_SHA256
        or sealed.tzinfo is None
    ):
        raise HandoffError("Excel handoff manifest contract drifted")
    sidecar = (package / "HANDOFF_MANIFEST.sha256").read_text(encoding="ascii")
    manifest_sha = sha256_file(manifest_path)
    if sidecar != f"{manifest_sha}  HANDOFF_MANIFEST.json\n":
        raise HandoffError("Excel handoff manifest sidecar failed")
    records = manifest.get("files")
    if not isinstance(records, list):
        raise HandoffError("Excel handoff manifest file inventory is malformed")
    expected_inventory = expected_files - {
        "HANDOFF_MANIFEST.json",
        "HANDOFF_MANIFEST.sha256",
    }
    if {str(record.get("path", "")) for record in records} != expected_inventory:
        raise HandoffError("Excel handoff manifest inventory drifted")
    for record in records:
        relative = PurePosixPath(str(record.get("path", "")))
        if relative.is_absolute() or ".." in relative.parts:
            raise HandoffError("Excel handoff manifest contains an unsafe path")
        path = package / relative.as_posix()
        if path.stat().st_size != int(record.get("size_bytes", -1)):
            raise HandoffError(f"Excel handoff file size mismatch: {relative}")
        if sha256_file(path) != str(record.get("sha256", "")):
            raise HandoffError(f"Excel handoff file hash mismatch: {relative}")

    if (package / "RUBRIC.md").read_text(encoding="utf-8") != source["rubric"]:
        raise HandoffError("Excel handoff rubric differs from frozen source")
    if (package / "PUBLIC_SALT_SHA256.txt").read_text(
        encoding="ascii"
    ) != source["salt"]:
        raise HandoffError("Excel handoff public salt differs from frozen source")
    if (package / "INSTRUCTIONS.md").read_text(
        encoding="utf-8"
    ) != excel_instructions():
        raise HandoffError("Excel handoff instructions drifted")
    for row in source["rows"]:
        name = f"montages/{row['blinded_code']}.jpg"
        if sha256_file(package / name) != sha256_file(SOURCE_PACKAGE / name):
            raise HandoffError(f"Excel handoff montage differs from source: {name}")
    workbook_payload = (package / WORKBOOK_NAME).read_bytes()
    workbook_check = validate_template_workbook(
        workbook_payload, source["rows"], source["rubric"]
    )
    canonical_package = AUDIT_ROOT / "public"
    if tree_inventory(package) != tree_inventory(canonical_package):
        raise HandoffError("Excel handoff delivery differs from its canonical mirror")
    if (DELIVERY_ROOT / "README_COORDINATOR.md").read_text(
        encoding="utf-8"
    ) != _coordinator_readme():
        raise HandoffError("Excel handoff coordinator instructions drifted")
    observed_status = json.loads(
        (DELIVERY_ROOT / "PRE_READER_STATUS.json").read_text(encoding="utf-8")
    )
    if observed_status != _pre_reader_status(
        manifest_sha=manifest_sha,
        sealed_utc=str(manifest["sealed_utc"]),
    ):
        raise HandoffError("Excel handoff pre-reader status drifted")
    return {
        "status": "PASS",
        "manifest_sha256": manifest_sha,
        "workbook_sha256": sha256_bytes(workbook_payload),
        "files": len(expected_files),
        "montages": 40,
        "source_manifest_sha256": PINNED_SOURCE_MANIFEST_SHA256,
        "source_form_sha256": PINNED_SOURCE_FORM_SHA256,
        "canonical_mirror": "PASS",
        "workbook": workbook_check,
    }


def _coordinator_readme() -> str:
    return """# FINAL-v14 Excel reader handoff — coordinator only

This is an interface-only revision of the sealed FINAL-v14 reader package. It preserves the same blinded codes, presentation order, montages, ontology, and public salt. Use it only if the original CSV package has not been distributed and no reader has begun. Do not switch formats during a reading session.

Give the pathologist only the `FOR_PATHOLOGIST` directory from this `v14_xlsx` package. Keep the original `reviews/v14` package as the immutable source record, and keep every file in the governed rerun's `reader_package/embargoed` directory inaccessible.

On return, preserve and SHA-256 seal the raw `completed_review_form.xlsx` before any conversion or key access. From the repository root run `.venv/bin/python tools/final_v14_excel_handoff.py ingest --completed /absolute/path/to/completed_review_form.xlsx`. The fail-closed tool writes the governed blinded stage to `reports/reruns/final_v14_additions_20260903/e4v_post_reader_xlsx/human_reader_1`, validates the workbook, and produces a canonical `completed_review_form.csv`. Only after that stage passes may the embargoed key be opened.
"""


def _pre_reader_status(*, manifest_sha: str, sealed_utc: str) -> dict[str, Any]:
    return {
        "schema_version": SCHEMA_VERSION,
        "component": "final_v14_xlsx_reader_handoff",
        "created_utc": sealed_utc,
        "status": "READY_FOR_PATHOLOGIST_IF_ELIGIBLE",
        "technical_validation": "PASS",
        "release_authorization": "CONDITIONAL",
        "interface_revision": INTERFACE_REVISION,
        "interface_only_change": True,
        "original_package_unchanged": True,
        "release_condition": (
            "Use only if the original CSV package has not been distributed and no "
            "reader has begun."
        ),
        "source_handoff_manifest_sha256": PINNED_SOURCE_MANIFEST_SHA256,
        "handoff_manifest_sha256": manifest_sha,
        "give_to_pathologist": "FOR_PATHOLOGIST only",
        "expected_return": COMPLETED_WORKBOOK_NAME,
        "next_action": (
            "Seal the returned XLSX, run blinded ingestion to canonical CSV, and only "
            "then open the embargoed key."
        ),
    }


def _build_delivery_stage(stage_root: Path, source: Mapping[str, Any], sealed_utc: str) -> None:
    package = stage_root / "FOR_PATHOLOGIST"
    montages = package / "montages"
    montages.mkdir(parents=True)
    instructions = excel_instructions().encode("utf-8")
    write_new(package / "INSTRUCTIONS.md", instructions)
    write_new(package / "RUBRIC.md", source["rubric"].encode("utf-8"))
    write_new(package / "PUBLIC_SALT_SHA256.txt", source["salt"].encode("ascii"))
    workbook_payload = build_workbook(
        source["rows"], source["rubric"], source["manifest"]["sealed_utc"]
    )
    validate_template_workbook(
        workbook_payload, source["rows"], source["rubric"], label="new workbook"
    )
    replay_payload = build_workbook(
        source["rows"], source["rubric"], source["manifest"]["sealed_utc"]
    )
    if workbook_payload != replay_payload:
        raise HandoffError("workbook generation is not byte-deterministic")
    write_new(package / WORKBOOK_NAME, workbook_payload)
    for row in source["rows"]:
        code = row["blinded_code"]
        shutil.copyfile(SOURCE_PACKAGE / f"montages/{code}.jpg", montages / f"{code}.jpg")
        os.chmod(montages / f"{code}.jpg", 0o600)

    included = sorted(
        path
        for path in package.rglob("*")
        if path.is_file()
        and path.name not in {"HANDOFF_MANIFEST.json", "HANDOFF_MANIFEST.sha256"}
    )
    manifest = {
        "schema_version": SCHEMA_VERSION,
        "package": PACKAGE_NAME,
        "interface_revision": INTERFACE_REVISION,
        "sealed_utc": sealed_utc,
        "n_montages": 40,
        "public_salt_sha256": source["salt"].strip(),
        "source_handoff_manifest_sha256": PINNED_SOURCE_MANIFEST_SHA256,
        "source_review_form_sha256": PINNED_SOURCE_FORM_SHA256,
        "source_independent_audit_receipt_sha256": PINNED_INDEPENDENT_AUDIT_SHA256,
        "files": [identity(path, relative_to=package) for path in included],
    }
    write_new(package / "HANDOFF_MANIFEST.json", json_bytes(manifest))
    manifest_sha = sha256_file(package / "HANDOFF_MANIFEST.json")
    write_new(
        package / "HANDOFF_MANIFEST.sha256",
        f"{manifest_sha}  HANDOFF_MANIFEST.json\n".encode("ascii"),
    )
    write_new(stage_root / "README_COORDINATOR.md", _coordinator_readme().encode("utf-8"))
    status = _pre_reader_status(manifest_sha=manifest_sha, sealed_utc=sealed_utc)
    write_new(stage_root / "PRE_READER_STATUS.json", json_bytes(status))


def _handoff_receipt_payload(
    delivery: Mapping[str, Any], *, created_utc: str
) -> dict[str, Any]:
    return {
        "schema_version": SCHEMA_VERSION,
        "component": "final_v14_xlsx_reader_handoff",
        "created_utc": created_utc,
        "status": "PASS",
        "technical_handoff_ready": True,
        "release_authorization": "CONDITIONAL",
        "release_condition": (
            "Use only if the original CSV package has not been distributed and no "
            "reader has begun."
        ),
        "interface_only_change": True,
        "original_v14_package_unchanged": True,
        "builder": identity(Path(__file__)),
        "dependencies": {
            "source_handoff_manifest": identity(
                SOURCE_PACKAGE / "HANDOFF_MANIFEST.json"
            ),
            "source_review_form": identity(SOURCE_PACKAGE / "review_form.csv"),
            "source_independent_audit": identity(INDEPENDENT_AUDIT_RECEIPT),
        },
        "artifacts": {
            "canonical_blank_workbook": identity(AUDIT_ROOT / "public" / WORKBOOK_NAME),
            "canonical_handoff_manifest": identity(
                AUDIT_ROOT / "public" / "HANDOFF_MANIFEST.json"
            ),
            "handoff_manifest": identity(DELIVERY_PACKAGE / "HANDOFF_MANIFEST.json"),
            "handoff_manifest_sidecar": identity(
                DELIVERY_PACKAGE / "HANDOFF_MANIFEST.sha256"
            ),
            "blank_workbook": identity(DELIVERY_PACKAGE / WORKBOOK_NAME),
            "coordinator_readme": identity(DELIVERY_ROOT / "README_COORDINATOR.md"),
            "pre_reader_status": identity(DELIVERY_ROOT / "PRE_READER_STATUS.json"),
        },
        "verification": dict(delivery),
    }


def _validate_handoff_receipt(delivery: Mapping[str, Any]) -> str:
    if AUDIT_ROOT.is_symlink() or not AUDIT_ROOT.is_dir():
        raise HandoffError("Excel handoff receipt directory is missing or is a symlink")
    if {path.name for path in AUDIT_ROOT.iterdir()} != {
        "public",
        HANDOFF_RECEIPT.name,
    }:
        raise HandoffError("Excel handoff receipt directory census drifted")
    if HANDOFF_RECEIPT.is_symlink() or not HANDOFF_RECEIPT.is_file():
        raise HandoffError("Excel handoff receipt is missing or is a symlink")
    try:
        receipt = json.loads(HANDOFF_RECEIPT.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, UnicodeDecodeError) as exc:
        raise HandoffError("Excel handoff receipt is malformed") from exc
    expected_keys = {
        "schema_version",
        "component",
        "created_utc",
        "status",
        "technical_handoff_ready",
        "release_authorization",
        "release_condition",
        "interface_only_change",
        "original_v14_package_unchanged",
        "builder",
        "dependencies",
        "artifacts",
        "verification",
    }
    try:
        created = dt.datetime.fromisoformat(str(receipt.get("created_utc", "")))
    except ValueError as exc:
        raise HandoffError("Excel handoff receipt timestamp is malformed") from exc
    if (
        set(receipt) != expected_keys
        or receipt.get("schema_version") != SCHEMA_VERSION
        or receipt.get("component") != "final_v14_xlsx_reader_handoff"
        or receipt.get("status") != "PASS"
        or receipt.get("technical_handoff_ready") is not True
        or receipt.get("release_authorization") != "CONDITIONAL"
        or receipt.get("release_condition")
        != "Use only if the original CSV package has not been distributed and no reader has begun."
        or receipt.get("interface_only_change") is not True
        or receipt.get("original_v14_package_unchanged") is not True
        or created.tzinfo is None
    ):
        raise HandoffError("Excel handoff receipt contract drifted")
    expected = _handoff_receipt_payload(
        delivery, created_utc=str(receipt["created_utc"])
    )
    if receipt != expected:
        raise HandoffError("Excel handoff receipt identities or content drifted")
    return sha256_file(HANDOFF_RECEIPT)


def build_delivery() -> dict[str, Any]:
    source = validate_source_package()
    if DELIVERY_ROOT.exists() or DELIVERY_ROOT.is_symlink():
        check = validate_delivery_package(source)
        reused_existing = True
    else:
        DELIVERY_ROOT.parent.mkdir(parents=True, exist_ok=True)
        stage = Path(tempfile.mkdtemp(prefix=".v14_xlsx.", dir=DELIVERY_ROOT.parent))
        try:
            sealed_utc = utc_now()
            _build_delivery_stage(stage, source, sealed_utc)
            os.replace(stage, DELIVERY_ROOT)
        finally:
            if stage.exists():
                shutil.rmtree(stage)
        reused_existing = False
        AUDIT_ROOT.mkdir(parents=True, exist_ok=True)
        canonical_package = AUDIT_ROOT / "public"
        if canonical_package.exists() or canonical_package.is_symlink():
            raise HandoffError(
                f"refusing to replace existing canonical Excel handoff: {canonical_package}"
            )
        shutil.copytree(DELIVERY_PACKAGE, canonical_package)
        for path in canonical_package.rglob("*"):
            if path.is_file():
                os.chmod(path, 0o600)
        check = validate_delivery_package(source)
    if not HANDOFF_RECEIPT.exists():
        receipt = _handoff_receipt_payload(check, created_utc=utc_now())
        write_new(HANDOFF_RECEIPT, json_bytes(receipt))
    _validate_handoff_receipt(check)
    return {**check, "reused_existing": reused_existing}


def _normalise_date(value: Any, label: str) -> str:
    if isinstance(value, dt.datetime):
        parsed = value.date()
    elif isinstance(value, dt.date):
        parsed = value
    else:
        text = _cell_value(value).strip()
        try:
            parsed = dt.date.fromisoformat(text)
        except ValueError as exc:
            raise HandoffError(f"{label}: review_date must be YYYY-MM-DD") from exc
        if text != parsed.isoformat():
            raise HandoffError(f"{label}: review_date must be canonical YYYY-MM-DD")
    if not dt.date(2000, 1, 1) <= parsed <= dt.date(2100, 12, 31):
        raise HandoffError(f"{label}: review_date is outside 2000-01-01 to 2100-12-31")
    return parsed.isoformat()


def _reject_formulas(workbook: Any, label: str) -> None:
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    raise HandoffError(f"{label}: formulas are forbidden ({cell.coordinate})")


def validate_completed_workbook(
    payload: bytes,
    rows: Sequence[Mapping[str, str]],
    rubric: str,
    *,
    label: str = "completed workbook",
) -> list[dict[str, Any]]:
    _safe_zip_members(payload, label)
    expected_links = {f"montages/{row['blinded_code']}.jpg" for row in rows}
    _validate_relative_hyperlink_relationships(payload, expected_links, label)
    try:
        workbook = load_workbook(io.BytesIO(payload), data_only=False, keep_links=True)
    except Exception as exc:
        raise HandoffError(f"{label}: openpyxl could not parse workbook") from exc
    if workbook.sheetnames != EXPECTED_SHEETS:
        raise HandoffError(f"{label}: worksheet census/order drifted")
    if getattr(workbook, "_external_links", []):
        raise HandoffError(f"{label}: external workbook links are forbidden")
    if not workbook.security.lockStructure:
        raise HandoffError(f"{label}: workbook structure protection changed")
    _reject_formulas(workbook, label)
    if workbook[OPTIONS_SHEET].sheet_state != "veryHidden":
        raise HandoffError(f"{label}: option worksheet is not veryHidden")
    _validate_option_definitions(workbook, label)
    _reject_undeclared_workbook_content(workbook, label)
    if workbook[INSTRUCTIONS_SHEET]["A2"].value != excel_instructions():
        raise HandoffError(f"{label}: instructions sheet changed")
    if workbook[RUBRIC_SHEET]["A2"].value != rubric:
        raise HandoffError(f"{label}: rubric sheet changed")
    review = workbook[REVIEW_SHEET]
    if review.freeze_panes != "D2" or not review.protection.sheet:
        raise HandoffError(f"{label}: review worksheet guardrails changed")
    headers = [review.cell(1, column).value for column in range(1, 14)]
    if headers != FORM_COLUMNS:
        raise HandoffError(f"{label}: form header schema changed")
    if set(review.tables) != {TABLE_NAME} or review.tables[TABLE_NAME].ref != "A1:M41":
        raise HandoffError(f"{label}: form table changed")
    _validate_data_validations(review, label)
    for row in review.iter_rows(min_row=42):
        if any(_cell_value(cell.value).strip() for cell in row):
            raise HandoffError(f"{label}: unexpected data below the 40 governed rows")

    completed: list[dict[str, Any]] = []
    for row_index, expected in enumerate(rows, start=2):
        code = expected["blinded_code"]
        fixed = [review.cell(row_index, column).value for column in range(1, 4)]
        wanted: list[Any] = [
            int(expected["presentation_order"]),
            code,
            int(expected["n_tiles"]),
        ]
        if fixed != wanted:
            raise HandoffError(f"{label}/{code}: fixed identity/order cells changed")
        if any(
            not review.cell(row_index, column).protection.locked
            for column in range(1, 4)
        ) or any(
            review.cell(row_index, column).protection.locked
            for column in range(4, 14)
        ):
            raise HandoffError(f"{label}/{code}: cell protection contract changed")
        hyperlink = review.cell(row_index, 2).hyperlink
        if hyperlink is None or hyperlink.target != f"montages/{code}.jpg":
            raise HandoffError(f"{label}/{code}: montage hyperlink changed")
        values = {
            FORM_COLUMNS[column - 1]: review.cell(row_index, column).value
            for column in range(1, 14)
        }
        status = _cell_value(values["review_status"]).strip()
        primary = _cell_value(values["primary_category"]).strip()
        secondary_1 = _cell_value(values["secondary_category_1"]).strip()
        secondary_2 = _cell_value(values["secondary_category_2"]).strip()
        description = _cell_value(values["free_text_description"]).strip()
        confidence_raw = values["confidence_1_to_5"]
        artifact = _cell_value(values["artifact_uninterpretable"]).strip()
        reviewer = _cell_value(values["reviewer_id"]).strip()
        review_date = _normalise_date(values["review_date"], f"{label}/{code}")
        attestation = _cell_value(values["blinding_attestation"]).strip()
        if status != "complete":
            raise HandoffError(f"{label}/{code}: review_status must be complete")
        if primary not in ONTOLOGY:
            raise HandoffError(f"{label}/{code}: invalid primary category")
        for name, secondary in (
            ("secondary_category_1", secondary_1),
            ("secondary_category_2", secondary_2),
        ):
            if secondary and secondary not in ONTOLOGY:
                raise HandoffError(f"{label}/{code}: invalid {name}")
        nonblank_categories = [primary, *(x for x in (secondary_1, secondary_2) if x)]
        if len(nonblank_categories) != len(set(nonblank_categories)):
            raise HandoffError(f"{label}/{code}: primary/secondary categories repeat")
        if not description:
            raise HandoffError(f"{label}/{code}: morphology description is blank")
        for field_name, text_value in (
            ("free_text_description", description),
            ("reviewer_id", reviewer),
        ):
            if text_value.startswith(("=", "+", "-", "@")):
                raise HandoffError(
                    f"{label}/{code}: {field_name} has an unsafe spreadsheet prefix"
                )
        try:
            confidence = int(confidence_raw)
        except (TypeError, ValueError) as exc:
            raise HandoffError(f"{label}/{code}: confidence is not an integer") from exc
        if str(confidence_raw).strip() not in {str(x) for x in CONFIDENCE_OPTIONS}:
            raise HandoffError(f"{label}/{code}: confidence must be 1 through 5")
        if artifact not in ARTIFACT_OPTIONS:
            raise HandoffError(f"{label}/{code}: artifact flag must be yes or no")
        if not reviewer:
            raise HandoffError(f"{label}/{code}: reviewer_id is blank")
        if attestation != ATTESTATION_OPTIONS[0]:
            raise HandoffError(f"{label}/{code}: blinding attestation is invalid")
        completed.append(
            {
                "presentation_order": int(expected["presentation_order"]),
                "blinded_code": code,
                "n_tiles": int(expected["n_tiles"]),
                "review_status": status,
                "primary_category": primary,
                "secondary_category_1": secondary_1,
                "secondary_category_2": secondary_2,
                "free_text_description": description,
                "confidence_1_to_5": confidence,
                "artifact_uninterpretable": artifact,
                "reviewer_id": reviewer,
                "review_date": review_date,
                "blinding_attestation": attestation,
            }
        )
    if len({row["reviewer_id"] for row in completed}) != 1:
        raise HandoffError(f"{label}: reviewer_id must be identical in all rows")
    if len({row["review_date"] for row in completed}) != 1:
        raise HandoffError(f"{label}: review_date must be identical in all rows")
    return completed


def _csv_payload(rows: Iterable[Mapping[str, Any]]) -> bytes:
    stream = io.StringIO(newline="")
    writer = csv.DictWriter(stream, fieldnames=FORM_COLUMNS, lineterminator="\n")
    writer.writeheader()
    writer.writerows(rows)
    return stream.getvalue().encode("utf-8")


def ingest_completed(
    completed_path: Path, output_root: Path | None = None
) -> dict[str, Any]:
    source = validate_source_package()
    validate_delivery_package(source)
    output_root = output_root or POST_READER_OUTPUT
    if output_root.resolve(strict=False) != POST_READER_OUTPUT.resolve(strict=False):
        raise HandoffError(
            f"completed-read output is hard-pinned to {POST_READER_OUTPUT}"
        )
    completed_path = completed_path.resolve(strict=True)
    if completed_path.suffix.casefold() != ".xlsx" or not completed_path.is_file():
        raise HandoffError("completed form must be an existing .xlsx file")
    if output_root.exists() or output_root.is_symlink():
        raise HandoffError(f"append-only output already exists: {output_root}")
    raw_payload = completed_path.read_bytes()
    raw_sha = sha256_bytes(raw_payload)
    completed_rows = validate_completed_workbook(
        raw_payload, source["rows"], source["rubric"]
    )
    output_root.parent.mkdir(parents=True, exist_ok=True)
    stage = Path(
        tempfile.mkdtemp(prefix=f".{output_root.name}.", dir=output_root.parent)
    )
    try:
        write_new(stage / COMPLETED_WORKBOOK_NAME, raw_payload, mode=0o400)
        csv_payload = _csv_payload(completed_rows)
        write_new(stage / COMPLETED_CSV_NAME, csv_payload, mode=0o400)
        readme = """# BLINDED completed FINAL-v14 pathology read

This directory contains the sealed raw Excel response and its validated canonical CSV export. It contains no montage-to-concept key and must remain unchanged. Open the embargoed key only after this receipt has been independently verified.
"""
        write_new(stage / "README_BLINDED.md", readme.encode("utf-8"), mode=0o400)
        receipt = {
            "schema_version": SCHEMA_VERSION,
            "component": "final_v14_xlsx_completed_read_ingest",
            "created_utc": utc_now(),
            "status": "PASS",
            "blinded_validation_complete": True,
            "safe_to_begin_unblinding": True,
            "rows": len(completed_rows),
            "reviewer_ids": sorted({row["reviewer_id"] for row in completed_rows}),
            "review_dates": sorted({row["review_date"] for row in completed_rows}),
            "dependencies": {
                "blank_workbook": identity(DELIVERY_PACKAGE / WORKBOOK_NAME),
                "handoff_manifest": identity(DELIVERY_PACKAGE / "HANDOFF_MANIFEST.json"),
                "source_handoff_manifest_sha256": PINNED_SOURCE_MANIFEST_SHA256,
                "validator": identity(Path(__file__)),
            },
            "artifacts": {
                "raw_completed_workbook": {
                    "path": COMPLETED_WORKBOOK_NAME,
                    "size_bytes": len(raw_payload),
                    "sha256": raw_sha,
                },
                "canonical_completed_csv": {
                    "path": COMPLETED_CSV_NAME,
                    "size_bytes": len(csv_payload),
                    "sha256": sha256_bytes(csv_payload),
                },
            },
            "promise": "No embargoed key was accessed by this ingestion stage.",
        }
        write_new(stage / "receipt.json", json_bytes(receipt), mode=0o400)
        os.replace(stage, output_root)
    finally:
        if stage.exists():
            shutil.rmtree(stage)
    return receipt


def verify_handoff() -> dict[str, Any]:
    source = validate_source_package()
    delivery = validate_delivery_package(source)
    receipt_sha = _validate_handoff_receipt(delivery)
    return {
        "schema_version": SCHEMA_VERSION,
        "component": "final_v14_xlsx_reader_handoff_verification",
        "status": "PASS",
        "technical_handoff_ready": True,
        "release_authorization": "CONDITIONAL",
        "release_condition": (
            "Use only if the original CSV package has not been distributed and no "
            "reader has begun."
        ),
        "source": {
            "manifest_sha256": source["manifest_sha256"],
            "form_sha256": source["form_sha256"],
            "original_v14_package_unchanged": True,
        },
        "delivery": delivery,
        "handoff_receipt_sha256": receipt_sha,
    }


def _parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    subparsers = parser.add_subparsers(dest="command", required=True)
    subparsers.add_parser("build", help="build or verify the sealed Excel handoff")
    subparsers.add_parser("verify", help="verify the sealed Excel handoff")
    ingest = subparsers.add_parser(
        "ingest", help="seal and validate a completed XLSX, then export canonical CSV"
    )
    ingest.add_argument("--completed", type=Path, required=True)
    return parser


def main() -> None:
    arguments = _parser().parse_args()
    try:
        if arguments.command == "build":
            result = build_delivery()
        elif arguments.command == "verify":
            result = verify_handoff()
        else:
            result = ingest_completed(arguments.completed)
    except HandoffError as exc:
        raise SystemExit(f"FINAL_V14_XLSX_FAIL: {exc}") from exc
    print(json.dumps(result, indent=2, sort_keys=True, ensure_ascii=False))


if __name__ == "__main__":
    main()
