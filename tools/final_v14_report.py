#!/usr/bin/env python3
"""Render the working Aim-4 report and mentor-annotated atlas from sealed data."""

from __future__ import annotations

import html
import json
import os
import sys
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from tools import final_v14_module1 as source  # noqa: E402

OUT = source.REPO / "reports/final_v14"


def value(item: dict, places: int = 3) -> str:
    if item.get("estimate") is None:
        return "Not estimable"
    lo, hi = item["ci95"]
    point = f"{item['estimate']:.{places}f}"
    return point + (f" [{lo:.{places}f}, {hi:.{places}f}]" if lo is not None else " [CI not estimable]")


def link(path: Path, label: str) -> str:
    return f"[{label}]({os.path.relpath(path, OUT)})"


def atlas(names: dict) -> None:
    cards = []
    for row in names["controlling_reads"]:
        response = row["response"]
        category = response["primary_category"] or "Primary category unrecorded"
        comment = response["free_text_description"] or "No free-text comment recorded."
        secondary = [response.get(f"secondary_category_{i}") for i in [1, 2]]
        secondary = ", ".join(x for x in secondary if x)
        image_path = source.REPO / "reviews/v14_xlsx/FOR_PATHOLOGIST/montages" / (row["controlling_code"] + ".jpg")
        image_url = html.escape(os.path.relpath(image_path, OUT), quote=True)
        cards.append(f"""<article data-search="{html.escape(category + ' ' + comment, quote=True)}">
<h2>Concept {row['prototype_id']:02d} <small>{row['controlling_code']}</small></h2>
<a href="{image_url}"><img loading="lazy" src="{image_url}" alt="Original twelve-tile montage for concept {row['prototype_id']:02d}"></a>
<h3>{html.escape(category)}</h3><p>{html.escape(comment)}</p>
{f'<p class="secondary">Secondary: {html.escape(secondary)}</p>' if secondary else ''}</article>""")
    page = """<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Aim 4 · Mentor-annotated concept atlas</title><style>
body{font:16px/1.55 system-ui,sans-serif;background:#f3f4f1;color:#18322d;margin:0}header,main{max-width:1280px;margin:auto;padding:28px}
h1{font-size:32px;margin:0}header p{max-width:900px}input{padding:12px;width:min(500px,90%);font:inherit;border:1px solid #8ba79e;border-radius:6px}
main{display:grid;grid-template-columns:repeat(auto-fit,minmax(420px,1fr));gap:24px;padding-top:0}article{background:white;padding:20px;border:1px solid #d1dad4;border-radius:8px}
img{width:100%;height:auto}h2{margin:0 0 12px}h3{font-size:17px;margin:12px 0 4px}small{font-size:13px;color:#667b72;font-weight:400}p{white-space:pre-wrap} .secondary{font-size:14px;color:#506259}
@media(max-width:500px){main{grid-template-columns:1fr;padding:14px}header{padding:18px}}[hidden]{display:none}</style>
<header><h1>Mentor-annotated concept atlas</h1><p>32 canonical source concepts · Original montages and controlling responses · FINAL-v14</p>
<p>Categories and comments are reproduced exactly from the accepted final review. Unrecorded fields remain unrecorded. These are descriptive annotations; the preregistered naming-quality gate was not met because required ratings were largely absent. The quantitative models use all 32 unsupervised concepts.</p>
<label for="search">Find a category or phrase</label><br><input id="search" type="search" placeholder="e.g. fibrosis, gland, mucin"><p id="count">32 concepts</p></header><main>""" + "\n".join(cards) + """</main><script>
document.querySelector('#search').addEventListener('input',e=>{let n=0;for(const card of document.querySelectorAll('article')){card.hidden=!card.dataset.search.toLowerCase().includes(e.target.value.toLowerCase());if(!card.hidden)n++}document.querySelector('#count').textContent=n+' concepts'});
</script></html>"""
    (OUT / "mentor_atlas.html").write_text(page)


def figure(results: dict) -> None:
    p = results["ALL32"]["pooled"]
    plt.rcParams.update({"font.size": 10, "axes.spines.top": False, "axes.spines.right": False,
                         "pdf.fonttype": 42, "svg.fonttype": "none"})
    fig, axes = plt.subplots(1, 2, figsize=(10, 4.3), gridspec_kw={"width_ratios": [1.3, 1]})
    labels = ["Full WSI model", "Concept classifier", "Reconstructed score", "Concept + clinical", "Clinical comparator"]
    keys = ["all/full/auroc", "all/concept/auroc", "fidelity/reconstructed_auroc", "all/joint/auroc", "all/clinical/auroc"]
    for i, key in enumerate(keys):
        d = p[key]
        axes[0].errorbar(d["estimate"], i, xerr=np.array([[d["estimate"] - d["ci95"][0]], [d["ci95"][1] - d["estimate"]]]),
                         fmt="o", color="#244d42" if i == 0 else "#437f70", capsize=4)
    axes[0].set(yticks=range(5), yticklabels=labels, xlabel="Patient-OOF AUROC · 95% CI", xlim=(.48, .76))
    axes[0].invert_yaxis()
    axes[0].axvline(.5, color="#9daba4", linestyle="--", linewidth=1)
    axes[0].set_title("A  Source discrimination", loc="left", weight="bold")
    keys2 = ["all/concept_minus_full/auroc", "all/concept_minus_clinical/auroc", "all/joint_minus_concept/auroc"]
    labels2 = ["Concept − full", "Concept − clinical", "Joint − concept"]
    for i, key in enumerate(keys2):
        d = p[key]
        axes[1].errorbar(d["estimate"], i, xerr=np.array([[d["estimate"] - d["ci95"][0]], [d["ci95"][1] - d["estimate"]]]),
                         fmt="o", color="#986448", capsize=4)
    axes[1].set(yticks=range(3), yticklabels=labels2, xlabel="Paired AUROC difference · 95% CI")
    axes[1].invert_yaxis()
    axes[1].axvline(0, color="#9daba4", linestyle="--", linewidth=1)
    axes[1].set_title("B  Paired comparisons", loc="left", weight="bold")
    fig.suptitle("Aim 4 · k=32 unsupervised abundance representation", weight="bold")
    fig.text(.5, .01, "1,239 source patients · fixed OOF predictions · 10,000 stratified paired patient bootstrap draws", ha="center", fontsize=9)
    fig.tight_layout(rect=(0, .04, 1, .94))
    (OUT / "figures").mkdir(exist_ok=True)
    for ext in ["pdf", "svg", "png"]:
        fig.savefig(OUT / "figures" / f"aim4_module1.{ext}", dpi=180, bbox_inches="tight")
    plt.close(fig)


def context_report() -> str:
    root = source.ROOT / "e4m1_context"
    summary = source.read_sealed(root / "results.json")
    for record in summary["artifacts"].values():
        source.verify_record(record)
    tests = pd.read_csv(root / "association_tests.csv")
    rows = []
    for (representation, covariate), block in tests.groupby(["representation", "covariate"], sort=True):
        rows.append(f"| {representation} | {covariate} | {int(block.n_complete_cases.iloc[0])} | {int((block.q_value_bh < .05).sum())}/{len(block)} |")
    fig, axes = plt.subplots(1, 2, figsize=(9, 8), sharey=True)
    for ax, covariate, title in zip(axes, ["oof_score", "kras"], ["A  Association with OOF score", "B  Association with KRAS"], strict=True):
        panel = source.read_sealed(root / "panels" / f"abundance__{covariate}.json")
        for i, row in enumerate(panel["rows"]):
            c = row["contrasts"][0]
            point, (low, high) = c["standardized_effect"], c["standardized_effect_bootstrap"]["ci95"]
            ax.errorbar(point, i, xerr=np.array([[point - low], [high - point]]), fmt="o", ms=4,
                        color="#244d42" if row["q_value_bh"] < .05 else "#a5b4ad", capsize=2)
        ax.axvline(0, linestyle="--", color="#a5b4ad", linewidth=1)
        ax.set_title(title, loc="left", weight="bold", fontsize=11)
        ax.set_xlabel("Adjusted effect / transformed-abundance SD")
    axes[0].set_yticks(range(len(panel["rows"])), [f"Concept {row['prototype_id']:02d}" for row in panel["rows"]])
    axes[0].invert_yaxis()
    fig.suptitle("All 25 concepts mapped in every outer fold", weight="bold")
    fig.text(.5, .015, "Source-subcohort fixed effects · 2,000 paired patient bootstrap draws · dark = BH q < 0.05", ha="center", fontsize=9)
    fig.tight_layout(rect=(0, .04, 1, .95))
    for ext in ["pdf", "svg", "png"]:
        fig.savefig(OUT / "figures" / f"aim4_context.{ext}", dpi=180, bbox_inches="tight")
    plt.close(fig)
    return """## Attribution and context

Twenty-five reference concepts map at cosine ≥0.80 in all five outer folds.
All 25 were retained in each prespecified association family. Eight abundance
coordinates associate with the OOF score and three with KRAS after BH correction
(q<0.05). The KRAS-associated coordinates are 22 and 17 (positive adjusted
differences) and 19 (negative). These v14 identifiers do not denote the legacy
p17/p28 hypotheses. Associations are reported under concept codes; no feature
was selected from them.

Cohort associations occur in 24/25 coordinates. Age, site, MSI, BRAF, and stage
also associate with subsets of the representation. Each covariate is evaluated
separately with source-subcohort fixed effects, except the source-subcohort
model itself. These patterns describe context dependence and do not establish
a biological mechanism or causal confounding effect.

| Representation | Covariate | Complete-case patients | BH q<0.05 / tested |
|---|---|---:|---:|
""" + "\n".join(rows) + f"""

HC3 Wald tests supply the p-values; descriptive intervals use 2,000 patient
bootstrap OLS refits per panel. Attention is a separate sensitivity with its own
score/KRAS BH families, no classifier, and no naming claim. A single pre-existing
FP32 attention mass exceeded one by 6.05×10⁻⁹; the sealed numerical amendment
clips only that boundary roundoff before the arcsine-square-root transform.

![All attributable concepts: score and KRAS associations](figures/aim4_context.png)

All {len(tests)} attempted tests, complete-case counts, contrasts, bootstrap intervals, and
25-ridge/five-logistic coefficient sign summaries are retained in
{link(root / 'results.json', 'the sealed context results')} and
{link(root / 'association_tests.csv', 'the association table')}.
Of these, 274 tests are estimable; the joint source-subcohort test for concept 13
is non-estimable and conservatively occupies its fixed BH family with p=1.

"""


def workbook(results: dict, names: dict) -> None:
    book = Workbook()
    book.remove(book.active)
    summary = book.create_sheet("Performance")
    summary.append(["Representation", "Estimand", "Estimate", "CI low", "CI high", "Finite draws", "Undefined draws"])
    for rep in ["ALL32"]:
        for key, d in results[rep]["pooled"].items():
            summary.append([rep, key, d["estimate"], *d["ci95"], d["finite_draws"], d["undefined_draws"]])
    restriction = book.create_sheet("Standardized restriction")
    restriction.append(["Reference", "Estimand", "Estimate", "CI low", "CI high", "Patients A", "Patients D"])
    for reference, block in results["ALL32"]["standardized_restriction"].items():
        for key in ["standardized_A_auroc", "standardized_D_auroc", "standardized_D_minus_A"]:
            d = block[key]
            restriction.append([reference, key, d["estimate"], *d["ci95"], block["patients_A"], block["patients_D"]])
    annotations = book.create_sheet("Mentor annotations")
    annotations.append(["Concept", "Controlling code", "Primary category", "Secondary 1", "Secondary 2", "Mentor comment", "Confidence", "Artifact flag", "Named-set eligible"])
    for row in names["controlling_reads"]:
        r = row["response"]
        annotations.append([row["prototype_id"], row["controlling_code"], r["primary_category"], r["secondary_category_1"], r["secondary_category_2"], r["free_text_description"], r["confidence_1_to_5"], r["artifact_uninterpretable"], row["eligible_for_named_ref"]])
    context = book.create_sheet("Context tests")
    table = pd.read_csv(source.ROOT / "e4m1_context/association_tests.csv").astype(object)
    table = table.where(pd.notna(table), None)
    context.append(table.columns.tolist())
    for row in table.itertuples(index=False, name=None):
        context.append(list(row))
    gates = book.create_sheet("Status and scope")
    gates.append(["Item", "Status or interpretation"])
    for key, status in results["gates"].items():
        gates.append([key, status])
    gates.append(["Naming gate", names["name_gate_status"]])
    gates.append(["Module I", "COMPLETE: source fits, performance/restriction bootstrap, attribution/context/attention"])
    gates.append(["Module II", "PENDING: archived target features unavailable"])
    gates.append(["Module III", "PENDING: archived task/control inputs unavailable"])
    gates.append(["Grid offset", "PENDING: canonical masks/coordinates unavailable; no reduced roster substituted"])
    gates.append(["Reader provenance", "Mentor response accepted as final per user; missing ratings/metadata remain unrecorded"])
    gates.append(["Inference", "Patient-level; fixed OOF predictions; conditional on fitted folds, vocabularies and algorithms"])
    gates.append(["Claims", "Unsupervised concept space; no mediation, causal claim, or AUROC decomposition"])
    gates.append(["Full Aim 4", "INCOMPLETE; working results workbook"])
    for sheet in book:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="244D42")
            cell.font = Font(color="FFFFFF", bold=True)
        for cells in sheet.columns:
            letter = cells[0].column_letter
            length = max(len(str(c.value or "")) for c in cells)
            sheet.column_dimensions[letter].width = min(65, max(14, length + 2))
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                elif isinstance(cell.value, float):
                    cell.number_format = "0.0000"
    book.save(OUT / "Aim4_Results.xlsx")


def main() -> None:
    OUT.mkdir(exist_ok=True)
    names = source.read_sealed(source.NAMING)
    result_path = source.ROOT / "e4m1_results/results.json"
    results = source.read_sealed(result_path)
    p = results["ALL32"]["pooled"]
    restriction = results["ALL32"]["standardized_restriction"]["A_all_primary"]
    atlas(names)
    figure(results)
    context = context_report()
    workbook(results, names)
    rows = [("Full WSI model AUROC", "all/full/auroc"), ("Concept-only AUROC", "all/concept/auroc"),
            ("Concept-only AUPRC", "all/concept/auprc"), ("Concept-minus-full AUROC", "all/concept_minus_full/auroc"),
            ("Reconstruction R²", "fidelity/r2_oof"), ("Reconstruction Pearson r", "fidelity/pearson"),
            ("Reconstructed-score AUROC", "fidelity/reconstructed_auroc"),
            ("Clinical-only AUROC", "all/clinical/auroc"), ("Concept+clinical AUROC", "all/joint/auroc"),
            ("Concept-minus-clinical AUROC", "all/concept_minus_clinical/auroc"),
            ("Concept+clinical-minus-clinical AUROC", "all/joint_minus_clinical/auroc")]
    table = "\n".join(f"| {title} | {value(p[key])} |" for title, key in rows)
    report = f"""# FINAL-v14 · Aim 4 analysis

**Working report. Module I, including the context/attention panels, is complete. The full Aim 4
campaign is not yet complete: archived inputs on `/mnt/wsl/oceanpath-hot` are
unavailable. No completed FINAL-v14 bundle has been declared.**

This report extends the frozen {link(source.REPO / 'reports/final_v13/Results.md', 'FINAL-v13 results')}.
All Aim 1–3 parent estimates and verdicts remain unchanged. The analysis follows
{link(source.REPO / 'reports/final_v14_PREREGISTRATION.md', 'the v14 preregistration')},
with the documented reader-acceptance and execution-order amendments below.
The tables and mentor annotations are also available in
{link(OUT / 'Aim4_Results.xlsx', 'the results workbook')}.

## Interpretation of Aim 1: concept fidelity and discrimination

The all-32-concept classifier achieved patient-OOF AUROC {value(p['all/concept/auroc'])}.
Its paired difference from the full WSI model was {value(p['all/concept_minus_full/auroc'])}.
The above-chance concept-discrimination gate is **SUPPORTED**. Ridge reconstruction
of the five-seed full-model native logits achieved R² {value(p['fidelity/r2_oof'])}.
These estimates describe fidelity and discrimination of the specified abundance
representation; they are not a biological mediation or AUROC decomposition.

| Estimand | Estimate [95% CI] |
|---|---:|
{table}

The MSS/pMMR plus BRAF-wild-type concept-score AUROC was {value(p['restriction/restricted_auroc'])}.
The common-support, composition-standardized restricted-minus-all difference was
{value(restriction['standardized_D_minus_A'])}. Its adjusted lower bound did not
exceed zero: restriction persistence is **NOT ESTABLISHED**. Common-support
weights and both all-patient/complete-case reference analyses are retained in
{link(result_path, 'the sealed numerical results')}.

The clinical comparison uses the inherited fold-local age/sex/site/stage encoder.
Clinical and joint scores were fitted only on outer-training patients; joint
penalties were selected with fold-local inner encoders. The inherited 1,060-patient
stage-known analysis is a fixed restriction of the same OOF predictions.
Its concept-only AUROC was {value(p['stage_known/concept/auroc'])}, compared with
{value(p['stage_known/clinical/auroc'])} for clinical-only and
{value(p['stage_known/full/auroc'])} for the full WSI model.

All 25 teacher-to-concept ridge mappings and five direct classifiers are sealed.
The source-only logistic penalty was C=0.01 in each outer fold. No target data
selected a coordinate, penalty, coefficient, or threshold. Confidence intervals
resample patients within source-subcohort×KRAS strata, pairing all scores in each
of 10,000 draws. H1.3 uses the inherited subcohort×site×KRAS strata. These intervals
condition on the fitted folds, vocabularies, and algorithms.

![Source concept discrimination and paired gaps](figures/aim4_module1.png)

{context}

## Reader assessment and montage atlas

The mentor's returned workbook is accepted as final under the user's explicit
instruction. Its original SHA-256 is
`1dc1cf7b38a3daf0c2bcf8eb2e07da9cd5c27092f9d035f835a1a699ed0673c9`.
All original response values and all 40 original montage bytes are preserved.
The user reports a same-day mentor response; this does not supply the unrecorded
row-level reviewer/date or signed blinding-attestation fields.

Only {names['named_ref_count']}/32 controlling reads meet every documented eligibility
requirement. Confidence and artifact status are absent for most rows, so
**NAME_GATE_FAIL** applies. No name-dependent classifier or formal named-morphology
claim is produced. This is a documentation-limited gate, not evidence that the
reader was unable to recognize the tissue. Among the eight duplicate pairs,
the literal categories show three exact agreements and one partial agreement;
four pairs lack a primary category. Missing artifact responses prevent these
from establishing the formal repeatability gate under the sealed conservative
amendment. No missing response is inferred from text or images.

The mentor distinguishes benign fibrosis from desmoplasia in free text. The
locked ontology combines these in one category, so that finer distinction remains
an annotation rather than a new governed feature or an analyst-assigned category.
Explore the {link(OUT / 'mentor_atlas.html', 'searchable mentor-annotated montage atlas')}.
The complete {link(source.NAMING, 'controlling-read and duplicate audit')} retains
all original comments, exclusions, mappings, and category counts.

## Interpretation of Aim 2: source-frozen cross-cohort discrimination

Pending archived target feature access. No v14 target AUROC, cross-cohort gate,
metastatic interaction, or Orion result is asserted. The frozen full-source
scoring bundle must pass its source dry run before target feature access; target
profiles and scores must be sealed before the locked outcome join. These will be
secondary analyses of previously opened cohorts, not fresh external validation.

## Interpretation of Aim 3: concept-space resolution consistency

Pending the sealed fine-task/control manifests and parent scores on the unmounted
archive. No concept-resolution boundary, control-adequacy, fine-signal, or paired
superiority verdict has been adjudicated. The five tasks and all three repeated
WT draws remain required. The operational dependency receipt is not a scientific
not-evaluable verdict and cannot substitute for completing the module.

## Remaining work and provenance

The remaining items are archived target scoring and its analyses, the
resolution ladder, the direct legacy
correspondence replay, and the prespecified raw-slide grid-offset sensitivity.
No unavailable component is treated as a negative result or silently omitted.

The accepted incomplete metadata is documented in
{link(source.ROOT / 'e4v_post_reader_xlsx/governance_amendment.json', 'the pre-unblinding amendment')}.
Independent Module I evaluation while the Module III archive is unavailable is
documented in {link(source.ROOT / 'e4m1_results/evaluation_contract.json', 'the pre-evaluation execution-order amendment')}.
No Module I result may change a later Module II/III method. The full report will
be sealed only after all required analyses and verifications finish.

The patient predictions, fit parameters, inner split identities, candidate
losses/failures, and bootstrap arrays are bound by
{link(source.OUT / 'source_fit_summary.json', 'the source fit receipt')} and
{link(result_path, 'the results seal')}. Publication figures are available as
{link(OUT / 'figures/aim4_module1.pdf', 'PDF')} and
{link(OUT / 'figures/aim4_module1.svg', 'SVG')}.
"""
    (OUT / "Results.md").write_text(report)
    status = {"status": "WORKING_FINAL_V14_AIM4_INCOMPLETE", "created_utc": source.io.utc_now(),
              "completed": ["accepted reader intake", "naming and repeatability audit", "Module I source fits", "Module I performance and restriction bootstrap", "context/attention attribution and coefficient stability"],
              "pending": ["Module II targets", "Module III ladder", "legacy correspondence replay", "grid-offset sensitivity", "final audit and complete bundle seal"],
              "archive_mount": "/mnt/wsl/oceanpath-hot", "user_input_needed": "Restore access to existing E:\\WSL\\oceanpath-hot.vhdx", "completed_bundle": False,
              "inputs": [source.io.identity(source.NAMING), source.io.identity(result_path), source.io.identity(source.ROOT / "e4m1_context/results.json")],
              "report_artifacts": [r for r in source.io.tree_inventory(OUT) if r["path"] != "STATUS.json"], "code": source.io.identity(Path(__file__))}
    (OUT / "STATUS.json").write_bytes(source.io.json_bytes(status))
    print(json.dumps({"status": status["status"], "report": str(OUT / 'Results.md')}, indent=2))


if __name__ == "__main__":
    main()
