from __future__ import annotations

import datetime as dt
import importlib.util
import io
import json
import shutil
import sys
import zipfile
from collections.abc import Callable
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.workbook.defined_name import DefinedName

REPOSITORY = Path(__file__).resolve().parents[1]
SCRIPT = REPOSITORY / "tools/final_v14_excel_handoff.py"
SPEC = importlib.util.spec_from_file_location("final_v14_excel_handoff_test", SCRIPT)
assert SPEC is not None and SPEC.loader is not None
HANDOFF = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = HANDOFF
SPEC.loader.exec_module(HANDOFF)


def _code(index: int) -> str:
    alphabet = "ABCDEFGHIJKLMNOPQRSTUVWXYZ234567"
    value = index
    output = []
    for _ in range(6):
        output.append(alphabet[value % len(alphabet)])
        value //= len(alphabet)
    return "".join(reversed(output))


def _rows() -> list[dict[str, str]]:
    return [
        {
            "presentation_order": str(index),
            "blinded_code": _code(index),
            "n_tiles": "12",
            **{column: "" for column in HANDOFF.RESPONSE_COLUMNS},
        }
        for index in range(1, 41)
    ]


def _rubric() -> str:
    categories = "\n".join(f"- {category}" for category in HANDOFF.ONTOLOGY)
    return f"# Fixed morphology rubric\n\n{categories}\n"


def _template() -> bytes:
    return HANDOFF.build_workbook(
        _rows(), _rubric(), "2026-09-04T05:15:42+00:00"
    )


def _completed_payload(
    mutate: Callable[[object], None] | None = None,
) -> bytes:
    workbook = load_workbook(io.BytesIO(_template()), data_only=False)
    review = workbook[HANDOFF.REVIEW_SHEET]
    for row in range(2, 42):
        review.cell(row, 4).value = "complete"
        review.cell(row, 5).value = HANDOFF.ONTOLOGY[0]
        review.cell(row, 6).value = HANDOFF.ONTOLOGY[1]
        review.cell(row, 7).value = ""
        review.cell(row, 8).value = "Gland-forming malignant epithelium."
        review.cell(row, 9).value = 4
        review.cell(row, 10).value = "no"
        review.cell(row, 11).value = "reviewer-01"
        review.cell(row, 12).value = dt.date(2026, 9, 4)
        review.cell(row, 13).value = "confirmed_no_key_access"
    if mutate is not None:
        mutate(review)
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


@pytest.fixture
def isolated_governed_handoff(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> dict[str, object]:
    source = HANDOFF.validate_source_package()
    delivery_root = tmp_path / "reviews/v14_xlsx"
    audit_root = tmp_path / "e4v_pre_reader_xlsx_handoff"
    monkeypatch.setattr(HANDOFF, "DELIVERY_ROOT", delivery_root)
    monkeypatch.setattr(
        HANDOFF, "DELIVERY_PACKAGE", delivery_root / "FOR_PATHOLOGIST"
    )
    monkeypatch.setattr(HANDOFF, "AUDIT_ROOT", audit_root)
    monkeypatch.setattr(
        HANDOFF, "HANDOFF_RECEIPT", audit_root / "handoff_receipt.json"
    )
    HANDOFF._build_delivery_stage(
        delivery_root, source, "2026-09-04T12:00:00+00:00"
    )
    audit_root.mkdir(parents=True)
    shutil.copytree(delivery_root / "FOR_PATHOLOGIST", audit_root / "public")
    delivery = HANDOFF.validate_delivery_package(source)
    HANDOFF.write_new(
        HANDOFF.HANDOFF_RECEIPT,
        HANDOFF.json_bytes(
            HANDOFF._handoff_receipt_payload(
                delivery, created_utc="2026-09-04T12:00:01+00:00"
            )
        ),
    )
    assert HANDOFF.verify_handoff()["status"] == "PASS"
    return {"source": source, "delivery": delivery, "tmp_path": tmp_path}


def test_workbook_is_deterministic_and_semantically_valid() -> None:
    first = _template()
    second = _template()
    assert first == second
    result = HANDOFF.validate_template_workbook(first, _rows(), _rubric())
    assert result == {
        "status": "PASS",
        "rows": 40,
        "dropdown_columns": 7,
        "date_validation_columns": 1,
        "montage_hyperlinks": 40,
        "macros": 0,
        "external_data_links": 0,
    }


def test_core_property_clock_does_not_change_workbook_bytes() -> None:
    first = _template()
    changed = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(first), "r") as source, zipfile.ZipFile(
        changed, "w"
    ) as destination:
        for name in source.namelist():
            payload = source.read(name)
            if name == "docProps/core.xml":
                assert b"2026-09-04T05:15:42Z" in payload
                payload = payload.replace(
                    b"2026-09-04T05:15:42Z", b"2099-12-31T23:59:59Z"
                )
            destination.writestr(name, payload)
    assert changed.getvalue() != first
    assert (
        HANDOFF._canonicalize_xlsx(
            changed.getvalue(), "2026-09-04T05:15:42+00:00"
        )
        == first
    )


def test_workbook_has_exact_table_dropdowns_and_protection() -> None:
    workbook = load_workbook(io.BytesIO(_template()), data_only=False)
    assert workbook.sheetnames == HANDOFF.EXPECTED_SHEETS
    assert workbook[HANDOFF.OPTIONS_SHEET].sheet_state == "veryHidden"
    review = workbook[HANDOFF.REVIEW_SHEET]
    assert review.freeze_panes == "D2"
    assert review.protection.sheet
    assert review.tables[HANDOFF.TABLE_NAME].ref == "A1:M41"
    assert [review.cell(1, column).value for column in range(1, 14)] == HANDOFF.FORM_COLUMNS
    assert all(review.cell(2, column).protection.locked for column in range(1, 4))
    assert all(not review.cell(2, column).protection.locked for column in range(4, 14))
    validations = {
        str(validation.sqref): (validation.type, validation.formula1)
        for validation in review.data_validations.dataValidation
    }
    assert validations == {
        "D2:D41": ("list", "=_review_status"),
        "E2:E41": ("list", "=_morphology_categories"),
        "F2:F41": ("list", "=_morphology_categories"),
        "G2:G41": ("list", "=_morphology_categories"),
        "I2:I41": ("list", "=_confidence"),
        "J2:J41": ("list", "=_artifact"),
        "L2:L41": ("date", "DATE(2000,1,1)"),
        "M2:M41": ("list", "=_attestation"),
    }
    assert review["B2"].hyperlink.target == f"montages/{_rows()[0]['blinded_code']}.jpg"


def test_valid_completed_workbook_exports_canonical_rows() -> None:
    completed = HANDOFF.validate_completed_workbook(
        _completed_payload(), _rows(), _rubric()
    )
    assert len(completed) == 40
    assert list(completed[0]) == HANDOFF.FORM_COLUMNS
    assert completed[0]["confidence_1_to_5"] == 4
    assert completed[0]["review_date"] == "2026-09-04"
    assert {row["reviewer_id"] for row in completed} == {"reviewer-01"}


@pytest.mark.parametrize(
    ("mutation", "message"),
    [
        (lambda sheet: setattr(sheet["A2"], "value", 40), "fixed identity/order"),
        (lambda sheet: setattr(sheet["E2"], "value", "invented category"), "primary"),
        (
            lambda sheet: setattr(sheet["F2"], "value", HANDOFF.ONTOLOGY[0]),
            "categories repeat",
        ),
        (lambda sheet: setattr(sheet["H2"], "value", "=1+1"), "formulas are forbidden"),
        (lambda sheet: setattr(sheet["I2"], "value", 6), "confidence"),
        (lambda sheet: setattr(sheet["J2"], "value", "maybe"), "artifact flag"),
        (lambda sheet: setattr(sheet["K2"], "value", "another-reviewer"), "reviewer_id"),
        (lambda sheet: setattr(sheet["K2"], "value", "+CMD"), "unsafe spreadsheet"),
        (lambda sheet: setattr(sheet["H2"], "value", "@external"), "unsafe spreadsheet"),
        (lambda sheet: setattr(sheet["L2"], "value", "09/04/2026"), "review_date"),
        (lambda sheet: setattr(sheet["L2"], "value", "1900-01-01"), "outside"),
        (lambda sheet: setattr(sheet["M2"], "value", "yes"), "attestation"),
        (lambda sheet: setattr(sheet["N1"], "value", "undeclared"), "undeclared"),
        (
            lambda sheet: setattr(sheet.data_validations, "dataValidation", []),
            "data-validation",
        ),
        (lambda sheet: setattr(sheet.protection, "sheet", False), "guardrails"),
    ],
)
def test_completed_workbook_tampering_fails_closed(
    mutation: Callable[[object], None], message: str
) -> None:
    with pytest.raises(HANDOFF.HandoffError, match=message):
        HANDOFF.validate_completed_workbook(
            _completed_payload(mutation), _rows(), _rubric()
        )


def test_canonical_csv_preserves_unicode_and_column_order() -> None:
    completed = HANDOFF.validate_completed_workbook(
        _completed_payload(), _rows(), _rubric()
    )
    payload = HANDOFF._csv_payload(completed)
    text = payload.decode("utf-8")
    assert text.splitlines()[0] == ",".join(HANDOFF.FORM_COLUMNS)
    assert "gland–lumen" in text
    assert len(text.splitlines()) == 41


def test_completed_workbook_rejects_extra_defined_name() -> None:
    workbook = load_workbook(io.BytesIO(_completed_payload()), data_only=False)
    workbook.defined_names.add(
        DefinedName(
            "undeclared_external_name",
            attr_text="'https://example.invalid/[book.xlsx]Sheet1'!$A$1",
        )
    )
    stream = io.BytesIO()
    workbook.save(stream)
    with pytest.raises(HANDOFF.HandoffError, match="defined-name census"):
        HANDOFF.validate_completed_workbook(stream.getvalue(), _rows(), _rubric())


def test_completed_workbook_rejects_extra_hidden_content() -> None:
    workbook = load_workbook(io.BytesIO(_completed_payload()), data_only=False)
    workbook[HANDOFF.OPTIONS_SHEET]["Z100"] = "undeclared"
    stream = io.BytesIO()
    workbook.save(stream)
    with pytest.raises(HANDOFF.HandoffError, match="option worksheet cell census"):
        HANDOFF.validate_completed_workbook(stream.getvalue(), _rows(), _rubric())


def test_xlsx_zip_rejects_active_content_family() -> None:
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as archive:
        archive.writestr("[Content_Types].xml", b"<Types/>")
        archive.writestr("xl/connections.xml", b"<connections/>")
    with pytest.raises(HANDOFF.HandoffError, match="active/external content"):
        HANDOFF._safe_zip_members(stream.getvalue(), "hostile workbook")


def test_governed_verifier_fails_closed_on_coordinator_and_receipt_tampering(
    isolated_governed_handoff: dict[str, object],
) -> None:
    readme = HANDOFF.DELIVERY_ROOT / "README_COORDINATOR.md"
    status = HANDOFF.DELIVERY_ROOT / "PRE_READER_STATUS.json"
    receipt = HANDOFF.HANDOFF_RECEIPT

    original = readme.read_bytes()
    readme.write_text("unsafe directions\n", encoding="utf-8")
    with pytest.raises(HANDOFF.HandoffError, match="coordinator instructions"):
        HANDOFF.verify_handoff()
    readme.write_bytes(original)

    original = status.read_bytes()
    status.write_text("{}\n", encoding="utf-8")
    with pytest.raises(HANDOFF.HandoffError, match="pre-reader status"):
        HANDOFF.verify_handoff()
    status.write_bytes(original)

    original = receipt.read_bytes()
    receipt.unlink()
    with pytest.raises(HANDOFF.HandoffError, match="receipt directory census"):
        HANDOFF.verify_handoff()
    receipt.write_bytes(original)

    changed = json.loads(original.decode("utf-8"))
    changed["builder"]["sha256"] = "0" * 64
    receipt.write_bytes(HANDOFF.json_bytes(changed))
    with pytest.raises(HANDOFF.HandoffError, match="identities or content"):
        HANDOFF.verify_handoff()


def test_ingest_seals_raw_xlsx_and_exports_canonical_csv_append_only(
    isolated_governed_handoff: dict[str, object],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = isolated_governed_handoff["source"]
    assert isinstance(source, dict)
    workbook = load_workbook(HANDOFF.DELIVERY_PACKAGE / HANDOFF.WORKBOOK_NAME)
    review = workbook[HANDOFF.REVIEW_SHEET]
    for row in range(2, 42):
        review.cell(row, 4).value = "complete"
        review.cell(row, 5).value = HANDOFF.ONTOLOGY[0]
        review.cell(row, 6).value = HANDOFF.ONTOLOGY[1]
        review.cell(row, 8).value = "Gland-forming malignant epithelium."
        review.cell(row, 9).value = 4
        review.cell(row, 10).value = "no"
        review.cell(row, 11).value = "reviewer-01"
        review.cell(row, 12).value = dt.date(2026, 9, 4)
        review.cell(row, 13).value = "confirmed_no_key_access"
    temporary_root = isolated_governed_handoff["tmp_path"]
    assert isinstance(temporary_root, Path)
    completed_path = temporary_root / HANDOFF.COMPLETED_WORKBOOK_NAME
    workbook.save(completed_path)
    post_reader = temporary_root / "post_reader/human_reader_1"
    monkeypatch.setattr(HANDOFF, "POST_READER_OUTPUT", post_reader)

    receipt = HANDOFF.ingest_completed(completed_path)
    assert receipt["status"] == "PASS"
    assert receipt["rows"] == 40
    assert receipt["safe_to_begin_unblinding"] is True
    assert (post_reader / HANDOFF.COMPLETED_WORKBOOK_NAME).read_bytes() == completed_path.read_bytes()
    csv_text = (post_reader / HANDOFF.COMPLETED_CSV_NAME).read_text(encoding="utf-8")
    assert len(csv_text.splitlines()) == 41
    assert csv_text.splitlines()[0] == ",".join(HANDOFF.FORM_COLUMNS)
    assert (post_reader / "receipt.json").stat().st_mode & 0o777 == 0o400
    with pytest.raises(HANDOFF.HandoffError, match="append-only output"):
        HANDOFF.ingest_completed(completed_path)
