"""Packaging invariants: scientific completion, lineage, and presentation drift."""
import json
import os
from pathlib import Path

import pytest
from openpyxl import Workbook

from tools import final_v14_bundle_receipt as bundle


def write_json(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value))
    return path


def sealed(path, value):
    write_json(path, value)
    write_json(Path(str(path) + ".seal.json"), {"artifact": bundle.identity(path)})
    return path


def completed_states():
    return {
        "report": {"status": "REPORT_READY_FOR_FINAL_AUDIT", "analysis_complete": True, "completed_bundle": False},
        "m1_source": {"status": "SOURCE_MODELS_SEALED"}, "m1": {"gates": {"H1.1": "NOT_ESTABLISHED"}, "ALL32": {}},
        "m1_audit": {"status": "PASS_SOURCE_COMPLETED_CAMPAIGN_INCOMPLETE", "source_models": {"status": "PASS"}, "source_inference": {"status": "PASS"}, "context": {"status": "PASS"}},
        "context": {"status": "CONTEXT_ASSOCIATIONS_COMPLETE"},
        "m2": {"status": "MODULE_II_ANALYSIS_COMPLETE"}, "m2_audit": {"status": "INDEPENDENT_MODULE_II_NUMERICAL_AUDIT_PASS"},
        "m3": {"status": "COMPLETE"}, "m3_source": {"status": "SOURCE_FITS_COMPLETE"}, "m3_audit": {"status": "PASS"},
        "correspondence": {"status": "CORRESPONDENCE_GEOMETRY_COMPLETE"}, "correspondence_audit": {"status": "PASS_INDEPENDENT_CORRESPONDENCE_REPLAY"},
        "variant_mapping": {"status": "CANONICAL_NINE_VARIANT_MAPPING_COMPLETE"},
        "variant_audit": {"status": "INDEPENDENT_NINE_VARIANT_MAPPING_AUDIT_PASS"},
        "grid": {"status": "GRID_OFFSET_SENSITIVITY_COMPLETE"}, "grid_audit": {"status": "GRID_OFFSET_INDEPENDENT_AUDIT_PASS"},
    }


def test_operational_grid_block_never_completes_campaign():
    data = completed_states()
    bundle.check_states(data)
    data["grid"] = {"status": "BLOCKED_MISSING_CANONICAL_ROOT"}
    with pytest.raises(bundle.AuditError, match="Actual grid sensitivity"):
        bundle.check_states(data)
    data["grid"] = {"status": "GRID_OFFSET_NOT_EVALUABLE", "eligible_counts": {"TCGA-COAD": 0}, "patient_roster_drawn": False, "slide_extraction_run": False}
    with pytest.raises(bundle.AuditError, match="Operational grid block"):
        bundle.check_states(data)
    data["grid_audit"].update(prespecified_insufficient_eligible_verified=True, operational_missing_dependencies=False)
    bundle.check_states(data)


@pytest.mark.parametrize("component", ["context", "m2_audit", "m3_source", "correspondence", "variant_mapping"])
def test_each_independent_stage_is_mandatory(component):
    data = completed_states()
    data[component]["status"] = "PENDING"
    with pytest.raises(bundle.AuditError):
        bundle.check_states(data)


def test_report_cannot_preempt_completion_receipt():
    data = completed_states()
    data["report"]["completed_bundle"] = True
    with pytest.raises(bundle.AuditError, match="pre-empt"):
        bundle.check_states(data)


def test_seal_rejects_conflicting_duplicate_sha(tmp_path):
    p = sealed(tmp_path / "result.json", {"status": "COMPLETE"})
    sp = Path(str(p) + ".seal.json")
    s = json.loads(sp.read_text())
    s["sha256"] = "0" * 64
    write_json(sp, s)
    with pytest.raises(bundle.AuditError, match="Conflicting seal"):
        bundle.verify_seal(p)


def test_closure_reaches_code_and_ci_and_detects_tampering(tmp_path):
    ci = tmp_path / "bootstrap.npz"
    ci.write_bytes(b"synthetic CI storage")
    code = tmp_path / "runner.py"
    code.write_text("# frozen source\n")
    result = sealed(tmp_path / "result.json", {"artifacts": [bundle.identity(ci)], "code": bundle.identity(code)})
    closure = bundle.IdentityClosure(tmp_path)
    closure.file(result, sealed=True)
    assert {r["path"] for r in closure.inventory()} == {str(x) for x in [ci, code, result, Path(str(result) + ".seal.json")]}
    ci.write_bytes(b"changed CI draws")
    with pytest.raises(bundle.AuditError, match="Hash mismatch"):
        bundle.IdentityClosure(tmp_path).file(result, sealed=True)


def test_historical_report_identity_resolves_only_exact_snapshot_hash(tmp_path):
    current = tmp_path / "current.md"
    current.write_text("working report")
    old = bundle.identity(current)
    snap = tmp_path / "snapshot.md"
    snap.write_bytes(current.read_bytes())
    current.write_text("completed report")
    mapping = {(str(current), old["sha256"]): snap}
    closure = bundle.IdentityClosure(tmp_path, path_map=mapping)
    row = closure.add(old)
    assert row["path"] == str(snap) and row["historical_path"] == str(current)
    with pytest.raises(bundle.AuditError, match="Hash mismatch"):
        bundle.IdentityClosure(tmp_path).add(old)


def test_old_source_audit_must_bind_current_results_through_preserved_status(tmp_path):
    result = write_json(tmp_path / "m1.json", {"estimate": .60})
    context = write_json(tmp_path / "context.json", {"status": "COMPLETE"})
    status = write_json(tmp_path / "old_STATUS.json", {"inputs": [bundle.identity(result)]})
    old_pin = bundle.identity(status)
    old_pin["path"] = str(tmp_path / "now_replaced_STATUS.json")
    audit = {"report": {"report_status": old_pin}, "context": {"results": bundle.identity(context)}}
    mapping = {(old_pin["path"], old_pin["sha256"]): status}
    paths = {"m1": result, "context": context}
    bundle.require_source_audit_bindings(audit, paths, mapping)
    write_json(result, {"estimate": .99})
    with pytest.raises(bundle.AuditError, match="Module I audited report status"):
        bundle.require_source_audit_bindings(audit, paths, mapping)


def test_inherited_pack_requires_current_exact_stat_and_explicit_method(tmp_path, monkeypatch):
    pack = tmp_path / "features.bin"
    pack.write_bytes(b"synthetic packed input")
    pin = bundle.identity(pack)
    inherited = {(str(pack), pin["sha256"]): {**pin, "mtime_ns": pack.stat().st_mtime_ns, "evidence": []}}
    monkeypatch.setattr(bundle, "verify_identity", lambda *_: pytest.fail("Inherited input must not be described as rehashed"))
    closure = bundle.IdentityClosure(tmp_path, inherited=inherited)
    assert closure.add(pin)["verification"].startswith("INHERITED_PRE_READER")
    os.utime(pack, ns=(pack.stat().st_atime_ns, pack.stat().st_mtime_ns + 1))
    with pytest.raises(bundle.AuditError, match="stat guard changed"):
        bundle.IdentityClosure(tmp_path, inherited=inherited).add(pin)


def make_report(tmp_path, rows):
    (tmp_path / "tables").mkdir()
    import csv
    with (tmp_path / "tables/metrics.csv").open("w", newline="") as handle:
        csv.writer(handle).writerows(rows)
    wb = Workbook()
    ws = wb.active
    ws.title = "Metrics"
    for row in rows:
        ws.append(row)
    wb.save(tmp_path / "Aim4_Results.xlsx")
    return {"table_bindings": [{"sheet": "Metrics", "csv": "tables/metrics.csv"}]}


def test_workbook_and_csv_match_scientific_numbers_and_gate_strings(tmp_path):
    rows = [["task", "estimate", "status"], ["codon", .519735577846958, "NOT_ESTABLISHED"]]
    status = make_report(tmp_path, rows)
    assert bundle.report_tables(tmp_path, status, expected_tables={"Metrics": rows})["table_rows"] == {"Metrics": 1}
    from openpyxl import load_workbook
    wb = load_workbook(tmp_path / "Aim4_Results.xlsx")
    wb["Metrics"]["C2"] = "SUPPORTED"
    wb.save(tmp_path / "Aim4_Results.xlsx")
    with pytest.raises(bundle.AuditError, match="Table mismatch"):
        bundle.report_tables(tmp_path, status, expected_tables={"Metrics": rows})


def test_shared_csv_workbook_numeric_corruption_fails_scientific_replay(tmp_path):
    wrong = [["task", "estimate"], ["codon", .99]]
    status = make_report(tmp_path, wrong)
    with pytest.raises(bundle.AuditError, match="scientific inputs"):
        bundle.report_tables(tmp_path, status, expected_tables={"Metrics": [["task", "estimate"], ["codon", .5197]]})


def test_shared_status_markdown_corruption_fails_scientific_replay(tmp_path):
    rows = [["task", "estimate"], ["codon", .5197]]
    status = make_report(tmp_path, rows)
    status["table_bindings"][0]["markdown_table"] = "Fine signal SUPPORTED; AUROC .99"
    (tmp_path / "Results.md").write_text(status["table_bindings"][0]["markdown_table"])
    with pytest.raises(bundle.AuditError, match="Markdown binding drift"):
        bundle.report_tables(tmp_path, status, expected_tables={"Metrics": rows},
                             expected_markdown={"Metrics": "Fine signal NOT_ESTABLISHED; AUROC .5197"})


def test_inherited_preflight_must_match_independent_fixed_pin(tmp_path):
    auditor = tmp_path / "auditor.py"
    auditor.write_text("# independent source\n")
    write_json(tmp_path / "e4v_pre_reader_independent_audit/audit_receipt.json",
               {"status": "PASS", "gates": {"core_identities": {"status": "PASS"}}, "auditor": bundle.identity(auditor)})
    write_json(tmp_path / "e4v_pre_reader/receipts/preflight.json", {"deep_hash": True})
    with pytest.raises(bundle.AuditError, match="independent auditor's fixed pin"):
        bundle.inherited_pack_pins(tmp_path)


def test_local_links_resolve_spaces_and_fail_missing_artifacts(tmp_path):
    image = tmp_path / "my figure.svg"
    image.write_text("<svg/>")
    report = tmp_path / "Results.md"
    report.write_text("[Figure](my%20figure.svg)\n[External](https://example.org)\n[Anchor](#result)\n")
    assert bundle.linked_files(report) == [image]
    image.unlink()
    with pytest.raises(bundle.AuditError, match="Broken report artifact link"):
        bundle.linked_files(report)


def test_final_receipt_is_last_acyclic_and_replayable(tmp_path):
    write_json(tmp_path / "STATUS.json", {"analysis_complete": True, "completed_bundle": False})
    result = {"status": "PASS_FINAL_CAMPAIGN_AUDIT", "artifacts": [], "local_report_links_verified": 0,
              "working_report_path_mapping": 10, "limitations": ["Synthetic packaging test only."]}
    receipt = bundle.publish(tmp_path, result)
    assert receipt["status"] == "COMPLETED"
    assert bundle.publish(tmp_path, result) == receipt
    assert bundle.verify_bundle(tmp_path)["status"] == "PASS"
    assert "final_bundle_receipt" not in (tmp_path / "source_manifest.json").read_text()
    (tmp_path / "Audit.md").write_text("tampered audit")
    with pytest.raises(bundle.AuditError, match="Hash mismatch"):
        bundle.verify_bundle(tmp_path)
